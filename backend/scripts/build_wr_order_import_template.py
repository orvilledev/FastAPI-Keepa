"""Turn the WR order import workbook into the blank template the app ships.

The source file is the warehouse's "wr order template save as txt <date>.xlsx",
which arrives with a worked example in rows 2-14 and a font element Excel wrote
outside the schema's range (`<family val="34"/>`), which openpyxl refuses to
read. This strips the example rows and clamps the font family, editing the parts
in place so every style, column width, comment and helper sheet survives.

Usage:
    python scripts/build_wr_order_import_template.py "<source .xlsx>"
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

DEST = (
    Path(__file__).resolve().parent.parent
    / "app"
    / "static"
    / "shipment_manager"
    / "WR_ORDER_IMPORT_TEMPLATE.xlsx"
)

SHEET_PART = "xl/worksheets/sheet1.xml"
_FAMILY = re.compile(r'<family val="(\d+)"\s*/>')
_ROW = re.compile(r'<row r="(\d+)"[^>]*?(?:/>|>.*?</row>)', re.DOTALL)
_DIMENSION = re.compile(r'<dimension ref="[^"]*"/>')
_MAX_FAMILY = 14


def _clamp_font_family(xml: str) -> str:
    def replace(match: re.Match[str]) -> str:
        value = int(match.group(1))
        return match.group(0) if value <= _MAX_FAMILY else '<family val="2"/>'

    return _FAMILY.sub(replace, xml)


def _keep_header_row_only(xml: str) -> str:
    kept = _ROW.sub(lambda m: m.group(0) if m.group(1) == "1" else "", xml)
    return _DIMENSION.sub('<dimension ref="A1:AI1"/>', kept)


def build(source: Path, dest: Path = DEST) -> Path:
    buffer = io.BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(
        buffer, "w", zipfile.ZIP_DEFLATED
    ) as out:
        for item in archive.infolist():
            data = archive.read(item.filename)
            if item.filename.endswith(".xml"):
                xml = _clamp_font_family(data.decode("utf-8"))
                if item.filename == SHEET_PART:
                    xml = _keep_header_row_only(xml)
                data = xml.encode("utf-8")
            out.writestr(item, data)

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(buffer.getvalue())
    return dest


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    source = Path(sys.argv[1])
    if not source.is_file():
        print(f"No such file: {source}")
        return 1
    written = build(source)

    from openpyxl import load_workbook

    workbook = load_workbook(written)
    sheet = workbook["Order Import Template"]
    print(f"Wrote {written} ({written.stat().st_size:,} bytes)")
    print("sheets:", workbook.sheetnames)
    print("rows:", sheet.max_row, "cols:", sheet.max_column)
    print("headers:", [cell.value for cell in sheet[1]])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
