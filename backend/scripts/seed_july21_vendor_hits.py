"""Seed July 21, 2026 daily analytics from itemized off-price Excel reports.

Aggregates Seller column hits per vendor (same display format as live Daily Run
archives), then force-rebuilds week / month / year.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl

from app.database import get_supabase
from app.repositories.off_price_analytics_snapshot_repository import (
    OffPriceAnalyticsSnapshotRepository,
)
from app.services.off_price_analytics_service import OffPriceAnalyticsService
from app.services.off_price_analytics_vendors import (
    VENDOR_DEFS,
    is_excluded_analytics_seller,
)

DAY = "2026-07-21"
START = "2026-07-21T00:00:00+00:00"
END = "2026-07-22T00:00:00+00:00"

# Downloads folder workbooks for 7.21.26
DOWNLOADS = Path(r"C:\Users\Administrator\Downloads")
VENDOR_FILES: Dict[str, Path] = {
    "obz": DOWNLOADS / "OBZ Off Price 7.21.26.xlsx",
    "dnk": DOWNLOADS / "DNK Off Price 7.21.26.xlsx",
    "clk": DOWNLOADS / "CLK Off Price 7.21.26.xlsx",
    "sff": DOWNLOADS / "SFF Off Price Report 7.21.26.xlsx",
    "tev": DOWNLOADS / "TEV Off Prive Report 7.21.26.xlsx",
    "bor": DOWNLOADS / "BOR Off Price Report 7.21.26.xlsx",
}


def _seller_counts_from_xlsx(path: Path) -> Tuple[int, List[Dict[str, object]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return 0, []
    cols = {str(c or "").strip().lower(): i for i, c in enumerate(header)}
    seller_idx = cols.get("seller")
    if seller_idx is None:
        raise ValueError(f"No Seller column in {path.name}: {header}")

    counts: Dict[str, int] = defaultdict(int)
    display: Dict[str, str] = {}
    total = 0
    for row in rows:
        if not row or seller_idx >= len(row):
            continue
        raw = row[seller_idx]
        seller = str(raw or "").strip()
        if not seller:
            continue
        if is_excluded_analytics_seller(seller):
            continue
        key = seller.lower()
        display.setdefault(key, seller)
        counts[key] += 1
        total += 1

    sellers = sorted(
        [{"seller_name": display[k], "hits": c} for k, c in counts.items()],
        key=lambda item: (-int(item["hits"]), str(item["seller_name"]).lower()),
    )
    return total, sellers


def main() -> None:
    for code, path in VENDOR_FILES.items():
        if not path.is_file():
            raise FileNotFoundError(path)

    vendor_data: Dict[str, Tuple[int, List[Dict[str, object]]]] = {}
    for code, path in VENDOR_FILES.items():
        total, sellers = _seller_counts_from_xlsx(path)
        vendor_data[code] = (total, sellers)
        print(f"{code}: {total} hits, {len(sellers)} sellers from {path.name}")

    db = get_supabase()
    snaps = OffPriceAnalyticsSnapshotRepository(db)
    svc = OffPriceAnalyticsService(db)

    vendors = []
    total = 0
    runs = 0
    all_sellers: set[str] = set()
    for code, name in VENDOR_DEFS:
        off, sellers = vendor_data.get(code, (0, []))
        run_count = 1 if off > 0 else 0
        total += off
        runs += run_count
        for s in sellers:
            sn = str(s.get("seller_name") or "").strip().lower()
            if sn:
                all_sellers.add(sn)
        vendors.append(
            {
                "code": code,
                "name": name,
                "off_price_count": off,
                "run_count": run_count,
                "scheduler_enabled": False,
                "tracking_enabled": True,
                "sellers": sellers,
            }
        )
    vendors.sort(key=lambda v: (-v["off_price_count"], v["code"]))
    vendors_with_hits = sum(1 for v in vendors if v["off_price_count"] > 0)

    row = {
        "period_type": "daily",
        "period_key": DAY,
        "period_label": "Jul 21, 2026",
        "period_start": START,
        "period_end": END,
        "total_off_price_count": total,
        "total_run_count": runs,
        "distinct_sellers": len(all_sellers),
        "vendors_with_hits": vendors_with_hits,
        "payload": {
            "vendors": vendors,
            "total_off_price_count": total,
            "total_run_count": runs,
            "distinct_sellers": len(all_sellers),
            "vendors_with_hits": vendors_with_hits,
            "recovery_note": (
                "Seeded July 21 from itemized off-price Excel reports "
                "(seller hit aggregates; ASIN-level rows not stored in Analytics)."
            ),
            "seeded_at": datetime.now(timezone.utc).isoformat(),
            "seed_source": "itemized_xlsx_7.21.26",
        },
        "source": "live",
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    snaps.upsert_snapshot(row)
    print(f"\nSeeded daily/{DAY}: hits={total} runs={runs} distinct_sellers={len(all_sellers)}")

    print("Rebuilding week / month / year…")
    now = datetime.now(timezone.utc)
    for period in ("weekly", "monthly", "yearly"):
        out = svc.get_off_price_summary(
            period,
            offset=0,
            persist=True,
            force_persist=True,
            user_id=None,
            reference=now,
        )
        print(
            f"  {period}/{out.get('period_key')}: "
            f"hits={out.get('total_off_price_count')} "
            f"runs={out.get('total_run_count')} "
            f"sellers={out.get('distinct_sellers')}"
        )
    print("Done.")


if __name__ == "__main__":
    main()
