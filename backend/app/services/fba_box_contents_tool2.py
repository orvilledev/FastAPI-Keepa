"""FBA Box Contents Tool #2 — spaced PO# carton dump → Box Contents workbook.

Input (Amazon carton export with blank spacer columns, .xls or .xlsx):
  - Row ``PO#:`` / shipment id
  - Header row ``Sku`` / ``UPC`` / ``Description`` / ``Qty`` / ``Weight`` /
    ``Carton Length`` / ``Carton Width`` / ``Carton Height`` (often with blanks
    between labels)
  - Repeating ``Carton#:`` blocks, item lines, and a ``Total`` row whose Weight
    cell is copied verbatim into Dimensions

Output (differs from Tool #1):
  - ``Box Contents``: UPC, Box Number, QTY + Sum-of-QTY pivot at column G
    (no Total QTY footer; no mauve header fill)
  - ``Dimensions``: Weight (text), Length, Width, Height — no Box # column
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_LEFT = Alignment(horizontal="left", vertical="center")

BOX_CONTENTS_SHEET = "Box Contents"
DIMENSIONS_SHEET = "Dimensions"
DEFAULT_OUTPUT_FILENAME = "FBA Box Contents Output.xlsx"
PIVOT_START_COL = 7  # column G


class FbaBoxContentsTool2Error(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class ContentRow:
    upc: str
    box_number: int
    qty: int | float


@dataclass
class Carton:
    box_number: int
    carton_id: str
    length: int | float | None
    width: int | float | None
    height: int | float | None
    weight_text: str = ""
    rows: list[ContentRow] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedCartonDetail:
    shipment_id: str
    cartons: tuple[Carton, ...]
    rows: tuple[ContentRow, ...]


@dataclass(frozen=True)
class FbaBoxContentsTool2Result:
    file_bytes: bytes
    filename: str
    row_count: int
    box_count: int
    upc_count: int
    total_qty: int | float
    shipment_id: str


def sanitize_download_filename(name: str | None, fallback: str = DEFAULT_OUTPUT_FILENAME) -> str:
    """Return ``{stem} Output.xlsx`` from the uploaded carton-detail filename."""
    cleaned = (name or "").replace('"', "").replace("\r", "").replace("\n", "").strip()
    base = Path(cleaned).name if cleaned else ""
    stem = Path(base).stem if base else ""
    if stem.lower().endswith(" output"):
        stem = stem[: -len(" output")].rstrip()
    stem = re.sub(r'[<>:"/\\|?*]', "", stem).strip(" .")
    if not stem:
        return fallback
    return f"{stem} Output.xlsx"


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return format(value, ".15g").strip()
    return str(value).strip()


def _normalize_label(value: object) -> str:
    return re.sub(r"\s+", " ", _cell_text(value)).strip().lower()


def _as_upc_string(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        if value < 0:
            return ""
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value) or value < 0:
            return ""
        if abs(value) < 2**53:
            return str(int(value)) if value.is_integer() else str(int(round(value)))
        return format(value, ".0f")
    text = str(value).strip()
    if not text:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _looks_like_upc(value: object) -> bool:
    text = _as_upc_string(value).replace(" ", "")
    return bool(text) and text.isdigit() and 8 <= len(text) <= 14


def _as_excel_number(value: object) -> int | float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer() and abs(value) < 2**53:
            return int(value)
        return value
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    if number.is_integer() and abs(number) < 2**53:
        return int(number)
    return number


def _as_qty(value: object) -> int | float:
    parsed = _as_excel_number(value)
    return 0 if parsed is None else parsed


def _is_carton_header(value: object) -> bool:
    label = _normalize_label(value).replace(" ", "")
    return label in {"carton#:", "carton#", "carton:"}


def _weight_as_text(value: object) -> str:
    """Preserve Total-row weight text (often a leading-space ``' 38.00'``)."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, int):
        return f" {value:.2f}"
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        return f" {value:.2f}"
    text = str(value)
    return text if text.strip() else ""


def _read_sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    if not content:
        raise FbaBoxContentsTool2Error("Uploaded file is empty.")
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FbaBoxContentsTool2Error("File is too large (max 15 MB).")

    stripped = content.lstrip()
    if stripped.lower().startswith(b"<html") or stripped.lower().startswith(b"<!doctype"):
        raise FbaBoxContentsTool2Error(
            "This looks like an HTML file saved as Excel. Export the carton detail as .xls or .xlsx."
        )

    if content.startswith(b"PK"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise FbaBoxContentsTool2Error(
                "Could not read the Excel file. Upload a valid .xlsx workbook."
            ) from exc
        try:
            sheet = workbook[workbook.sheetnames[0]]
            return [tuple(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

    try:
        import xlrd
    except ImportError as exc:
        raise FbaBoxContentsTool2Error(
            "Reading .xls files requires xlrd. Upload .xlsx or install xlrd on the server."
        ) from exc

    try:
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows: list[tuple[object, ...]] = []
        for row_index in range(sheet.nrows):
            values: list[object] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    values.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    values.append(bool(cell.value))
                else:
                    values.append(cell.value)
            rows.append(tuple(values))
        return rows
    except FbaBoxContentsTool2Error:
        raise
    except Exception as exc:
        raise FbaBoxContentsTool2Error(
            "Could not read the Excel file. Upload a valid carton detail .xls or .xlsx."
        ) from exc


def _cell_at(row: Sequence[object], index: int) -> object:
    if index < 0 or index >= len(row):
        return None
    return row[index]


def _find_header_map(rows: Sequence[Sequence[object]]) -> tuple[int, dict[str, int]]:
    """Return (header_row_index, name→column) for Sku/UPC/Qty/Weight/dims."""
    wanted = {
        "sku": "sku",
        "upc": "upc",
        "qty": "qty",
        "quantity": "qty",
        "weight": "weight",
        "carton length": "length",
        "carton width": "width",
        "carton height": "height",
        "length": "length",
        "width": "width",
        "height": "height",
    }
    for index, row in enumerate(rows):
        mapping: dict[str, int] = {}
        for col, value in enumerate(row):
            key = wanted.get(_normalize_label(value))
            if key and key not in mapping:
                mapping[key] = col
        if "sku" in mapping and "upc" in mapping and "qty" in mapping:
            return index, mapping
    raise FbaBoxContentsTool2Error(
        'Could not find a header row with "Sku", "UPC", and "Qty" columns.'
    )


def _row_has_total_label(row: Sequence[object]) -> bool:
    for value in row:
        if _normalize_label(value) == "total":
            return True
    return False


def parse_carton_detail_tool2(content: bytes) -> ParsedCartonDetail:
    rows = _read_sheet_rows(content)
    header_index, cols = _find_header_map(rows)

    shipment_id = ""
    for row in rows[: max(header_index, 1)]:
        first = _normalize_label(_cell_at(row, 0)).replace(" ", "")
        second = _cell_text(_cell_at(row, 1))
        if first in {"po#:", "po#", "po:"} and second:
            shipment_id = second
            break
        candidate = _cell_text(_cell_at(row, 2))
        if candidate.upper().startswith("FBA"):
            shipment_id = candidate
            break

    sku_col = cols["sku"]
    upc_col = cols["upc"]
    qty_col = cols["qty"]
    weight_col = cols.get("weight", -1)
    length_col = cols.get("length", -1)
    width_col = cols.get("width", -1)
    height_col = cols.get("height", -1)

    cartons: list[Carton] = []
    current: Carton | None = None

    for row in rows[header_index + 1 :]:
        first = _cell_at(row, 0)
        if _is_carton_header(first):
            current = Carton(
                box_number=len(cartons) + 1,
                carton_id=_cell_text(_cell_at(row, 1)),
                length=_as_excel_number(_cell_at(row, length_col)) if length_col >= 0 else None,
                width=_as_excel_number(_cell_at(row, width_col)) if width_col >= 0 else None,
                height=_as_excel_number(_cell_at(row, height_col)) if height_col >= 0 else None,
            )
            cartons.append(current)
            continue

        if current is None:
            continue

        if _row_has_total_label(row):
            if weight_col >= 0:
                current.weight_text = _weight_as_text(_cell_at(row, weight_col))
            continue

        sku = _cell_text(_cell_at(row, sku_col))
        if not sku:
            continue
        upc_raw = _cell_at(row, upc_col)
        if not _looks_like_upc(upc_raw):
            continue
        upc = _as_upc_string(upc_raw)
        if not upc:
            continue
        qty = _as_qty(_cell_at(row, qty_col))
        current.rows.append(ContentRow(upc=upc, box_number=current.box_number, qty=qty))

    if not cartons:
        raise FbaBoxContentsTool2Error(
            'The file must include "Carton#:" rows under a Sku/UPC/Qty header.'
        )

    all_rows = tuple(row for carton in cartons for row in carton.rows)
    if not all_rows:
        raise FbaBoxContentsTool2Error("No item rows with a UPC were found under any carton.")

    return ParsedCartonDetail(
        shipment_id=shipment_id,
        cartons=tuple(cartons),
        rows=all_rows,
    )


def _set_text_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str | None,
    *,
    bold: bool = False,
    align_left: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align_left:
        cell.alignment = _LEFT
    if value is None or value == "":
        cell.value = None
    else:
        cell.value = str(value)
        cell.data_type = "s"


def _set_number_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | float | None,
    *,
    bold: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if value is None or isinstance(value, bool):
        cell.value = None
        return
    if isinstance(value, int):
        cell.value = int(value)
    elif isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            cell.value = int(value)
        else:
            cell.value = float(value)
    else:
        raise TypeError(f"Expected a number, got {type(value).__name__}")


def _write_box_contents(sheet: Worksheet, rows: Sequence[ContentRow]) -> int | float:
    sheet.sheet_view.showGridLines = True
    _set_text_cell(sheet, 1, 1, "UPC", bold=True)
    _set_text_cell(sheet, 1, 2, "Box Number", bold=True)
    _set_text_cell(sheet, 1, 3, "QTY", bold=True)

    total_qty: int | float = 0
    for index, row in enumerate(rows, start=2):
        _set_text_cell(sheet, index, 1, row.upc)
        _set_number_cell(sheet, index, 2, row.box_number)
        _set_number_cell(sheet, index, 3, row.qty)
        total_qty += row.qty

    boxes = sorted({row.box_number for row in rows})
    counts: dict[tuple[str, int], int | float] = {}
    upcs: set[str] = set()
    for row in rows:
        upcs.add(row.upc)
        key = (row.upc, row.box_number)
        counts[key] = counts.get(key, 0) + row.qty

    sorted_upcs = sorted(upcs)
    grand_total_col = PIVOT_START_COL + 1 + len(boxes)

    _set_text_cell(sheet, 1, PIVOT_START_COL, "Sum of QTY")
    _set_text_cell(sheet, 1, PIVOT_START_COL + 1, "Column Labels")
    _set_text_cell(sheet, 2, PIVOT_START_COL, "Row Labels")
    for offset, box in enumerate(boxes):
        _set_number_cell(sheet, 2, PIVOT_START_COL + 1 + offset, box)
    _set_text_cell(sheet, 2, grand_total_col, "Grand Total")

    column_totals: list[int | float] = [0] * len(boxes)
    current_row = 3
    for upc in sorted_upcs:
        _set_text_cell(sheet, current_row, PIVOT_START_COL, upc, align_left=True)
        row_total: int | float = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((upc, box), 0)
            if qty:
                _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
                row_total += qty
                column_totals[offset] += qty
        _set_number_cell(sheet, current_row, grand_total_col, row_total)
        current_row += 1

    _set_text_cell(sheet, current_row, PIVOT_START_COL, "Grand Total", align_left=True)
    for offset, qty in enumerate(column_totals):
        _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
    _set_number_cell(sheet, current_row, grand_total_col, total_qty)

    sheet.column_dimensions["A"].width = 13.11
    sheet.column_dimensions["B"].width = 13.33
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL)].width = 13.11
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL + 1)].width = 15.55
    for column in range(PIVOT_START_COL + 2, grand_total_col):
        sheet.column_dimensions[get_column_letter(column)].width = 3.0
    sheet.column_dimensions[get_column_letter(grand_total_col)].width = 10.78
    return total_qty


def _write_dimensions(sheet: Worksheet, cartons: Sequence[Carton]) -> None:
    sheet.sheet_view.showGridLines = True
    for column, header in enumerate(("Weight", "Length", "Width", "Height"), start=1):
        _set_text_cell(sheet, 1, column, header)

    for carton in cartons:
        row = carton.box_number + 1
        _set_text_cell(sheet, row, 1, carton.weight_text or None)
        _set_number_cell(sheet, row, 2, carton.length)
        _set_number_cell(sheet, row, 3, carton.width)
        _set_number_cell(sheet, row, 4, carton.height)

    sheet.column_dimensions["A"].width = 10.0
    sheet.column_dimensions["B"].width = 10.0
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions["D"].width = 10.0


def build_output_workbook(parsed: ParsedCartonDetail) -> bytes:
    workbook = Workbook()
    contents = workbook.active
    contents.title = BOX_CONTENTS_SHEET
    dimensions = workbook.create_sheet(DIMENSIONS_SHEET)
    _write_box_contents(contents, parsed.rows)
    _write_dimensions(dimensions, parsed.cartons)
    dimensions.sheet_view.tabSelected = False
    contents.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_fba_box_contents_tool2(
    content: bytes,
    filename: str | None = None,
) -> FbaBoxContentsTool2Result:
    parsed = parse_carton_detail_tool2(content)
    file_bytes = build_output_workbook(parsed)
    total_qty: int | float = 0
    for row in parsed.rows:
        total_qty += row.qty
    upc_count = len({row.upc for row in parsed.rows})
    output_name = sanitize_download_filename(filename, DEFAULT_OUTPUT_FILENAME)
    return FbaBoxContentsTool2Result(
        file_bytes=file_bytes,
        filename=output_name,
        row_count=len(parsed.rows),
        box_count=len(parsed.cartons),
        upc_count=upc_count,
        total_qty=total_qty,
        shipment_id=parsed.shipment_id,
    )
