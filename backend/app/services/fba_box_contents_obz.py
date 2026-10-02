"""FBA Box Contents OBZ Tool — Oboz packing slip by carton → Box Contents workbook.

Input (Oboz ``PackingSlipByCartonReport`` .xlsx / .xls):
  - A ``Carton`` row with the carton number
  - A header row ``PO`` / ``UPC/GTIN`` / ``Qty``
  - One item line per unit under that carton (GTIN stored as text, often with
    leading zeros)

Output (``Box Contents`` only, named ``{PO} Box Contents.xlsx``):
  - Columns A–C: Box Number (number), UPC (text, leading zeros removed), Qty
    (number), then a bold ``Total`` row whose Qty cell is ``=SUM(...)``
  - Pivot at column F: Sum of Qty by UPC (text) and carton number (number)
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass
from typing import Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_CENTER = Alignment(horizontal="center")
_LEFT = Alignment(horizontal="left")

BOX_CONTENTS_SHEET = "Box Contents"
DEFAULT_OUTPUT_FILENAME = "Box Contents.xlsx"
PIVOT_START_COL = 6  # column F
_LINE_START_ROW = 2


class FbaBoxContentsObzError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class ContentRow:
    upc: str
    box_number: int
    qty: int | float


@dataclass(frozen=True)
class ParsedPackingSlip:
    shipment_id: str
    rows: tuple[ContentRow, ...]


@dataclass(frozen=True)
class FbaBoxContentsObzResult:
    file_bytes: bytes
    filename: str
    row_count: int
    box_count: int
    upc_count: int
    total_qty: int | float
    shipment_id: str


def sanitize_download_filename(shipment_id: str | None) -> str:
    """Return ``{PO} Box Contents.xlsx``."""
    cleaned = (shipment_id or "").replace('"', "").replace("\r", "").replace("\n", "").strip()
    cleaned = re.sub(r'[<>:"/\\|?*]', "", cleaned).strip(" .")
    if not cleaned:
        return DEFAULT_OUTPUT_FILENAME
    return f"{cleaned} Box Contents.xlsx"


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return format(value, ".15g").strip()
    return str(value).strip()


def _normalize_label(value: object) -> str:
    return re.sub(r"\s+", " ", _cell_text(value)).strip().lower()


def _as_int(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer() and abs(value) < 2**53:
            return int(value)
        return None
    text = _cell_text(value)
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    return None


def _as_qty(value: object) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer() and abs(value) < 2**53:
            return int(value)
        return float(value)
    text = _cell_text(value).replace(",", "")
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    if re.fullmatch(r"-?\d+\.\d+", text):
        number = float(text)
        if number.is_integer() and abs(number) < 2**53:
            return int(number)
        return number
    return None


def normalize_upc(value: object) -> str:
    """GTIN/UPC as Excel text, with leading zeros removed (``00840…`` → ``840…``)."""
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, int):
        if value < 0:
            return ""
        text = str(value)
    elif isinstance(value, float):
        if not math.isfinite(value) or value < 0 or not value.is_integer() or abs(value) >= 2**53:
            return ""
        text = str(int(value))
    else:
        text = str(value).strip()
    if not text or text.lower() == "upc/gtin":
        return ""
    if re.fullmatch(r"\d+", text):
        text = text.lstrip("0") or "0"
    return text


def _labeled_columns(row: Sequence[object]) -> dict[int, str]:
    found: dict[int, str] = {}
    for index, value in enumerate(row):
        label = _normalize_label(value)
        if label:
            found[index] = label
    return found


def _read_sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    if not content:
        raise FbaBoxContentsObzError("Uploaded file is empty.")
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FbaBoxContentsObzError("File is too large (max 15 MB).")

    stripped = content.lstrip()
    if stripped.lower().startswith(b"<html") or stripped.lower().startswith(b"<!doctype"):
        raise FbaBoxContentsObzError(
            "This looks like an HTML file saved as Excel. Export the packing slip as .xls or .xlsx."
        )

    if content.startswith(b"PK"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise FbaBoxContentsObzError(
                "Could not read the Excel file. Upload a valid .xlsx workbook."
            ) from exc
        try:
            sheet = workbook[workbook.sheetnames[0]]
            return [tuple(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

    try:
        import xlrd
    except ImportError as exc:
        raise FbaBoxContentsObzError(
            "Reading .xls files requires xlrd. Upload .xlsx or install xlrd on the server."
        ) from exc

    try:
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows: list[tuple[object, ...]] = []
        for row_index in range(sheet.nrows):
            values: list[object] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    values.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    values.append(bool(cell.value))
                else:
                    values.append(cell.value)
            rows.append(tuple(values))
        return rows
    except FbaBoxContentsObzError:
        raise
    except Exception as exc:
        raise FbaBoxContentsObzError(
            "Could not read the Excel file. Upload a valid packing slip .xls or .xlsx."
        ) from exc


def _column_index(labels: dict[int, str], name: str) -> int | None:
    for index, label in labels.items():
        if label == name:
            return index
    return None


def parse_packing_slip(content: bytes) -> ParsedPackingSlip:
    """Read carton blocks from an Oboz packing slip by carton."""
    sheet_rows = _read_sheet_rows(content)
    if not any(any(cell is not None and cell != "" for cell in row) for row in sheet_rows):
        raise FbaBoxContentsObzError("The packing slip has no rows.")

    upc_col: int | None = None
    qty_col: int | None = None
    po_col: int | None = None
    current_box: int | None = None
    shipment_id = ""
    rows: list[ContentRow] = []

    for row in sheet_rows:
        labels = _labeled_columns(row)
        carton_col = _column_index(labels, "carton")
        if carton_col is not None:
            box_number = None
            for index in range(carton_col + 1, len(row)):
                box_number = _as_int(row[index])
                if box_number is not None:
                    break
            if box_number is None or box_number < 1:
                raise FbaBoxContentsObzError(
                    "A Carton row is missing its carton number."
                )
            current_box = box_number
            continue

        if "upc/gtin" in labels.values() and "qty" in labels.values():
            upc_col = _column_index(labels, "upc/gtin")
            qty_col = _column_index(labels, "qty")
            po_col = _column_index(labels, "po")
            continue

        if current_box is None or upc_col is None or qty_col is None:
            continue
        if upc_col >= len(row):
            continue

        upc = normalize_upc(row[upc_col])
        if not upc:
            continue
        qty_raw = row[qty_col] if qty_col < len(row) else None
        qty = _as_qty(qty_raw)
        if qty is None:
            raise FbaBoxContentsObzError(
                f"Item {upc} in carton {current_box} is missing a numeric Qty."
            )
        if qty < 0:
            raise FbaBoxContentsObzError(
                f"Item {upc} in carton {current_box} has a negative Qty."
            )
        if not shipment_id and po_col is not None and po_col < len(row):
            shipment_id = _cell_text(row[po_col])
        rows.append(ContentRow(upc=upc, box_number=current_box, qty=qty))

    if current_box is None:
        raise FbaBoxContentsObzError(
            "No Carton rows were found. Upload an Oboz Packing Slip by Carton."
        )
    if not rows:
        raise FbaBoxContentsObzError("No item rows with a UPC/GTIN were found under any carton.")

    return ParsedPackingSlip(shipment_id=shipment_id, rows=tuple(rows))


def _set_text_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str,
    *,
    bold: bool = False,
    align_left: bool = False,
    align_center: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align_center:
        cell.alignment = _CENTER
    elif align_left:
        cell.alignment = _LEFT
    cell.value = value
    cell.data_type = "s"


def _set_number_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | float,
    *,
    bold: bool = False,
    align_center: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align_center:
        cell.alignment = _CENTER
    if isinstance(value, bool):
        raise TypeError("Qty and box numbers must be numbers.")
    if isinstance(value, int):
        cell.value = int(value)
    elif isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            cell.value = int(value)
        else:
            cell.value = float(value)
    else:
        raise TypeError(f"Expected a number, got {type(value).__name__}")


def _set_formula_cell(sheet: Worksheet, row: int, column: int, formula: str) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD
    cell.alignment = _CENTER
    cell.value = formula


def _write_box_contents(sheet: Worksheet, rows: Sequence[ContentRow]) -> int | float:
    sheet.sheet_view.showGridLines = True
    _set_text_cell(sheet, 1, 1, "Box Number", bold=True, align_center=True)
    _set_text_cell(sheet, 1, 2, "UPC", bold=True, align_center=True)
    _set_text_cell(sheet, 1, 3, "Qty", bold=True, align_center=True)

    total_qty: int | float = 0
    for offset, row in enumerate(rows):
        excel_row = _LINE_START_ROW + offset
        _set_number_cell(sheet, excel_row, 1, row.box_number, align_center=True)
        _set_text_cell(sheet, excel_row, 2, row.upc, align_center=True)
        _set_number_cell(sheet, excel_row, 3, row.qty, align_center=True)
        total_qty += row.qty

    last_line = _LINE_START_ROW + len(rows) - 1
    total_row = last_line + 1
    _set_text_cell(sheet, total_row, 2, "Total", bold=True, align_center=True)
    _set_formula_cell(sheet, total_row, 3, f"=SUM(C{_LINE_START_ROW}:C{last_line})")

    boxes = sorted({row.box_number for row in rows})
    counts: dict[tuple[str, int], int | float] = {}
    upcs: set[str] = set()
    for row in rows:
        upcs.add(row.upc)
        key = (row.upc, row.box_number)
        counts[key] = counts.get(key, 0) + row.qty

    sorted_upcs = sorted(upcs)
    grand_total_col = PIVOT_START_COL + 1 + len(boxes)

    _set_text_cell(sheet, 1, PIVOT_START_COL, "Sum of Qty")
    _set_text_cell(sheet, 1, PIVOT_START_COL + 1, "Column Labels")
    _set_text_cell(sheet, 2, PIVOT_START_COL, "Row Labels")
    for offset, box in enumerate(boxes):
        _set_number_cell(sheet, 2, PIVOT_START_COL + 1 + offset, box)
    _set_text_cell(sheet, 2, grand_total_col, "Grand Total")

    column_totals: list[int | float] = [0] * len(boxes)
    current_row = 3
    for upc in sorted_upcs:
        _set_text_cell(sheet, current_row, PIVOT_START_COL, upc, align_left=True)
        row_total: int | float = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((upc, box), 0)
            if qty:
                _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
                row_total += qty
                column_totals[offset] += qty
        _set_number_cell(sheet, current_row, grand_total_col, row_total)
        current_row += 1

    _set_text_cell(sheet, current_row, PIVOT_START_COL, "Grand Total", align_left=True)
    for offset, qty in enumerate(column_totals):
        _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
    _set_number_cell(sheet, current_row, grand_total_col, total_qty)

    sheet.column_dimensions["A"].width = 11.11
    sheet.column_dimensions["B"].width = 15.11
    sheet.column_dimensions["C"].width = 8.89
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL)].width = 13.11
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL + 1)].width = 15.55
    for column in range(PIVOT_START_COL + 2, grand_total_col):
        sheet.column_dimensions[get_column_letter(column)].width = 2.0
    sheet.column_dimensions[get_column_letter(grand_total_col)].width = 10.78
    return total_qty


def build_output_workbook(parsed: ParsedPackingSlip) -> bytes:
    workbook = Workbook()
    contents = workbook.active
    contents.title = BOX_CONTENTS_SHEET
    _write_box_contents(contents, parsed.rows)
    contents.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_fba_box_contents_obz(
    content: bytes,
    filename: str | None = None,
) -> FbaBoxContentsObzResult:
    _ = filename
    parsed = parse_packing_slip(content)
    file_bytes = build_output_workbook(parsed)
    total_qty: int | float = 0
    for row in parsed.rows:
        total_qty += row.qty
    upc_count = len({row.upc for row in parsed.rows})
    box_count = len({row.box_number for row in parsed.rows})
    return FbaBoxContentsObzResult(
        file_bytes=file_bytes,
        filename=sanitize_download_filename(parsed.shipment_id),
        row_count=len(parsed.rows),
        box_count=box_count,
        upc_count=upc_count,
        total_qty=total_qty,
        shipment_id=parsed.shipment_id,
    )
