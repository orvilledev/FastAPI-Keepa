"""Convert an FBA Carton Detail report into a Box Contents + Dimensions workbook.

Input (Amazon FBA carton dump, .xls or .xlsx): repeating blocks of
``Carton#:`` + item lines (SKU / UPC / Description / Qty / Weight) + ``Total``.

Output:
  - ``Box Contents``: UPC (text), Box # (number), QTY (number), plus a Sum-of-QTY
    pivot by UPC and box starting at column G
  - ``Dimensions``: Box #, Weight (item weights rounded up to 1 decimal), Length,
    Width, Height — all true Excel numbers
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_UP
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_HEADER_FONT = Font(name="Calibri", size=11, bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E0B0FF")
_LEFT = Alignment(horizontal="left", vertical="center")
_CENTER = Alignment(horizontal="center", vertical="center")
_TEXT_FORMAT = "@"
_GENERAL = "General"

BOX_CONTENTS_SHEET = "Box Contents"
DIMENSIONS_SHEET = "Dimensions"
DEFAULT_OUTPUT_FILENAME = "FBA Box Contents Output.xlsx"

PIVOT_START_COL = 7  # column G — three blank columns after UPC / Box # / QTY

_SKIP_LABELS = frozenset(
    {
        "sku",
        "size",
        "upc",
        "description",
        "qty",
        "weight",
        "carton#:",
        "carton#",
        "fba carton detail",
    }
)


class FbaBoxContentsError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class ContentRow:
    upc: str
    box_number: int
    qty: int | float


@dataclass
class Carton:
    box_number: int
    carton_id: str
    length: int | float | None
    width: int | float | None
    height: int | float | None
    item_weights: list[float] = field(default_factory=list)
    rows: list[ContentRow] = field(default_factory=list)

    @property
    def weight(self) -> int | float:
        return _round_up_one_decimal(sum(self.item_weights))


@dataclass(frozen=True)
class ParsedCartonDetail:
    shipment_id: str
    cartons: tuple[Carton, ...]
    rows: tuple[ContentRow, ...]


@dataclass(frozen=True)
class FbaBoxContentsResult:
    file_bytes: bytes
    filename: str
    row_count: int
    box_count: int
    upc_count: int
    total_qty: int | float
    shipment_id: str


def sanitize_download_filename(name: str | None, fallback: str = DEFAULT_OUTPUT_FILENAME) -> str:
    """Return ``{stem} Output.xlsx`` from the uploaded carton-detail filename."""
    cleaned = (name or "").replace('"', "").replace("\r", "").replace("\n", "").strip()
    base = Path(cleaned).name if cleaned else ""
    stem = Path(base).stem if base else ""
    if stem.lower().endswith(" output"):
        stem = stem[: -len(" output")].rstrip()
    stem = re.sub(r'[<>:"/\\|?*]', "", stem).strip(" .")
    if not stem:
        return fallback
    return f"{stem} Output.xlsx"


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return format(value, ".15g").strip()
    return str(value).strip()


def _as_upc_string(value: object) -> str:
    """Keep UPC as a digit string so Excel will not coerce it to a number."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        if value < 0:
            return ""
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value) or value < 0:
            return ""
        if abs(value) < 2**53:
            return str(int(value)) if value.is_integer() else str(int(round(value)))
        return format(value, ".0f")
    text = str(value).strip()
    if not text:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _looks_like_upc(value: object) -> bool:
    text = _as_upc_string(value).replace(" ", "")
    return bool(text) and text.isdigit() and 8 <= len(text) <= 14


def _as_excel_number(value: object) -> int | float | None:
    """Return a true int/float, or None when the cell is empty/non-numeric."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer() and abs(value) < 2**53:
            return int(value)
        return value
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    if number.is_integer() and abs(number) < 2**53:
        return int(number)
    return number


def _as_qty(value: object) -> int | float:
    parsed = _as_excel_number(value)
    return 0 if parsed is None else parsed


def _as_weight(value: object) -> float | None:
    parsed = _as_excel_number(value)
    if parsed is None:
        return None
    return float(parsed)


def _round_up_one_decimal(value: float) -> int | float:
    """Excel ROUNDUP(value, 1) — ceiling away from zero to one decimal place."""
    if value == 0:
        return 0
    quantized = Decimal(str(value)).quantize(Decimal("0.1"), rounding=ROUND_UP)
    as_float = float(quantized)
    if quantized == quantized.to_integral_value() and abs(as_float) < 2**53:
        return int(quantized)
    return as_float


def _normalize_label(value: object) -> str:
    return re.sub(r"\s+", " ", _cell_text(value)).strip().lower()


def _is_carton_header(value: object) -> bool:
    label = _normalize_label(value).replace(" ", "")
    return label in {"carton#:", "carton#", "carton:"}


def _is_skip_row(value: object) -> bool:
    label = _normalize_label(value)
    if not label:
        return True
    if label in _SKIP_LABELS:
        return True
    if label.startswith("total"):
        return True
    if label.startswith("page"):
        return True
    if label.startswith("fba carton"):
        return True
    return False


def _read_sheet_rows(content: bytes) -> list[tuple[object, ...]]:
    if not content:
        raise FbaBoxContentsError("Uploaded file is empty.")
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FbaBoxContentsError("File is too large (max 15 MB).")

    stripped = content.lstrip()
    if stripped.lower().startswith(b"<html") or stripped.lower().startswith(b"<!doctype"):
        raise FbaBoxContentsError(
            "This looks like an HTML file saved as Excel. Export the FBA Carton Detail as .xls or .xlsx."
        )

    if content.startswith(b"PK"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise FbaBoxContentsError(
                "Could not read the Excel file. Upload a valid .xlsx workbook."
            ) from exc
        try:
            sheet = workbook[workbook.sheetnames[0]]
            return [tuple(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

    try:
        import xlrd
    except ImportError as exc:
        raise FbaBoxContentsError(
            "Reading .xls files requires xlrd. Upload .xlsx or install xlrd on the server."
        ) from exc

    try:
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows: list[tuple[object, ...]] = []
        for row_index in range(sheet.nrows):
            values: list[object] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    values.append(None)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    values.append(bool(cell.value))
                else:
                    values.append(cell.value)
            rows.append(tuple(values))
        return rows
    except FbaBoxContentsError:
        raise
    except Exception as exc:
        raise FbaBoxContentsError(
            "Could not read the Excel file. Upload a valid FBA Carton Detail .xls or .xlsx."
        ) from exc


def _cell_at(row: Sequence[object], index: int) -> object:
    if index < 0 or index >= len(row):
        return None
    return row[index]


def _parse_item_fields(row: Sequence[object]) -> tuple[str, int | float, float | None] | None:
    """Return (upc, qty, weight) for an item line, or None if it is not an item."""
    sku = _cell_at(row, 0)
    if _is_skip_row(sku) or _is_carton_header(sku):
        return None
    if not _cell_text(sku):
        return None

    col_b = _cell_at(row, 1)
    col_c = _cell_at(row, 2)
    if _looks_like_upc(col_b):
        upc = _as_upc_string(col_b)
        qty = _as_qty(_cell_at(row, 3))
        weight = _as_weight(_cell_at(row, 4))
        return upc, qty, weight
    if _looks_like_upc(col_c):
        upc = _as_upc_string(col_c)
        qty = _as_qty(_cell_at(row, 4))
        weight = _as_weight(_cell_at(row, 5))
        return upc, qty, weight
    return None


def parse_carton_detail(content: bytes) -> ParsedCartonDetail:
    rows = _read_sheet_rows(content)
    shipment_id = ""
    if rows:
        title = _normalize_label(_cell_at(rows[0], 0))
        candidate = _cell_text(_cell_at(rows[0], 2))
        if title.startswith("fba carton") and candidate:
            shipment_id = candidate
        elif candidate.upper().startswith("FBA"):
            shipment_id = candidate

    cartons: list[Carton] = []
    current: Carton | None = None

    for row in rows:
        first = _cell_at(row, 0)
        if _is_carton_header(first):
            current = Carton(
                box_number=len(cartons) + 1,
                carton_id=_cell_text(_cell_at(row, 1)),
                length=_as_excel_number(_cell_at(row, 2)),
                width=_as_excel_number(_cell_at(row, 3)),
                height=_as_excel_number(_cell_at(row, 4)),
            )
            cartons.append(current)
            continue

        if current is None:
            continue
        if _is_skip_row(first):
            continue

        parsed = _parse_item_fields(row)
        if parsed is None:
            continue
        upc, qty, weight = parsed
        if not upc:
            continue
        current.rows.append(ContentRow(upc=upc, box_number=current.box_number, qty=qty))
        if weight is not None:
            current.item_weights.append(weight)

    if not cartons:
        raise FbaBoxContentsError(
            'The file must be an FBA Carton Detail report with "Carton#:" rows.'
        )

    all_rows = tuple(row for carton in cartons for row in carton.rows)
    if not all_rows:
        raise FbaBoxContentsError("No item rows with a UPC were found under any carton.")

    return ParsedCartonDetail(
        shipment_id=shipment_id,
        cartons=tuple(cartons),
        rows=all_rows,
    )


def _style_header_cell(cell) -> None:
    """Bold mauve header look — does not change value or number format."""
    cell.font = _HEADER_FONT
    cell.fill = _HEADER_FILL
    cell.alignment = _CENTER


def _set_text_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str | None,
    *,
    bold: bool = False,
    header: bool = False,
    align_left: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = _TEXT_FORMAT
    if align_left:
        cell.alignment = _LEFT
    if value is None or value == "":
        cell.value = None
    else:
        cell.value = str(value)
        cell.data_type = "s"
    if header:
        _style_header_cell(cell)
    else:
        cell.font = _CALIBRI_BOLD if bold else _CALIBRI


def _set_number_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | float | None,
    *,
    bold: bool = False,
    header: bool = False,
    align_left: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = _GENERAL
    if align_left:
        cell.alignment = _LEFT
    if value is None or isinstance(value, bool):
        cell.value = None
    elif isinstance(value, int):
        cell.value = int(value)
    elif isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            cell.value = int(value)
        else:
            cell.value = float(value)
    else:
        raise TypeError(f"Expected a number, got {type(value).__name__}")
    if header:
        _style_header_cell(cell)
    else:
        cell.font = _CALIBRI_BOLD if bold else _CALIBRI


def _write_box_contents(sheet: Worksheet, rows: Sequence[ContentRow]) -> int | float:
    sheet.sheet_view.showGridLines = True
    _set_text_cell(sheet, 1, 1, "UPC", header=True)
    _set_text_cell(sheet, 1, 2, "Box #", header=True)
    _set_text_cell(sheet, 1, 3, "QTY", header=True)

    total_qty: int | float = 0
    for index, row in enumerate(rows, start=2):
        _set_text_cell(sheet, index, 1, row.upc)
        _set_number_cell(sheet, index, 2, row.box_number)
        _set_number_cell(sheet, index, 3, row.qty)
        total_qty += row.qty

    last_data_row = 1 + len(rows)
    total_row = last_data_row + 2
    _set_text_cell(sheet, total_row, 2, "Total QTY", bold=True)
    _set_number_cell(sheet, total_row, 3, total_qty, bold=True)

    boxes = sorted({row.box_number for row in rows})
    counts: dict[tuple[str, int], int | float] = {}
    upcs: set[str] = set()
    for row in rows:
        upcs.add(row.upc)
        key = (row.upc, row.box_number)
        counts[key] = counts.get(key, 0) + row.qty

    sorted_upcs = sorted(upcs)
    grand_total_col = PIVOT_START_COL + 1 + len(boxes)

    _set_text_cell(sheet, 1, PIVOT_START_COL, "Sum of QTY", header=True)
    _set_text_cell(sheet, 1, PIVOT_START_COL + 1, "Column Labels", header=True)
    _set_text_cell(sheet, 2, PIVOT_START_COL, "Row Labels", header=True)
    for offset, box in enumerate(boxes):
        _set_number_cell(sheet, 2, PIVOT_START_COL + 1 + offset, box, header=True)
    _set_text_cell(sheet, 2, grand_total_col, "Grand Total", header=True)

    column_totals: list[int | float] = [0] * len(boxes)
    current_row = 3
    for upc in sorted_upcs:
        _set_text_cell(sheet, current_row, PIVOT_START_COL, upc)
        row_total: int | float = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((upc, box), 0)
            if qty:
                _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
                row_total += qty
                column_totals[offset] += qty
        _set_number_cell(sheet, current_row, grand_total_col, row_total)
        current_row += 1

    _set_text_cell(sheet, current_row, PIVOT_START_COL, "Grand Total")
    for offset, qty in enumerate(column_totals):
        _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
    _set_number_cell(sheet, current_row, grand_total_col, total_qty)

    sheet.column_dimensions["A"].width = 13.11
    sheet.column_dimensions["B"].width = 9.0
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL)].width = 13.11
    for column in range(PIVOT_START_COL + 1, grand_total_col):
        sheet.column_dimensions[get_column_letter(column)].width = 4.0
    sheet.column_dimensions[get_column_letter(grand_total_col)].width = 12.0
    return total_qty


def _write_dimensions(sheet: Worksheet, cartons: Sequence[Carton]) -> None:
    sheet.sheet_view.showGridLines = True
    headers = ("Box #", "Weight", "Length", "Width", "Height")
    for column, header in enumerate(headers, start=1):
        _set_text_cell(sheet, 1, column, header, header=True)

    for carton in cartons:
        row = carton.box_number + 1
        _set_number_cell(sheet, row, 1, carton.box_number, align_left=True)
        _set_number_cell(sheet, row, 2, carton.weight)
        _set_number_cell(sheet, row, 3, carton.length)
        _set_number_cell(sheet, row, 4, carton.width)
        _set_number_cell(sheet, row, 5, carton.height)

    sheet.column_dimensions["A"].width = 8.89
    sheet.column_dimensions["B"].width = 8.0
    sheet.column_dimensions["C"].width = 8.0
    sheet.column_dimensions["D"].width = 8.0
    sheet.column_dimensions["E"].width = 8.0


def build_output_workbook(parsed: ParsedCartonDetail) -> bytes:
    workbook = Workbook()
    contents = workbook.active
    contents.title = BOX_CONTENTS_SHEET
    dimensions = workbook.create_sheet(DIMENSIONS_SHEET)
    _write_box_contents(contents, parsed.rows)
    _write_dimensions(dimensions, parsed.cartons)
    dimensions.sheet_view.tabSelected = False
    contents.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_fba_box_contents(
    content: bytes,
    filename: str | None = None,
) -> FbaBoxContentsResult:
    parsed = parse_carton_detail(content)
    file_bytes = build_output_workbook(parsed)
    total_qty: int | float = 0
    for row in parsed.rows:
        total_qty += row.qty
    upc_count = len({row.upc for row in parsed.rows})
    output_name = sanitize_download_filename(filename, DEFAULT_OUTPUT_FILENAME)
    return FbaBoxContentsResult(
        file_bytes=file_bytes,
        filename=output_name,
        row_count=len(parsed.rows),
        box_count=len(parsed.cartons),
        upc_count=upc_count,
        total_qty=total_qty,
        shipment_id=parsed.shipment_id,
    )
