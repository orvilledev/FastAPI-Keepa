"""NFA and SMW Output vs Amazon manifest → ``{shipment} - Corrected.xlsx``.

Takes the NFA and SMW Tool Output (``Box Contents`` + ``Dimensions``) and an
Amazon Seller Central manifest (``Box packing information``). The manifest is
the source of truth: Old SKUs are remapped via the catalog, units the manifest
does not expect are removed, and every shortfall is appended as a new row in
the last box.

The output mirrors the manual ``… - Corrected.xlsx`` layout:
  - Box Contents: UPC, Box Number, QTY with the Output rows in their original
    order; remapped and appended rows highlighted yellow
  - an E–G summary of the appended units (shipment id, Box #, Qty)
  - a Sum-of-QTY pivot at column I (column G when nothing was appended)
  - Dimensions in the Box Number layout (Box Number, Weight, Length, Width,
    Height). Zero values are given a suggestion + a ``Zero Correction`` reason
    (suggestions already made in the Output are carried over unchanged)
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.fba_nfa_smw_dimensions import (
    DimensionLine,
    apply_zero_corrections,
    build_items_by_box,
    parse_dimension_lines,
    write_dimension_lines,
)
from app.services.fba_upload_compare import (
    AdjustmentRow,
    ContentRow,
    FbaUploadCompareError,
    ParsedAmz,
    ParsedOutput,
    apply_old_sku_remaps,
    build_upc_to_old_sku_map,
    parse_amz_upload,
    parse_output_workbook,
    reconcile_to_amz,
)

_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_CENTER = Alignment(horizontal="center")
_LEFT = Alignment(horizontal="left")
_YELLOW = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")

BOX_CONTENTS_SHEET = "Box Contents"
DIMENSIONS_SHEET = "Dimensions"
DEFAULT_OUTPUT_FILENAME = "Box Contents - Corrected.xlsx"

PIVOT_START_COL_DEFAULT = 7  # G — same as the Output
PIVOT_START_COL_WITH_SUMMARY = 9  # I — room for the E–G summary
_SUMMARY_COL = 5  # E


class FbaBoxContentsTool2CompareError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class FbaBoxContentsTool2CompareResult:
    file_bytes: bytes
    filename: str
    row_count: int
    box_count: int
    upc_count: int
    total_qty: int | float
    remapped_count: int
    added_count: int
    added_qty: int | float
    removed_count: int
    removed_qty: int | float
    last_box: int
    shipment_id: str
    zero_corrected_count: int = 0


def sanitize_download_filename(shipment_id: str | None) -> str:
    """Return ``{shipment_id} - Corrected.xlsx``."""
    cleaned = (shipment_id or "").replace('"', "").replace("\r", "").replace("\n", "").strip()
    cleaned = re.sub(r'[<>:"/\\|?*]', "", cleaned).strip(" .")
    if not cleaned:
        return DEFAULT_OUTPUT_FILENAME
    return f"{cleaned} - Corrected.xlsx"


def _shipment_from_filename(filename: str | None) -> str:
    if not filename:
        return ""
    stem = Path(filename).stem
    for suffix in (" - Corrected", " Output", " Box Contents"):
        if stem.lower().endswith(suffix.lower()):
            stem = stem[: -len(suffix)].rstrip()
            break
    return stem.strip()


def _reword(exc: Exception, fallback: str) -> str:
    message = str(exc) or fallback
    return (
        message.replace("AMZ upload", "manifest")
        .replace("AMZ ", "manifest ")
        .replace('"AMZ', '"manifest')
    )


def parse_tool2_output(content: bytes) -> ParsedOutput:
    """Parse the NFA and SMW Tool Output (Box Contents + Dimensions)."""
    try:
        return parse_output_workbook(content)
    except FbaUploadCompareError as exc:
        raise FbaBoxContentsTool2CompareError(str(exc)) from exc


def parse_manifest(content: bytes) -> ParsedAmz:
    """Parse the Amazon Seller Central manifest (Box packing information)."""
    try:
        return parse_amz_upload(content)
    except FbaUploadCompareError as exc:
        raise FbaBoxContentsTool2CompareError(
            _reword(exc, "Could not read the manifest Excel file.")
        ) from exc


def _set_text_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str | None,
    *,
    bold: bool = False,
    align: Alignment | None = None,
    fill: PatternFill | None = None,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align is not None:
        cell.alignment = align
    if fill is not None:
        cell.fill = fill
    if value is None or value == "":
        cell.value = None
        return
    cell.value = str(value)
    cell.data_type = "s"


def _set_number_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | float | None,
    *,
    align: Alignment | None = None,
    fill: PatternFill | None = None,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI
    if align is not None:
        cell.alignment = align
    if fill is not None:
        cell.fill = fill
    if value is None or isinstance(value, bool):
        cell.value = None
        return
    if isinstance(value, float) and value.is_integer() and abs(value) < 2**53:
        cell.value = int(value)
    else:
        cell.value = value


def _write_added_summary(
    sheet: Worksheet,
    shipment_id: str,
    added: Sequence[AdjustmentRow],
) -> None:
    _set_text_cell(sheet, 1, _SUMMARY_COL, shipment_id, bold=True, align=_CENTER)
    _set_text_cell(sheet, 1, _SUMMARY_COL + 1, "Box #", bold=True, align=_CENTER)
    _set_text_cell(sheet, 1, _SUMMARY_COL + 2, "Qty", bold=True, align=_CENTER)
    for excel_row, item in enumerate(added, start=2):
        # Text so 12-digit UPCs never render in scientific notation.
        _set_text_cell(sheet, excel_row, _SUMMARY_COL, item.identifier, align=_CENTER, fill=_YELLOW)
        _set_number_cell(sheet, excel_row, _SUMMARY_COL + 1, item.box_number, align=_CENTER, fill=_YELLOW)
        _set_number_cell(sheet, excel_row, _SUMMARY_COL + 2, item.qty, align=_CENTER, fill=_YELLOW)

    sheet.column_dimensions["E"].width = 13.89
    sheet.column_dimensions["F"].width = 11.44


def _write_box_contents(
    sheet: Worksheet,
    rows: Sequence[ContentRow],
    *,
    shipment_id: str,
    added: Sequence[AdjustmentRow],
    box_numbers: Sequence[int],
) -> int | float:
    sheet.sheet_view.showGridLines = True
    _set_text_cell(sheet, 1, 1, "UPC", bold=True)
    _set_text_cell(sheet, 1, 2, "Box Number", bold=True)
    _set_text_cell(sheet, 1, 3, "QTY", bold=True)

    total_qty: int | float = 0
    for excel_row, row in enumerate(rows, start=2):
        fill = _YELLOW if (row.remapped or row.added) else None
        _set_text_cell(sheet, excel_row, 1, row.identifier, fill=fill)
        _set_number_cell(sheet, excel_row, 2, row.box_number, fill=fill)
        _set_number_cell(sheet, excel_row, 3, row.qty, fill=fill)
        total_qty += row.qty

    if added:
        _write_added_summary(sheet, shipment_id, added)
    pivot_start = PIVOT_START_COL_WITH_SUMMARY if added else PIVOT_START_COL_DEFAULT

    boxes = sorted({row.box_number for row in rows} | set(box_numbers))
    counts: dict[tuple[str, int], int | float] = {}
    for row in rows:
        key = (row.identifier, row.box_number)
        counts[key] = counts.get(key, 0) + row.qty
    identifiers = sorted({row.identifier for row in rows})
    grand_total_col = pivot_start + 1 + len(boxes)

    _set_text_cell(sheet, 1, pivot_start, "Sum of QTY")
    _set_text_cell(sheet, 1, pivot_start + 1, "Column Labels")
    _set_text_cell(sheet, 2, pivot_start, "Row Labels")
    for offset, box in enumerate(boxes):
        _set_number_cell(sheet, 2, pivot_start + 1 + offset, box)
    _set_text_cell(sheet, 2, grand_total_col, "Grand Total")

    column_totals: list[int | float] = [0] * len(boxes)
    current_row = 3
    for ident in identifiers:
        _set_text_cell(sheet, current_row, pivot_start, ident, align=_LEFT)
        row_total: int | float = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((ident, box), 0)
            if qty:
                _set_number_cell(sheet, current_row, pivot_start + 1 + offset, qty)
                row_total += qty
                column_totals[offset] += qty
        _set_number_cell(sheet, current_row, grand_total_col, row_total)
        current_row += 1

    _set_text_cell(sheet, current_row, pivot_start, "Grand Total", align=_LEFT)
    for offset, qty in enumerate(column_totals):
        _set_number_cell(sheet, current_row, pivot_start + 1 + offset, qty)
    _set_number_cell(sheet, current_row, grand_total_col, total_qty)

    sheet.column_dimensions["A"].width = 13.11
    sheet.column_dimensions["B"].width = 13.33
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions[get_column_letter(pivot_start)].width = 13.11
    sheet.column_dimensions[get_column_letter(pivot_start + 1)].width = 15.55
    for column in range(pivot_start + 2, grand_total_col):
        sheet.column_dimensions[get_column_letter(column)].width = 3.0
    sheet.column_dimensions[get_column_letter(grand_total_col)].width = 10.78
    return total_qty


def read_dimension_lines(content: bytes) -> list[DimensionLine]:
    """Dimensions rows of an NFA and SMW Output (new Box Number layout or the old one)."""
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise FbaBoxContentsTool2CompareError("Could not read the Output Excel file.") from exc
    try:
        if DIMENSIONS_SHEET not in workbook.sheetnames:
            return []
        return parse_dimension_lines(workbook[DIMENSIONS_SHEET])
    finally:
        workbook.close()


def build_corrected_workbook(
    rows: Sequence[ContentRow],
    dimensions: Sequence[DimensionLine],
    *,
    shipment_id: str,
    added: Sequence[AdjustmentRow],
    box_numbers: Sequence[int],
) -> bytes:
    workbook = Workbook()
    contents = workbook.active
    contents.title = BOX_CONTENTS_SHEET
    dimensions_sheet = workbook.create_sheet(DIMENSIONS_SHEET)
    _write_box_contents(
        contents,
        rows,
        shipment_id=shipment_id,
        added=added,
        box_numbers=box_numbers,
    )
    write_dimension_lines(dimensions_sheet, dimensions)
    dimensions_sheet.sheet_view.tabSelected = False
    contents.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_tool2_manifest_compare(
    output_content: bytes,
    manifest_content: bytes,
    catalog_rows: Sequence[Mapping[str, object]],
    *,
    output_filename: str | None = None,
) -> FbaBoxContentsTool2CompareResult:
    """Compare the NFA and SMW Output to the manifest and return the corrected workbook."""
    parsed_output = parse_tool2_output(output_content)
    parsed_manifest = parse_manifest(manifest_content)
    shipment_id = parsed_manifest.shipment_id or _shipment_from_filename(output_filename)

    amz_ids = {row.sku_id for row in parsed_manifest.skus}
    upc_to_old = build_upc_to_old_sku_map(catalog_rows, amz_ids=amz_ids)
    remapped_rows = apply_old_sku_remaps(parsed_output, parsed_manifest, upc_to_old)
    try:
        reconciled = reconcile_to_amz(
            remapped_rows, parsed_manifest, upc_to_old, preserve_rows=True
        )
    except FbaUploadCompareError as exc:
        raise FbaBoxContentsTool2CompareError(str(exc)) from exc

    # Zero weights/dimensions are judged against the cartons as they were packed
    # (the Output rows), not the manifest-adjusted rows.
    dimension_lines, zero_corrected = apply_zero_corrections(
        read_dimension_lines(output_content),
        build_items_by_box(
            (row.identifier, row.box_number, row.qty) for row in parsed_output.rows
        ),
    )

    file_bytes = build_corrected_workbook(
        reconciled.rows,
        dimension_lines,
        shipment_id=shipment_id,
        added=reconciled.added,
        box_numbers=reconciled.box_numbers,
    )

    return FbaBoxContentsTool2CompareResult(
        file_bytes=file_bytes,
        filename=sanitize_download_filename(shipment_id),
        row_count=len(reconciled.rows),
        box_count=len(reconciled.box_numbers),
        upc_count=len({row.identifier for row in reconciled.rows}),
        total_qty=sum(row.qty for row in reconciled.rows),
        remapped_count=sum(1 for row in reconciled.rows if row.remapped),
        added_count=len(reconciled.added),
        added_qty=sum(item.qty for item in reconciled.added),
        removed_count=len(reconciled.removed),
        removed_qty=sum(item.qty for item in reconciled.removed),
        last_box=reconciled.last_box,
        shipment_id=shipment_id,
        zero_corrected_count=zero_corrected,
    )
