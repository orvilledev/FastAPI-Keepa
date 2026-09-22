"""Parse a bulk SKU list and build an existence-check workbook."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_MAX_SKUS = 25_000
_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)

_HEADER_ALIASES = {
    "sku": "sku",
    "skus": "sku",
    "merchant sku": "sku",
    "merchant_sku": "sku",
    "product sku": "sku",
}


class WarehouseSkuCheckError(ValueError):
    """User-correctable SKU check input problems."""


@dataclass(frozen=True)
class SkuCheckRow:
    sku: str
    exists: bool
    upc: str = ""
    fnsku: str = ""
    style_name: str = ""
    condition: str = ""
    match_count: int = 0


@dataclass(frozen=True)
class SkuCheckResult:
    file_bytes: bytes
    filename: str
    total: int
    found: int
    missing: int
    unique_skus: int


def _cell_str(raw: Any) -> str:
    if raw is None or isinstance(raw, bool):
        return ""
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if raw.is_integer() and abs(raw) < 2**53:
            return str(int(raw))
        return format(raw, ".15g").strip()
    return str(raw).strip()


def _normalize_header(cell: Any) -> str | None:
    if cell is None:
        return None
    key = re.sub(r"\s+", " ", str(cell).strip().lower())
    return _HEADER_ALIASES.get(key)


def _split_tokens(text: str) -> list[str]:
    parts = re.split(r"[\n\r,;\t]+", text)
    return [part.strip() for part in parts if part.strip()]


def _is_sku_header(value: str) -> bool:
    return _normalize_header(value) == "sku"


def _skus_from_table(rows: Sequence[Sequence[Any]]) -> list[str]:
    if not rows:
        return []
    header = rows[0]
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header):
        normalized = _normalize_header(cell)
        if normalized and normalized not in mapping:
            mapping[normalized] = idx

    if "sku" in mapping:
        col = mapping["sku"]
        values: list[str] = []
        for row in rows[1:]:
            if col < len(row):
                text = _cell_str(row[col])
                if text:
                    values.append(text)
        return values

    # No SKU header — take the first non-empty cell per row (including row 0).
    values: list[str] = []
    for row in rows:
        for cell in row:
            text = _cell_str(cell)
            if text:
                values.append(text)
                break
    return values


def _parse_text(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig", errors="replace")
    return _split_tokens(text)


def _parse_csv_bytes(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [tuple(row) for row in reader]
    return _skus_from_table(rows)


def _parse_excel(content: bytes) -> list[str]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise WarehouseSkuCheckError(
            "Could not read the Excel file. Upload a valid .xlsx."
        ) from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        return _skus_from_table(rows)
    finally:
        workbook.close()


def _parse_xls(content: bytes) -> list[str]:
    try:
        import xlrd
    except ImportError as exc:
        raise WarehouseSkuCheckError(
            "Reading .xls requires xlrd. Save as .xlsx or .txt and try again."
        ) from exc
    try:
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        rows = [
            tuple(sheet.cell_value(r, c) for c in range(sheet.ncols))
            for r in range(sheet.nrows)
        ]
        return _skus_from_table(rows)
    except WarehouseSkuCheckError:
        raise
    except Exception as exc:
        raise WarehouseSkuCheckError(
            "Could not read the .xls file. Save as .xlsx or .txt."
        ) from exc


def parse_sku_list_file(filename: str | None, content: bytes) -> list[str]:
    """Extract unique SKUs from .txt / .csv / .xlsx / .xls, preserving order."""
    if not content:
        raise WarehouseSkuCheckError("Uploaded file is empty.")
    if len(content) > _MAX_UPLOAD_BYTES:
        raise WarehouseSkuCheckError("File is too large (max 15 MB).")

    name = (filename or "upload.txt").lower()
    if name.endswith((".xlsx", ".xlsm")) or content.startswith(b"PK"):
        skus = _parse_excel(content)
    elif name.endswith(".xls"):
        skus = _parse_xls(content)
    elif name.endswith(".csv"):
        skus = _parse_csv_bytes(content)
    else:
        skus = _parse_text(content)

    cleaned: list[str] = []
    seen: set[str] = set()
    for sku in skus:
        key = sku.strip()
        if not key or _is_sku_header(key):
            continue
        if key in seen:
            continue
        seen.add(key)
        cleaned.append(key)
        if len(cleaned) > _MAX_SKUS:
            raise WarehouseSkuCheckError(
                f"Too many SKUs (max {_MAX_SKUS:,}). Split the file and try again."
            )

    if not cleaned:
        raise WarehouseSkuCheckError(
            "No SKUs found. Use a .txt (one SKU per line), .csv, or .xlsx with a SKU column."
        )
    return cleaned


def build_check_rows(
    skus: Sequence[str],
    found_by_sku: Mapping[str, Sequence[Mapping[str, Any]]],
) -> list[SkuCheckRow]:
    rows: list[SkuCheckRow] = []
    for sku in skus:
        matches = list(found_by_sku.get(sku) or [])
        if not matches:
            rows.append(SkuCheckRow(sku=sku, exists=False, match_count=0))
            continue
        first = matches[0]
        rows.append(
            SkuCheckRow(
                sku=sku,
                exists=True,
                upc=str(first.get("upc") or ""),
                fnsku=str(first.get("fnsku") or ""),
                style_name=str(first.get("style_name") or ""),
                condition=str(first.get("condition") or ""),
                match_count=len(matches),
            )
        )
    return rows


def build_sku_check_workbook(rows: Sequence[SkuCheckRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKU Check"
    headers = ("SKU", "Exists", "UPC", "FNSKU", "Style Name", "Condition", "Match Count")
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = _CALIBRI_BOLD

    for index, row in enumerate(rows, start=2):
        values = (
            row.sku,
            "Yes" if row.exists else "No",
            row.upc or None,
            row.fnsku or None,
            row.style_name or None,
            row.condition or None,
            row.match_count,
        )
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=index, column=col, value=value)
            cell.font = _CALIBRI
            if col == 1:
                cell.number_format = "@"

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 14
    sheet.column_dimensions["E"].width = 40
    sheet.column_dimensions["F"].width = 12
    sheet.column_dimensions["G"].width = 12

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_sku_check_workbook(
    skus: Sequence[str],
    found_by_sku: Mapping[str, Sequence[Mapping[str, Any]]],
    *,
    filename: str = "SKU Existence Check.xlsx",
) -> SkuCheckResult:
    rows = build_check_rows(skus, found_by_sku)
    found = sum(1 for row in rows if row.exists)
    missing = len(rows) - found
    return SkuCheckResult(
        file_bytes=build_sku_check_workbook(rows),
        filename=filename,
        total=len(rows),
        found=found,
        missing=missing,
        unique_skus=len(rows),
    )
