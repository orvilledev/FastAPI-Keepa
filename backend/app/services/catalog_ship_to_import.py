"""Parse the Ship To Address workbook (Code / Full Address / Address 1 / City / State / Postal Code)."""
from __future__ import annotations

import io
import logging
from datetime import date, datetime
from typing import Any, Dict, List, Sequence, Tuple

from openpyxl import load_workbook

from app.services.catalog_ship_to_headers import HEADERS, REQUIRED_HEADER, SHEET_NAME

logger = logging.getLogger(__name__)


def _cell_str(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return "TRUE" if raw else "FALSE"
    if isinstance(raw, datetime):
        if raw.hour == 0 and raw.minute == 0 and raw.second == 0 and raw.microsecond == 0:
            return raw.date().isoformat()
        return raw.isoformat(sep=" ", timespec="seconds")
    if isinstance(raw, date):
        return raw.isoformat()
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if raw.is_integer():
            return str(int(raw))
        return format(raw, ".15g").strip()
    text = str(raw).strip()
    if text.startswith("="):
        return ""
    return text


def _normalize_header(cell: Any) -> str:
    return _cell_str(cell).strip()


def _find_sheet(wb, preferred: str):
    wanted = preferred.strip().upper()
    for name in wb.sheetnames:
        if name.strip().upper() == wanted:
            return wb[name]
    return wb[wb.sheetnames[0]] if wb.sheetnames else None


def compose_full_address(address_1: str, city: str, state: str, postal_code: str) -> str:
    """Match the source workbook CONCAT(C, \", \", D, \", \", E, \", \", F)."""
    return f"{address_1}, {city}, {state}, {postal_code}"


def _header_index_map(
    header_row: Sequence[Any],
    expected: Sequence[str],
    required: str,
) -> Dict[str, int]:
    found: Dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = _normalize_header(cell)
        if not name or name in found:
            continue
        found[name] = idx

    if required not in found:
        raise ValueError(
            f'Spreadsheet must include a "{required}" column. '
            "Download the template and match the exact headers."
        )

    missing = [h for h in expected if h not in found]
    if missing:
        preview = ", ".join(missing)
        raise ValueError(
            f"Missing required column(s): {preview}. "
            "Use the downloadable template (exact headers from the ship-to address list)."
        )
    return {h: found[h] for h in expected}


def parse_ship_to_spreadsheet(filename: str, content: bytes) -> Tuple[List[Dict[str, str]], int]:
    lower = (filename or "").lower()
    if not lower.endswith((".xlsx", ".xlsm", ".xls")):
        raise ValueError("Upload an .xlsx file matching the Ship To Address template.")

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = _find_sheet(wb, SHEET_NAME)
        if sheet is None:
            raise ValueError("Workbook has no sheets.")
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], 0

        mapping = _header_index_map(header_row, HEADERS, REQUIRED_HEADER)
        valid: List[Dict[str, str]] = []
        invalid = 0
        for row in rows_iter:
            row_data: Dict[str, str] = {}
            any_value = False
            for header in HEADERS:
                idx = mapping[header]
                value = _cell_str(row[idx]) if idx < len(row) else ""
                if value:
                    any_value = True
                row_data[header] = value
            if not any_value:
                continue
            key = (row_data.get(REQUIRED_HEADER) or "").strip()
            if not key:
                invalid += 1
                continue
            row_data["Full Address"] = compose_full_address(
                row_data.get("Address 1") or "",
                row_data.get("City") or "",
                row_data.get("State") or "",
                row_data.get("Postal Code") or "",
            )
            valid.append(row_data)
        return valid, invalid
    finally:
        wb.close()


def ship_to_row_to_record(row_data: Dict[str, str]) -> Dict[str, Any]:
    return {
        "code": (row_data.get("Code") or "").strip(),
        "full_address": row_data.get("Full Address") or "",
        "address_1": row_data.get("Address 1") or "",
        "city": row_data.get("City") or "",
        "state": row_data.get("State") or "",
        "postal_code": row_data.get("Postal Code") or "",
        "row_data": row_data,
    }


def dedupe_by_code(rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    """Last row wins for duplicate codes within one import file."""
    seen: Dict[str, Dict[str, str]] = {}
    for row in rows:
        key = (row.get("Code") or "").strip()
        if key:
            seen[key] = row
    return list(seen.values())
