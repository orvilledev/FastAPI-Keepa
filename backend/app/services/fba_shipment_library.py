"""Parse the official Amazon FBA shipment TSV.

The preamble Name and Shipment ID are kept, plus the yellow item fields:
Merchant SKU, external-id, and Shipped.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from io import StringIO

_UPC_PREFIX = re.compile(r"^UPC:\s*", re.IGNORECASE)


def strip_upc_label(value: str) -> str:
    """Drop a leading ``UPC:`` label so external-id stores the code only."""
    return _UPC_PREFIX.sub("", (value or "").strip()).strip()


STATUS_MISSING = "Missing Items"
STATUS_ACCOUNTED = "All Accounted"
STATUS_OVERAGE = "Overage"
STATUS_NOT_IN_LIST = "Not in the list"


def is_fully_accounted(shipped_and_counted: list[tuple[int, int]]) -> bool:
    """True when every line has been counted for exactly its shipped quantity."""
    if not shipped_and_counted:
        return False
    return all(counted == shipped for shipped, counted in shipped_and_counted)


def count_progress(shipped: int, counted: int) -> tuple[int, str]:
    """Missing is what is still left of Shipped. Status follows the scan total."""
    missing = max(0, shipped - counted)
    if counted > shipped:
        return missing, STATUS_OVERAGE
    if counted == shipped:
        return missing, STATUS_ACCOUNTED
    return missing, STATUS_MISSING


def share_audience(upload: dict) -> list[str] | None:
    """None means every worker. A list means only those user ids, plus the owner.

    A missing value or database NULL stays open to everyone, which is how
    shipments shared before custom sharing was added are stored.
    """
    if "shared_with" not in upload or upload.get("shared_with") is None:
        return None
    raw = upload.get("shared_with")
    if isinstance(raw, str):
        text = raw.strip()
        if not text or text.lower() == "null":
            return None
        text = text.strip("{}")
        if not text:
            return []
        return [part.strip().strip('"') for part in text.split(",") if part.strip()]
    if not isinstance(raw, list):
        return None
    return [str(item) for item in raw if item]


def shipment_open_to(upload: dict, user_id: str, *, can_manage_any: bool = False) -> bool:
    """True when this person may open the shipment."""
    if can_manage_any:
        return True
    if str(upload.get("created_by") or "") == str(user_id or ""):
        return True
    if upload.get("trashed_at"):
        return False
    if upload.get("is_locked"):
        return False
    audience = share_audience(upload)
    if audience is None:
        return True
    return str(user_id or "") in audience


def upload_visibility_filter(viewer_id: str) -> str:
    """PostgREST filter: owner, or an unlocked shipment shared with everyone or this person."""
    return (
        f"created_by.eq.{viewer_id},"
        f"and(is_locked.eq.false,or(shared_with.is.null,shared_with.cs.{{{viewer_id}}}))"
    )


def choose_line_for_scan(matches: list[dict]) -> dict:
    """Add the scan to the first matching line that still has units left."""
    ordered = sorted(matches, key=lambda row: int(row.get("line_number") or 0))
    for row in ordered:
        if int(row.get("counted") or 0) < int(row.get("shipped") or 0):
            return row
    return ordered[0]


class FbaShipmentParseError(ValueError):
    """The file is not the official FBA shipment template."""


@dataclass(frozen=True)
class FbaShipmentLine:
    line_number: int
    merchant_sku: str
    external_id: str
    shipped: int


@dataclass(frozen=True)
class ParsedFbaShipment:
    shipment_id: str
    lines: list[FbaShipmentLine]
    shipment_name: str = ""

    @property
    def total_shipped(self) -> int:
        return sum(line.shipped for line in self.lines)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FbaShipmentParseError("Could not read the file as text. Upload the official .tsv template.")


def _norm_header(value: str) -> str:
    return " ".join((value or "").strip().lower().replace("_", " ").replace("-", " ").split())


def _parse_shipped(raw: str, merchant_sku: str, row_number: int) -> int:
    text = (raw or "").strip().replace(",", "")
    if not text:
        raise FbaShipmentParseError(
            f"Row {row_number} ({merchant_sku}) is missing a Shipped quantity."
        )
    try:
        qty = int(text)
    except ValueError as exc:
        raise FbaShipmentParseError(
            f"Row {row_number} ({merchant_sku}) has a Shipped value that is not a whole number: {raw!r}."
        ) from exc
    if qty < 0:
        raise FbaShipmentParseError(
            f"Row {row_number} ({merchant_sku}) has a negative Shipped quantity."
        )
    return qty


def parse_fba_shipment_tsv(content: bytes) -> ParsedFbaShipment:
    """Read Shipment ID plus Merchant SKU, external-id, and Shipped from each item row."""
    if not content or not content.strip():
        raise FbaShipmentParseError("Uploaded file is empty.")

    text = _decode(content)
    if text.count("\t") < 3 and "Merchant SKU" not in text and "merchant sku" not in text.lower():
        raise FbaShipmentParseError(
            "This file does not look like the official FBA shipment .tsv. "
            "Upload the tab-separated template with Shipment ID, Merchant SKU, external-id, and Shipped."
        )

    rows = list(csv.reader(StringIO(text), delimiter="\t"))
    header_index = None
    columns: dict[str, int] = {}
    for index, row in enumerate(rows):
        normalized = [_norm_header(cell) for cell in row]
        if "merchant sku" in normalized and "shipped" in normalized and "external id" in normalized:
            header_index = index
            columns = {name: pos for pos, name in enumerate(normalized) if name}
            break

    if header_index is None:
        raise FbaShipmentParseError(
            "Could not find the item header row. The official template needs Merchant SKU, external-id, and Shipped."
        )

    shipment_id = ""
    shipment_name = ""
    for row in rows[:header_index]:
        if len(row) < 2:
            continue
        label = _norm_header(row[0])
        value = (row[1] or "").strip()
        if label == "shipment id" and not shipment_id:
            shipment_id = value
        elif label in {"name", "shipment name"} and not shipment_name:
            shipment_name = value
    if not shipment_id:
        raise FbaShipmentParseError("Shipment ID is missing from the top of the file.")

    sku_col = columns["merchant sku"]
    external_col = columns["external id"]
    shipped_col = columns["shipped"]
    width = max(sku_col, external_col, shipped_col)

    lines: list[FbaShipmentLine] = []
    for offset, row in enumerate(rows[header_index + 1 :], start=1):
        if not any((cell or "").strip() for cell in row):
            continue
        padded = list(row) + [""] * (width + 1 - len(row))
        merchant_sku = (padded[sku_col] or "").strip()
        if not merchant_sku:
            continue
        external_id = strip_upc_label(padded[external_col] or "")
        shipped = _parse_shipped(padded[shipped_col], merchant_sku, header_index + 1 + offset)
        lines.append(
            FbaShipmentLine(
                line_number=len(lines) + 1,
                merchant_sku=merchant_sku,
                external_id=external_id,
                shipped=shipped,
            )
        )

    if not lines:
        raise FbaShipmentParseError("The file has a Shipment ID but no item rows to store.")

    return ParsedFbaShipment(shipment_id=shipment_id, lines=lines, shipment_name=shipment_name)


def choose_blocking_shipment_copy(rows: list[dict]) -> dict | None:
    """The copy that should stop another upload of the same Shipment ID.

    An active copy is preferred, then a completed one, then one in Trash.
    Existing extra copies are left in place; this only decides which one to mention.
    """
    if not rows:
        return None

    def rank(row: dict) -> tuple[int, str]:
        if row.get("trashed_at"):
            place = 2
        elif row.get("archived_at"):
            place = 1
        else:
            place = 0
        return (place, str(row.get("created_at") or ""))

    return min(rows, key=rank)


def duplicate_shipment_message(existing: dict) -> str:
    """Popup text when a Shipment ID is already stored."""
    if existing.get("trashed_at"):
        place = " The existing copy is in Trash."
    elif existing.get("archived_at"):
        place = " The existing copy is in Completed shipments."
    else:
        place = ""
    return (
        "A copy of this shipment is already uploaded."
        f"{place} Uploading this second copy was blocked."
    )


_UNSAFE_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def _filename_stem(value: str) -> str:
    """A single path segment that is safe to put in a download header."""
    cleaned = _UNSAFE_FILENAME.sub(" ", value or "")
    cleaned = " ".join(cleaned.split()).strip(" .")
    try:
        cleaned.encode("latin-1")
    except UnicodeEncodeError:
        cleaned = cleaned.encode("latin-1", "replace").decode("latin-1")
        cleaned = cleaned.replace("?", " ")
        cleaned = " ".join(cleaned.split()).strip(" .")
    return cleaned[:180]


def accounting_report_filename(shipment_name: str, shipment_id: str = "") -> str:
    """Use the upload's Name as the download filename, not the Shipment ID.

    Shipments stored before that field was saved keep the previous
    ``{Shipment ID} Accounting.xlsx`` name.
    """
    stem = _filename_stem(shipment_name)
    if stem:
        return f"{stem}.xlsx"
    fallback = _filename_stem(shipment_id) or "shipment"
    return f"{fallback} Accounting.xlsx"


BOXED_SHEET = "Boxed"
BOXED_HEADER = ["external-id", "Counted", "Box Number"]
MAX_BOX_NUMBER = 99999
MISSING_BOX_FILL = "FF0000"

BoxedRow = tuple[str, "int | None"]


def _occupied_boxes(counts: dict[int, int]) -> set[int]:
    return {int(box) for box, count in counts.items() if int(count) > 0}


def _missing_box_numbers(boxes: set[int]) -> set[int]:
    if not boxes:
        return set()
    highest = max(boxes)
    return {number for number in range(1, highest + 1) if number not in boxes}


def _next_box_number(boxes: set[int]) -> int:
    """Lowest box number that is not already in use, starting at 1."""
    number = 1
    while number in boxes:
        number += 1
    return number


def assert_box_number_follows(
    box_number: int,
    counts: dict[int, int],
    *,
    previous_box: int | None = None,
) -> None:
    """Refuse a box number that would skip 1, 2, 3, and so on.

    A number that already has units can be used again. A new number has to be
    the lowest one that is still open. Moving the last unit out of a box cannot
    leave a gap under a higher box.
    """
    box_number = int(box_number)
    if previous_box is not None and int(previous_box) == box_number:
        return

    before = _occupied_boxes(counts)
    without_this_unit = set(before)
    if previous_box is not None and int(counts.get(int(previous_box), 0)) <= 1:
        without_this_unit.discard(int(previous_box))

    if box_number not in without_this_unit:
        expected = _next_box_number(without_this_unit)
        if box_number != expected:
            raise ValueError(
                f"Box {box_number} skips a box number. The next box is {expected}."
            )

    after = set(without_this_unit)
    after.add(box_number)
    created = _missing_box_numbers(after) - _missing_box_numbers(before)
    if created:
        skipped = min(created)
        raise ValueError(
            f"Box {skipped} cannot be left empty while a higher box still has units."
        )


def box_count_and_missing(rows: list[BoxedRow]) -> tuple[int, int]:
    """Highest box number, and how many of 1 through that number have no units."""
    used = {box for _, box in rows if box is not None}
    if not used:
        return 0, 0
    highest = max(used)
    missing = sum(1 for number in range(1, highest + 1) if number not in used)
    return highest, missing


def accounting_totals_by_external_id(lines: list[dict], unknown_lines: list[dict]) -> dict[str, int]:
    """Counted units per external-id exactly as the Accounting sheet reports them."""
    totals: dict[str, int] = {}
    for row in [*lines, *unknown_lines]:
        external_id = strip_upc_label(str(row.get("external_id") or ""))
        counted = int(row.get("counted") or 0)
        if not external_id or counted <= 0:
            continue
        totals[external_id] = totals.get(external_id, 0) + counted
    return totals


def counts_kept_in_box(
    scans: list[dict],
    box_number: int,
    limits: dict[str, int],
    key: str,
) -> dict[str, int]:
    """How many counted units of each item sit in one box.

    Only the first ``limits`` scans for an item are counted units, in scan order.
    Later scans and units in other boxes are left out.
    """
    grouped: dict[str, list[dict]] = {}
    for scan in scans:
        raw = str(scan.get(key) or "")
        item = strip_upc_label(raw) if key == "external_id" else raw
        if not item:
            continue
        if key == "external_id" and scan.get("line_id"):
            continue
        grouped.setdefault(item, []).append(scan)
    target = int(box_number)
    counts: dict[str, int] = {}
    for item, item_scans in grouped.items():
        limit = max(0, int(limits.get(item) or 0))
        if limit == 0:
            continue
        ordered = sorted(
            item_scans,
            key=lambda row: (str(row.get("created_at") or ""), str(row.get("id") or "")),
        )
        kept = ordered[:limit]
        count = sum(
            1
            for scan in kept
            if scan.get("box_number") is not None and int(scan["box_number"]) == target
        )
        if count:
            counts[item] = count
    return counts


def units_for_count(counted: int, scans: list[dict]) -> list[dict]:
    """One slot per counted unit. Saved scans fill the slots in scan order; the rest have no box."""
    counted = max(0, int(counted or 0))
    kept = list(scans)[:counted]
    units: list[dict] = []
    for index, scan in enumerate(kept, start=1):
        box_number = scan.get("box_number")
        units.append(
            {
                "unit_number": index,
                "scan_id": str(scan.get("id") or "") or None,
                "box_number": int(box_number) if box_number is not None else None,
            }
        )
    for index in range(len(kept) + 1, counted + 1):
        units.append({"unit_number": index, "scan_id": None, "box_number": None})
    return units


def build_boxed_rows(
    *,
    lines: list[dict],
    unknown_lines: list[dict],
    box_scans: list[dict],
) -> list[BoxedRow]:
    """One (external-id, box number) row per counted unit, matching the Accounting sheet.

    ``box_scans`` must already be in scan order. A scan is kept only while the
    external-id still has units in the Accounting count, so Boxed can never exceed
    Accounting. Units counted without a box (for example before boxes existed) are
    added at the end with a blank box number, so the two sheets always agree.
    """
    expected = accounting_totals_by_external_id(lines, unknown_lines)
    placed: dict[str, int] = {}
    rows: list[BoxedRow] = []
    for scan in box_scans:
        external_id = strip_upc_label(str(scan.get("external_id") or ""))
        if not external_id or placed.get(external_id, 0) >= expected.get(external_id, 0):
            continue
        box_number = scan.get("box_number")
        if box_number is None:
            continue
        rows.append((external_id, int(box_number)))
        placed[external_id] = placed.get(external_id, 0) + 1
    for external_id in sorted(expected):
        for _ in range(expected[external_id] - placed.get(external_id, 0)):
            rows.append((external_id, None))
    rows.sort(key=lambda row: (row[1] is None, row[1] or 0, row[0]))
    return rows


def _matches_label_search(line: dict, needle: str) -> bool:
    """Same match as the shipment table: merchant SKU or external-id contains the term."""
    sku = str(line.get("merchant_sku") or "").lower()
    code = strip_upc_label(str(line.get("external_id") or "")).lower()
    return needle in sku or needle in code


def labels_by_box(
    *,
    lines: list[dict],
    box_scans: list[dict],
    box_number: int | None = None,
    search: str | None = None,
    unknown_lines: list[dict] | None = None,
) -> list[tuple[int, list[tuple[str, int]]]]:
    """Counted units grouped for printing: box 1, then box 2, and so on.

    Within a box, identical external-ids are one quantity. ``box_number`` and
    ``search`` keep only the rows the shipment table is showing. A box filter
    does not require units outside that box to be boxed. With no filter, units
    that have no box still stop the print, and scans that are not on the
    shipment are left out. A filter includes those scans when they match.
    """
    needle = (search or "").strip().lower()
    selected_lines = lines
    selected_unknown: list[dict] = []
    if needle:
        selected_lines = [line for line in lines if _matches_label_search(line, needle)]
        selected_unknown = [
            line for line in (unknown_lines or []) if _matches_label_search(line, needle)
        ]
    elif box_number is not None:
        selected_unknown = list(unknown_lines or [])
    rows = build_boxed_rows(
        lines=selected_lines,
        unknown_lines=selected_unknown,
        box_scans=box_scans,
    )
    if box_number is not None:
        rows = [row for row in rows if row[1] == box_number]
    if not rows:
        if needle or box_number is not None:
            raise ValueError("Nothing in this filter has a counted unit to print.")
        raise ValueError("There are no counted units to print.")
    if any(box is None for _, box in rows):
        raise ValueError(
            "Every counted unit needs a box number before FNSKU labels can be printed."
        )
    grouped: dict[int, dict[str, int]] = {}
    for external_id, box in rows:
        items = grouped.setdefault(int(box), {})
        items[external_id] = items.get(external_id, 0) + 1
    return [(box, list(items.items())) for box, items in grouped.items()]


def box_summary(rows: list[BoxedRow]) -> dict[str, str]:
    """Per external-id, the boxes its counted units are in, e.g. ``Box 1: 2, Box 4: 2``."""
    per_item: dict[str, dict["int | None", int]] = {}
    for external_id, box in rows:
        boxes = per_item.setdefault(external_id, {})
        boxes[box] = boxes.get(box, 0) + 1
    summary: dict[str, str] = {}
    for external_id, boxes in per_item.items():
        parts = [f"Box {box}: {boxes[box]}" for box in sorted(b for b in boxes if b is not None)]
        if None in boxes:
            parts.append(f"No box: {boxes[None]}")
        summary[external_id] = ", ".join(parts)
    return summary


def _verify_boxed_matches_accounting(
    rows: list[BoxedRow],
    lines: list[dict],
    unknown_lines: list[dict],
) -> None:
    expected = accounting_totals_by_external_id(lines, unknown_lines)
    boxed: dict[str, int] = {}
    for external_id, _ in rows:
        boxed[external_id] = boxed.get(external_id, 0) + 1
    if boxed != expected:
        raise ValueError(
            "The Boxed totals do not match the Accounting totals, so the report was not created. "
            "Scan the shipment again or contact support."
        )


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _add_boxed_pivot(sheet, rows: list[BoxedRow], first_column: int = 7) -> None:
    """Excel PivotTable beside the Boxed rows: external-id down, Box Number across, sum of Counted."""
    from datetime import datetime, timezone

    from openpyxl.pivot.cache import (
        CacheDefinition,
        CacheField,
        CacheSource,
        SharedItems,
        WorksheetSource,
    )
    from openpyxl.pivot.fields import Index, Missing, Number, Text
    from openpyxl.pivot.record import Record, RecordList
    from openpyxl.pivot.table import (
        DataField,
        FieldItem,
        Location,
        PivotField,
        PivotTableStyle,
        RowColField,
        RowColItem,
        TableDefinition,
    )
    from openpyxl.styles import Font

    external_ids = sorted({external_id for external_id, _ in rows})
    box_numbers = sorted({box for _, box in rows if box is not None})
    has_blank = any(box is None for _, box in rows)
    box_items: list["int | None"] = [*box_numbers, *([None] if has_blank else [])]
    ext_index = {value: pos for pos, value in enumerate(external_ids)}
    box_index = {value: pos for pos, value in enumerate(box_items)}

    cache = CacheDefinition(
        refreshOnLoad=True,
        refreshedBy="MSW Overwatch",
        refreshedDate=(
            datetime.now(timezone.utc).replace(tzinfo=None) - datetime(1899, 12, 30)
        ).total_seconds()
        / 86400,
        createdVersion=8,
        refreshedVersion=8,
        minRefreshableVersion=3,
        recordCount=len(rows),
        cacheSource=CacheSource(
            type="worksheet",
            worksheetSource=WorksheetSource(ref=f"A1:C{len(rows) + 1}", sheet=BOXED_SHEET),
        ),
        cacheFields=[
            CacheField(
                name=BOXED_HEADER[0],
                numFmtId=0,
                sharedItems=SharedItems(
                    _fields=[Text(v=value) for value in external_ids],
                    count=len(external_ids),
                ),
            ),
            CacheField(
                name=BOXED_HEADER[1],
                numFmtId=0,
                sharedItems=SharedItems(
                    containsSemiMixedTypes=False,
                    containsString=False,
                    containsNumber=True,
                    containsInteger=True,
                    minValue=1,
                    maxValue=1,
                ),
            ),
            CacheField(
                name=BOXED_HEADER[2],
                numFmtId=0,
                sharedItems=SharedItems(
                    _fields=[Missing() if box is None else Number(v=box) for box in box_items],
                    containsSemiMixedTypes=None if has_blank else False,
                    containsString=False,
                    containsBlank=True if has_blank else None,
                    containsNumber=True if box_numbers else None,
                    containsInteger=True if box_numbers else None,
                    minValue=min(box_numbers) if box_numbers else None,
                    maxValue=max(box_numbers) if box_numbers else None,
                    count=len(box_items),
                ),
            ),
        ],
    )
    cache.records = RecordList(
        count=len(rows),
        r=[
            Record(
                _fields=[
                    Index(v=ext_index[external_id]),
                    Number(v=1),
                    Index(v=box_index[box]),
                ]
            )
            for external_id, box in rows
        ],
    )

    last_column = first_column + 1 + len(box_items)
    last_row = 2 + len(external_ids) + 1
    pivot = TableDefinition(
        name="PivotTable1",
        cacheId=1,
        dataCaption="Values",
        updatedVersion=8,
        minRefreshableVersion=3,
        createdVersion=8,
        useAutoFormatting=True,
        itemPrintTitles=True,
        indent=0,
        outline=True,
        outlineData=True,
        multipleFieldFilters=False,
        applyWidthHeightFormats=True,
        location=Location(
            ref=f"{_column_letter(first_column)}1:{_column_letter(last_column)}{last_row}",
            firstHeaderRow=1,
            firstDataRow=2,
            firstDataCol=1,
        ),
        pivotFields=[
            PivotField(
                axis="axisRow",
                showAll=False,
                items=[
                    *[FieldItem(x=pos) for pos in range(len(external_ids))],
                    FieldItem(t="default"),
                ],
            ),
            PivotField(dataField=True, showAll=False),
            PivotField(
                axis="axisCol",
                showAll=False,
                items=[
                    *[FieldItem(x=pos) for pos in range(len(box_items))],
                    FieldItem(t="default"),
                ],
            ),
        ],
        rowFields=[RowColField(x=0)],
        rowItems=[
            *[RowColItem(t="data", x=[Index(v=pos)]) for pos in range(len(external_ids))],
            RowColItem(t="grand", x=[Index(v=0)]),
        ],
        colFields=[RowColField(x=2)],
        colItems=[
            *[RowColItem(t="data", x=[Index(v=pos)]) for pos in range(len(box_items))],
            RowColItem(t="grand", x=[Index(v=0)]),
        ],
        dataFields=[DataField(name="Sum of Counted", fld=1, baseField=0, baseItem=0)],
        pivotTableStyleInfo=PivotTableStyle(
            name="PivotStyleLight16",
            showRowHeaders=True,
            showColHeaders=True,
            showRowStripes=False,
            showColStripes=False,
            showLastColumn=True,
        ),
    )
    pivot.cache = cache
    sheet.add_pivot(pivot)

    # The cells Excel shows until the pivot refreshes. refreshOnLoad rebuilds it on open.
    bold = Font(bold=True)
    grid: dict[tuple[str, "int | None"], int] = {}
    for external_id, box in rows:
        grid[(external_id, box)] = grid.get((external_id, box), 0) + 1

    def put(row: int, offset: int, value, *, is_bold: bool = False) -> None:
        cell = sheet.cell(row=row, column=first_column + offset, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        if is_bold:
            cell.font = bold

    put(1, 0, "Sum of Counted", is_bold=True)
    put(1, 1, "Column Labels", is_bold=True)
    put(2, 0, "Row Labels", is_bold=True)
    for pos, box in enumerate(box_items):
        put(2, 1 + pos, "(blank)" if box is None else box, is_bold=True)
    put(2, 1 + len(box_items), "Grand Total", is_bold=True)
    for row_pos, external_id in enumerate(external_ids):
        row_number = 3 + row_pos
        put(row_number, 0, external_id)
        row_total = 0
        for pos, box in enumerate(box_items):
            count = grid.get((external_id, box), 0)
            if count:
                put(row_number, 1 + pos, count)
                row_total += count
        put(row_number, 1 + len(box_items), row_total, is_bold=True)
    total_row = 3 + len(external_ids)
    put(total_row, 0, "Grand Total", is_bold=True)
    for pos, box in enumerate(box_items):
        put(total_row, 1 + pos, sum(c for (_, b), c in grid.items() if b == box), is_bold=True)
    put(total_row, 1 + len(box_items), len(rows), is_bold=True)

    sheet.column_dimensions[_column_letter(first_column)].width = 18
    for pos in range(len(box_items) + 1):
        sheet.column_dimensions[_column_letter(first_column + 1 + pos)].width = 12


def _add_boxed_sheet(workbook, rows: list[BoxedRow]) -> None:
    from openpyxl.styles import Font

    sheet = workbook.create_sheet(BOXED_SHEET)
    sheet.append(BOXED_HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for external_id, box in rows:
        sheet.append([external_id, 1, box])
        sheet.cell(row=sheet.max_row, column=1).data_type = "s"
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 10
    sheet.column_dimensions["C"].width = 13
    if rows:
        _add_boxed_pivot(sheet, rows)


def build_accounting_workbook(
    *,
    shipment_id: str,
    counted_by: str,
    lines: list[dict],
    unknown_lines: list[dict],
    box_scans: list[dict] | None = None,
    shipment_name: str = "",
) -> bytes:
    """Excel report of one user's count against the shipment.

    Sheet 1 (Accounting) is the per-SKU summary. Sheet 2 (Boxed) lists every counted
    unit with its box number plus a pivot, and always adds up to the Accounting sheet.
    """
    from io import BytesIO
    from datetime import datetime, timezone

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fills = {
        STATUS_ACCOUNTED: PatternFill("solid", fgColor="BBF7D0"),
        STATUS_MISSING: PatternFill("solid", fgColor="FDE68A"),
        STATUS_OVERAGE: PatternFill("solid", fgColor="FDBA74"),
        STATUS_NOT_IN_LIST: PatternFill("solid", fgColor="FECACA"),
    }
    boxed_rows = build_boxed_rows(lines=lines, unknown_lines=unknown_lines, box_scans=box_scans or [])
    _verify_boxed_matches_accounting(boxed_rows, lines, unknown_lines)
    boxes_for = box_summary(boxed_rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accounting"
    sheet.append(["Name", shipment_name])
    sheet.append(["Shipment ID", shipment_id])
    sheet.append(["Counted by", counted_by])
    sheet.append(["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    box_count, missing_boxes = box_count_and_missing(boxed_rows)
    sheet["D1"] = "Box count"
    sheet["E1"] = box_count
    sheet["D2"] = "Missing box"
    if missing_boxes:
        missing_cell = sheet["E2"]
        missing_cell.value = missing_boxes
        missing_cell.fill = PatternFill("solid", fgColor=MISSING_BOX_FILL)
    sheet.column_dimensions["D"].width = 16
    sheet.column_dimensions["E"].width = 12
    sheet.append([])
    header = ["Merchant SKU", "external-id", "Shipped", "Counted", "Missing", "Status", "Boxes"]
    sheet.append(header)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for line in lines:
        shipped = int(line.get("shipped") or 0)
        counted = int(line.get("counted") or 0)
        missing, status = count_progress(shipped, counted)
        sheet.append([
            line.get("merchant_sku") or "",
            line.get("external_id") or "",
            shipped,
            counted,
            missing,
            status,
            boxes_for.get(strip_upc_label(str(line.get("external_id") or "")), "") if counted > 0 else "",
        ])
        fill = fills.get(status)
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill
    if unknown_lines:
        sheet.append([])
        sheet.append(["Not in the list"])
        sheet.append(["external-id", "Counted", "Status", "Boxes"])
        for row in unknown_lines:
            sheet.append([
                row.get("external_id") or "",
                int(row.get("counted") or 0),
                STATUS_NOT_IN_LIST,
                boxes_for.get(strip_upc_label(str(row.get("external_id") or "")), ""),
            ])
            for cell in sheet[sheet.max_row]:
                cell.fill = fills[STATUS_NOT_IN_LIST]
    sheet.column_dimensions["G"].width = 32
    _add_boxed_sheet(workbook, boxed_rows)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
