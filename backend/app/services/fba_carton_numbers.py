"""NFA and SMW Tool — carton dumps → Carton Numbers workbook.

Input: any number of FBA carton dumps (the same spaced ``PO#:`` exports the
Output generator reads, ``.xls`` / ``.xlsx`` / ``.xlsm``), loose or inside ``.zip``
archives.

Output: one workbook, one sheet per shipment (named after its ``PO#`` / FBA id):

    Original Carton | (blank)               | Box Number
    Carton#:        | 00001928250570412008  | 1
    Carton#:        | 00008273992079477459  | 2
    ...

Every ``Carton#:`` row in the source becomes the next box number, in the order
the cartons appear. The carton barcode is kept as text so its leading zeros and
all 20 digits survive.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath, PureWindowsPath

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.fba_box_contents_tool2 import (
    FbaBoxContentsTool2Error,
    _as_excel_number,
    _cell_at,
    _cell_text,
    _is_carton_header,
    _normalize_label,
    _read_sheet_rows,
)

OUTPUT_FILENAME = "Carton Numbers.xlsx"
HEADER_ORIGINAL = "Original Carton"
HEADER_BOX = "Box Number"
CARTON_LABEL = "Carton#:"

MAX_FILES = 100
MAX_WORKBOOK_BYTES = 15 * 1024 * 1024
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000
MAX_EXPANDED_BYTES = 200 * 1024 * 1024

_EXCEL_SUFFIXES = (".xls", ".xlsx", ".xlsm")
_SHEET_NAME_LIMIT = 31
_FOOTER_LABELS = {"total ctns:", "total ctns", "total cartons:", "total cartons"}
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_CARTON_COLUMN_WIDTH = 21.33203125


class FbaCartonNumbersError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class ShipmentCartons:
    shipment_id: str
    source_name: str
    carton_ids: tuple[str, ...]


@dataclass(frozen=True)
class FbaCartonNumbersResult:
    file_bytes: bytes
    filename: str
    shipments: tuple[ShipmentCartons, ...]
    skipped_count: int

    @property
    def shipment_count(self) -> int:
        return len(self.shipments)

    @property
    def box_count(self) -> int:
        return sum(len(shipment.carton_ids) for shipment in self.shipments)


def _base_name(name: str) -> str:
    cleaned = (name or "").replace("\x00", "")
    return PureWindowsPath(PurePosixPath(cleaned).name).name


def _is_excel_name(name: str) -> bool:
    return name.lower().endswith(_EXCEL_SUFFIXES)


def _is_zip_name(name: str) -> bool:
    return name.lower().endswith(".zip")


def _is_junk_entry(path: str) -> bool:
    parts = [part for part in re.split(r"[\\/]+", path) if part]
    if not parts:
        return True
    if any(part == "__MACOSX" for part in parts):
        return True
    leaf = parts[-1]
    return leaf.startswith(".") or leaf.startswith("~$")


def _expand_zip(
    archive_name: str, content: bytes, budget: list[int]
) -> tuple[list[tuple[str, bytes]], int]:
    """Return (workbooks, skipped_count) from one archive. ``budget`` is remaining bytes."""
    if len(content) > MAX_ARCHIVE_BYTES:
        raise FbaCartonNumbersError(f'"{archive_name}" is too large (max 50 MB per zip).')
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise FbaCartonNumbersError(
            f'"{archive_name}" is not a valid zip file.'
        ) from exc

    workbooks: list[tuple[str, bytes]] = []
    skipped = 0
    with archive:
        infos = archive.infolist()
        if len(infos) > MAX_ARCHIVE_ENTRIES:
            raise FbaCartonNumbersError(
                f'"{archive_name}" has too many files (max {MAX_ARCHIVE_ENTRIES}).'
            )
        for info in infos:
            if info.is_dir() or _is_junk_entry(info.filename):
                continue
            leaf = _base_name(info.filename)
            label = f"{archive_name} > {leaf}"
            if _is_zip_name(leaf):
                raise FbaCartonNumbersError(
                    f'"{label}" is a zip inside a zip. Unzip it and upload the files instead.'
                )
            if not _is_excel_name(leaf):
                skipped += 1
                continue
            if info.flag_bits & 0x1:
                raise FbaCartonNumbersError(
                    f'"{label}" is password protected. Remove the password and upload it again.'
                )
            if info.file_size > MAX_WORKBOOK_BYTES:
                raise FbaCartonNumbersError(f'"{label}" is too large (max 15 MB).')
            with archive.open(info) as handle:
                raw = handle.read(MAX_WORKBOOK_BYTES + 1)
            if len(raw) > MAX_WORKBOOK_BYTES:
                raise FbaCartonNumbersError(f'"{label}" is too large (max 15 MB).')
            budget[0] -= len(raw)
            if budget[0] < 0:
                raise FbaCartonNumbersError(
                    "The uploaded zips expand to too much data. Upload fewer files."
                )
            if not raw:
                raise FbaCartonNumbersError(f'"{label}" is empty.')
            workbooks.append((label, raw))
    return workbooks, skipped


def expand_uploads(
    uploads: list[tuple[str, bytes]],
) -> tuple[list[tuple[str, bytes]], int]:
    """Flatten loose workbooks and zip archives into ``(name, bytes)`` workbooks."""
    if not uploads:
        raise FbaCartonNumbersError("No files were uploaded.")

    budget = [MAX_EXPANDED_BYTES]
    workbooks: list[tuple[str, bytes]] = []
    skipped = 0
    for name, content in uploads:
        display = _base_name(name) or "upload"
        if _is_zip_name(display):
            extracted, extra_skipped = _expand_zip(display, content, budget)
            workbooks.extend(extracted)
            skipped += extra_skipped
        elif _is_excel_name(display):
            if not content:
                raise FbaCartonNumbersError(f'"{display}" is empty.')
            if len(content) > MAX_WORKBOOK_BYTES:
                raise FbaCartonNumbersError(f'"{display}" is too large (max 15 MB).')
            workbooks.append((display, content))
        else:
            raise FbaCartonNumbersError(
                f'"{display}" is not supported. Upload .xls, .xlsx, .xlsm, or .zip files.'
            )

    if not workbooks:
        raise FbaCartonNumbersError(
            "No Excel files were found. Upload .xls, .xlsx, or .xlsm files, or a zip that contains them."
        )
    if len(workbooks) > MAX_FILES:
        raise FbaCartonNumbersError(f"Upload at most {MAX_FILES} carton files at a time.")
    return workbooks, skipped


def _find_shipment_id(rows: list[tuple[object, ...]], fallback: str) -> str:
    for row in rows[:12]:
        first = _normalize_label(_cell_at(row, 0)).replace(" ", "")
        second = _cell_text(_cell_at(row, 1))
        if first in {"po#:", "po#", "po:"} and second:
            return second
        candidate = _cell_text(_cell_at(row, 2))
        if candidate.upper().startswith("FBA"):
            return candidate
    return fallback


def _footer_carton_count(rows: list[tuple[object, ...]]) -> int | None:
    for row in reversed(rows):
        if _normalize_label(_cell_at(row, 0)) not in _FOOTER_LABELS:
            continue
        for value in row[1:]:
            number = _as_excel_number(value)
            if number is not None:
                return int(number)
    return None


def _carton_id(row: tuple[object, ...], label: str, row_number: int) -> str:
    for value in row[1:]:
        if value is None or isinstance(value, bool):
            continue
        if isinstance(value, (int, float)):
            raise FbaCartonNumbersError(
                f'"{label}" row {row_number}: the carton number is stored as a number, '
                "so its leading zeros and digits can be lost. Export the file again "
                "with the carton numbers as text."
            )
        text = str(value).strip()
        if text:
            return text
    raise FbaCartonNumbersError(f'"{label}" row {row_number}: Carton#: has no carton number.')


def parse_carton_numbers(content: bytes, source_name: str) -> ShipmentCartons:
    """Read one carton dump and return its shipment id and carton barcodes in order."""
    try:
        rows = _read_sheet_rows(content)
    except FbaBoxContentsTool2Error as exc:
        raise FbaCartonNumbersError(f'"{source_name}": {exc}') from exc

    stem = re.sub(r"\.(xls|xlsx|xlsm)$", "", _base_name(source_name.split(" > ")[-1]), flags=re.I)
    shipment_id = _find_shipment_id(rows, stem)

    carton_ids: list[str] = []
    for index, row in enumerate(rows, start=1):
        if _is_carton_header(_cell_at(row, 0)):
            carton_ids.append(_carton_id(row, source_name, index))

    if not carton_ids:
        raise FbaCartonNumbersError(
            f'"{source_name}" has no "Carton#:" rows. Upload the FBA carton dump for the shipment.'
        )

    footer = _footer_carton_count(rows)
    if footer is not None and footer != len(carton_ids):
        raise FbaCartonNumbersError(
            f'"{source_name}": the file says {footer} cartons in its Total Ctns footer '
            f"but {len(carton_ids)} Carton#: rows were found."
        )

    return ShipmentCartons(
        shipment_id=shipment_id,
        source_name=source_name,
        carton_ids=tuple(carton_ids),
    )


def sheet_title(shipment_id: str) -> str:
    """Excel-safe sheet name: no ``[]:*?/\\``, no edge apostrophes, 31 characters."""
    cleaned = re.sub(r"[\[\]:*?/\\]", "-", shipment_id).strip().strip("'")
    cleaned = cleaned[:_SHEET_NAME_LIMIT].strip().strip("'")
    return cleaned or "Shipment"


def build_carton_numbers_workbook(shipments: list[ShipmentCartons]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    seen: dict[str, ShipmentCartons] = {}
    for shipment in shipments:
        title = sheet_title(shipment.shipment_id)
        key = title.lower()
        if key in seen:
            raise FbaCartonNumbersError(
                f"Shipment {title} is in both \"{seen[key].source_name}\" and "
                f"\"{shipment.source_name}\". Upload each shipment once."
            )
        seen[key] = shipment

        sheet = workbook.create_sheet(title)
        header_original = sheet.cell(row=1, column=1, value=HEADER_ORIGINAL)
        header_original.font = _CALIBRI_BOLD
        header_box = sheet.cell(row=1, column=3, value=HEADER_BOX)
        header_box.font = _CALIBRI_BOLD

        for box_number, carton_id in enumerate(shipment.carton_ids, start=1):
            row = box_number + 1
            sheet.cell(row=row, column=1, value=CARTON_LABEL)
            sheet.cell(row=row, column=2, value=carton_id).data_type = "s"
            sheet.cell(row=row, column=3, value=box_number)

        sheet.column_dimensions["B"].width = _CARTON_COLUMN_WIDTH

    workbook.active = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_carton_numbers(uploads: list[tuple[str, bytes]]) -> FbaCartonNumbersResult:
    workbooks, skipped = expand_uploads(uploads)
    shipments = [parse_carton_numbers(content, name) for name, content in workbooks]
    file_bytes = build_carton_numbers_workbook(shipments)
    return FbaCartonNumbersResult(
        file_bytes=file_bytes,
        filename=OUTPUT_FILENAME,
        shipments=tuple(shipments),
        skipped_count=skipped,
    )
