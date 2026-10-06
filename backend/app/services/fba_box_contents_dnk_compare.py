"""DNK Box Contents vs Amazon manifest → corrected Box Contents workbook.

Takes a DNK-style Box Contents workbook (Contents + Dimensions) and an Amazon
Seller Central manifest (``Box packing information``), remaps Old SKUs via the
catalog, and forces quantities to match the manifest (source of truth).

Missing / shortfall units are appended to the last box and highlighted yellow.
The output mirrors the manual ``… Box Contents - corrected.xlsx`` layout:
Contents (Box Number, UPC, Quantity), an E–G last-box adjustment summary, a
Sum-of-Quantity pivot starting at column I when adjustments exist, and the
original Dimensions sheet layout.
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.fba_upload_compare import (
    ContentRow,
    ParsedAmz,
    apply_old_sku_remaps,
    build_upc_to_old_sku_map,
    parse_amz_upload,
    reconcile_to_amz,
)

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_CENTER = Alignment(horizontal="center")
_LEFT = Alignment(horizontal="left")

CONTENTS_SHEET = "Contents"
DIMENSIONS_SHEET = "Dimensions"
DEFAULT_OUTPUT_FILENAME = "Box Contents - corrected.xlsx"

_YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
# Matches the peach / theme Accent fill used on the E–G summary in the sample.
_SUMMARY_FILL = PatternFill(
    fill_type="solid",
    fgColor=Color(theme=9, tint=0.7999816888943144),
)

PIVOT_START_COL_DEFAULT = 7  # G — no last-box summary
PIVOT_START_COL_WITH_SUMMARY = 9  # I — room for E–G summary
_SUMMARY_COL = 5  # E
_LINE_START_ROW = 2
_QTY_NUMBER_FORMAT = r"#,##0;(#,##0)"
_WEIGHT_NUMBER_FORMAT = r"#,##0.0;(#,##0.0)"
_DIM_INT_NUMBER_FORMAT = r"#,##0;(#,##0)"


class FbaBoxContentsDnkCompareError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class CartonDimension:
    box_number: int
    weight: int | float | None = None
    length: int | float | None = None
    width: int | float | None = None
    height: int | float | None = None


@dataclass
class ParsedDnkBoxContents:
    shipment_id: str
    rows: list[ContentRow] = field(default_factory=list)
    dimensions: list[CartonDimension] = field(default_factory=list)


@dataclass(frozen=True)
class FbaBoxContentsDnkCompareResult:
    file_bytes: bytes
    filename: str
    row_count: int
    box_count: int
    upc_count: int
    total_qty: int | float
    remapped_count: int
    added_count: int
    added_qty: int | float
    removed_count: int
    removed_qty: int | float
    last_box: int
    shipment_id: str


def sanitize_download_filename(shipment_id: str | None) -> str:
    """Return ``{shipment_id} Box Contents - corrected.xlsx``."""
    cleaned = (shipment_id or "").replace('"', "").replace("\r", "").replace("\n", "").strip()
    cleaned = re.sub(r'[<>:"/\\|?*]', "", cleaned).strip(" .")
    if not cleaned:
        return DEFAULT_OUTPUT_FILENAME
    return f"{cleaned} Box Contents - corrected.xlsx"


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return str(value).rstrip("0").rstrip(".") if "." in str(value) else str(value)
    return str(value).strip()


def _as_number(value: object) -> int | float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return int(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer() and abs(value) < 2**53:
            return int(value)
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    if not math.isfinite(num):
        return None
    if num.is_integer() and abs(num) < 2**53:
        return int(num)
    return float(num)


def _as_qty(value: object) -> int | float:
    num = _as_number(value)
    if num is None:
        raise FbaBoxContentsDnkCompareError("Quantity must be a number.")
    return num


def _shipment_from_filename(filename: str | None) -> str:
    if not filename:
        return ""
    stem = Path(filename).stem
    for suffix in (
        " Box Contents - corrected",
        " Box Contents",
        " Output",
        " - corrected",
    ):
        if stem.lower().endswith(suffix.lower()):
            stem = stem[: -len(suffix)].rstrip()
            break
    return stem.strip()


def parse_dnk_box_contents(
    content: bytes,
    *,
    filename: str | None = None,
) -> ParsedDnkBoxContents:
    """Parse a DNK Box Contents workbook (Contents + Dimensions)."""
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FbaBoxContentsDnkCompareError("Box Contents file is too large (max 15 MB).")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise FbaBoxContentsDnkCompareError(
            "Could not read the Box Contents Excel file."
        ) from exc

    sheet_name = None
    for name in workbook.sheetnames:
        if name.strip().lower() == CONTENTS_SHEET.lower():
            sheet_name = name
            break
    if sheet_name is None:
        raise FbaBoxContentsDnkCompareError(
            f'Box Contents file must include a "{CONTENTS_SHEET}" sheet.'
        )

    sheet = workbook[sheet_name]
    h1 = _cell_text(sheet.cell(1, 1).value).lower()
    h2 = _cell_text(sheet.cell(1, 2).value).lower()
    h3 = _cell_text(sheet.cell(1, 3).value).lower()
    # DNK order: Box Number | UPC | Quantity
    if "box" in h1 and "upc" in h2 and h3 in ("qty", "quantity"):
        box_col, upc_col, qty_col = 1, 2, 3
    # Also accept Upload Compare order: UPC | Box Number | QTY
    elif "upc" in h1 and "box" in h2 and h3 in ("qty", "quantity"):
        upc_col, box_col, qty_col = 1, 2, 3
    else:
        raise FbaBoxContentsDnkCompareError(
            'Contents sheet must have columns "Box Number", "UPC", "Quantity" '
            "(or UPC / Box Number / QTY)."
        )

    rows: list[ContentRow] = []
    for r in range(2, sheet.max_row + 1):
        upc = _cell_text(sheet.cell(r, upc_col).value)
        box = _as_number(sheet.cell(r, box_col).value)
        qty_raw = sheet.cell(r, qty_col).value
        if not upc and box is None and qty_raw is None:
            continue
        if not upc:
            continue
        if box is None:
            raise FbaBoxContentsDnkCompareError(f"Contents row {r} is missing a Box Number.")
        rows.append(
            ContentRow(identifier=upc, box_number=int(box), qty=_as_qty(qty_raw))
        )

    if not rows:
        raise FbaBoxContentsDnkCompareError("Box Contents file has no Contents rows.")

    shipment_id = _shipment_from_filename(filename)
    dimensions: list[CartonDimension] = []
    if DIMENSIONS_SHEET in workbook.sheetnames:
        dim = workbook[DIMENSIONS_SHEET]
        # Detect header row (Order / Carton / Weight…)
        header_row = None
        cols: dict[str, int] = {}
        for r in range(1, min(6, dim.max_row + 1)):
            labels = {
                c: _cell_text(dim.cell(r, c).value).lower()
                for c in range(1, min(10, dim.max_column + 1))
            }
            carton_cols = [c for c, v in labels.items() if v in ("carton", "box #", "box number", "box")]
            weight_cols = [c for c, v in labels.items() if v == "weight"]
            if carton_cols and weight_cols:
                header_row = r
                cols["carton"] = carton_cols[0]
                cols["weight"] = weight_cols[0]
                for key, aliases in (
                    ("order", ("order",)),
                    ("length", ("length",)),
                    ("width", ("width",)),
                    ("height", ("height",)),
                ):
                    found = [c for c, v in labels.items() if v in aliases]
                    if found:
                        cols[key] = found[0]
                break

        if header_row is not None:
            for r in range(header_row + 1, dim.max_row + 1):
                if "order" in cols:
                    order_text = _cell_text(dim.cell(r, cols["order"]).value)
                    if order_text and not shipment_id:
                        shipment_id = order_text
                    elif order_text and order_text.upper().startswith("FBA"):
                        shipment_id = order_text
                box = _as_number(dim.cell(r, cols["carton"]).value)
                if box is None:
                    continue
                dimensions.append(
                    CartonDimension(
                        box_number=int(box),
                        weight=_as_number(dim.cell(r, cols["weight"]).value)
                        if "weight" in cols
                        else None,
                        length=_as_number(dim.cell(r, cols["length"]).value)
                        if "length" in cols
                        else None,
                        width=_as_number(dim.cell(r, cols["width"]).value)
                        if "width" in cols
                        else None,
                        height=_as_number(dim.cell(r, cols["height"]).value)
                        if "height" in cols
                        else None,
                    )
                )

    close = getattr(workbook, "close", None)
    if callable(close):
        close()

    return ParsedDnkBoxContents(
        shipment_id=shipment_id,
        rows=rows,
        dimensions=dimensions,
    )


def parse_manifest(content: bytes) -> ParsedAmz:
    """Parse the Amazon Seller Central manifest (Box packing information)."""
    try:
        return parse_amz_upload(content)
    except Exception as exc:
        # Re-wrap compare errors / generic failures with DNK-facing wording.
        message = str(exc) or "Could not read the manifest Excel file."
        message = (
            message.replace("AMZ upload", "manifest")
            .replace("AMZ ", "manifest ")
            .replace('"AMZ', '"manifest')
        )
        raise FbaBoxContentsDnkCompareError(message) from exc


def _set_text_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str | None,
    *,
    bold: bool = False,
    align_left: bool = False,
    align_center: bool = False,
    fill: PatternFill | None = None,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align_center:
        cell.alignment = _CENTER
    elif align_left:
        cell.alignment = _LEFT
    if fill is not None:
        cell.fill = fill
    if value is None:
        cell.value = None
        return
    cell.value = value
    cell.data_type = "s"


def _set_number_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | float | None,
    *,
    bold: bool = False,
    align_center: bool = False,
    number_format: str = "General",
    fill: PatternFill | None = None,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = number_format
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align_center:
        cell.alignment = _CENTER
    if fill is not None:
        cell.fill = fill
    if value is None or isinstance(value, bool):
        cell.value = None
        return
    if isinstance(value, int):
        cell.value = int(value)
    elif isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            cell.value = int(value)
        else:
            cell.value = float(value)
    else:
        raise TypeError(f"Expected a number, got {type(value).__name__}")


def _number_format_for_dim(value: int | float | None, *, weight: bool = False) -> str:
    if value is None:
        return "General"
    if weight or isinstance(value, float):
        return _WEIGHT_NUMBER_FORMAT
    return _DIM_INT_NUMBER_FORMAT


def _write_last_box_summary(
    sheet: Worksheet,
    shipment_id: str,
    last_box: int,
    added: Sequence,
) -> None:
    """Write the E–G summary of units appended to the last box."""
    last_box_adds = [a for a in added if a.box_number == last_box]
    if not last_box_adds:
        return

    _set_text_cell(sheet, 1, _SUMMARY_COL, shipment_id or "")
    _set_text_cell(sheet, 2, _SUMMARY_COL, "Box Number")
    _set_text_cell(sheet, 2, _SUMMARY_COL + 1, "UPC")
    _set_text_cell(sheet, 2, _SUMMARY_COL + 2, "QTY")

    for index, item in enumerate(last_box_adds):
        excel_row = 3 + index
        _set_number_cell(
            sheet,
            excel_row,
            _SUMMARY_COL,
            item.box_number,
            fill=_SUMMARY_FILL,
        )
        # Always text so 12-digit UPCs never render as scientific notation.
        _set_text_cell(
            sheet,
            excel_row,
            _SUMMARY_COL + 1,
            item.identifier,
            fill=_SUMMARY_FILL,
        )
        _set_number_cell(
            sheet,
            excel_row,
            _SUMMARY_COL + 2,
            item.qty,
            fill=_SUMMARY_FILL,
        )

    sheet.column_dimensions["E"].width = 11.44
    sheet.column_dimensions["F"].width = 13.66
    sheet.column_dimensions["G"].width = 8.33
    sheet.column_dimensions["H"].width = 3.0


def _write_contents(
    sheet: Worksheet,
    rows: Sequence[ContentRow],
    *,
    shipment_id: str,
    last_box: int,
    added: Sequence,
    box_numbers: Sequence[int] = (),
) -> int | float:
    sheet.sheet_view.showGridLines = True
    _set_text_cell(sheet, 1, 1, "Box Number", bold=True, align_center=True)
    _set_text_cell(sheet, 1, 2, "UPC", bold=True, align_center=True)
    _set_text_cell(sheet, 1, 3, "Quantity", bold=True, align_center=True)

    last_box_adds = [a for a in added if a.box_number == last_box]
    has_summary = bool(last_box_adds)
    if has_summary:
        _write_last_box_summary(sheet, shipment_id, last_box, added)

    pivot_start = PIVOT_START_COL_WITH_SUMMARY if has_summary else PIVOT_START_COL_DEFAULT

    total_qty: int | float = 0
    for offset, row in enumerate(rows):
        excel_row = _LINE_START_ROW + offset
        # Both Old-SKU remaps and last-box top-ups are yellow (sample corrected file).
        fill = _YELLOW if (row.remapped or row.added) else None
        _set_number_cell(
            sheet,
            excel_row,
            1,
            row.box_number,
            align_center=True,
            fill=fill,
        )
        # Keep UPC / Old SKU as text in column B (matches DNK + corrected sample).
        _set_text_cell(
            sheet,
            excel_row,
            2,
            row.identifier,
            align_center=True,
            fill=fill,
        )
        _set_number_cell(
            sheet,
            excel_row,
            3,
            row.qty,
            align_center=True,
            number_format=_QTY_NUMBER_FORMAT,
            fill=fill,
        )
        total_qty += row.qty

    boxes = sorted({row.box_number for row in rows} | set(box_numbers))
    counts: dict[tuple[str, int], int | float] = {}
    identifiers: set[str] = set()
    for row in rows:
        identifiers.add(row.identifier)
        key = (row.identifier, row.box_number)
        counts[key] = counts.get(key, 0) + row.qty

    sorted_ids = sorted(identifiers)
    grand_total_col = pivot_start + 1 + len(boxes)

    _set_text_cell(sheet, 1, pivot_start, "Sum of Quantity")
    _set_text_cell(sheet, 1, pivot_start + 1, "Column Labels")
    _set_text_cell(sheet, 2, pivot_start, "Row Labels")
    for offset, box in enumerate(boxes):
        _set_number_cell(sheet, 2, pivot_start + 1 + offset, box)
    _set_text_cell(sheet, 2, grand_total_col, "Grand Total")

    column_totals: list[int | float] = [0] * len(boxes)
    current_row = 3
    for ident in sorted_ids:
        _set_text_cell(sheet, current_row, pivot_start, ident, align_left=True)
        row_total: int | float = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((ident, box), 0)
            if qty:
                _set_number_cell(
                    sheet,
                    current_row,
                    pivot_start + 1 + offset,
                    qty,
                    number_format=_QTY_NUMBER_FORMAT,
                )
                row_total += qty
                column_totals[offset] += qty
        _set_number_cell(
            sheet,
            current_row,
            grand_total_col,
            row_total,
            number_format=_QTY_NUMBER_FORMAT,
        )
        current_row += 1

    _set_text_cell(sheet, current_row, pivot_start, "Grand Total", align_left=True)
    for offset, qty in enumerate(column_totals):
        _set_number_cell(
            sheet,
            current_row,
            pivot_start + 1 + offset,
            qty,
            number_format=_QTY_NUMBER_FORMAT,
        )
    _set_number_cell(
        sheet,
        current_row,
        grand_total_col,
        total_qty,
        number_format=_QTY_NUMBER_FORMAT,
    )

    sheet.column_dimensions["A"].width = 11.44
    sheet.column_dimensions["B"].width = 13.66
    sheet.column_dimensions["C"].width = 8.33
    sheet.column_dimensions[get_column_letter(pivot_start)].width = 14.89
    sheet.column_dimensions[get_column_letter(pivot_start + 1)].width = 15.55
    for column in range(pivot_start + 2, grand_total_col):
        sheet.column_dimensions[get_column_letter(column)].width = 3.0
    sheet.column_dimensions[get_column_letter(grand_total_col)].width = 10.78
    return total_qty


def _write_dimensions(
    sheet: Worksheet,
    shipment_id: str,
    dimensions: Sequence[CartonDimension],
) -> None:
    sheet.sheet_view.showGridLines = True
    sheet.merge_cells("A1:G1")
    _set_text_cell(sheet, 1, 1, "Carton Dimensions", bold=True)
    sheet["A1"].font = Font(name="Arial", size=12, bold=True)

    headers = ("Order", "Carton", "Weight", "Length", "Width", "Height")
    for column, header in enumerate(headers, start=1):
        _set_text_cell(sheet, 2, column, header, bold=True)

    if not dimensions:
        sheet.column_dimensions["A"].width = 20.55
        sheet.column_dimensions["B"].width = 8.33
        return

    first_row = 3
    last_row = 2 + len(dimensions)
    if last_row > first_row:
        sheet.merge_cells(start_row=first_row, start_column=1, end_row=last_row, end_column=1)

    for offset, dim in enumerate(dimensions):
        row = first_row + offset
        if offset == 0:
            _set_text_cell(sheet, row, 1, shipment_id, align_left=True)
        _set_number_cell(sheet, row, 2, dim.box_number)
        _set_number_cell(
            sheet,
            row,
            3,
            dim.weight,
            number_format=_number_format_for_dim(dim.weight, weight=True),
        )
        _set_number_cell(
            sheet,
            row,
            4,
            dim.length,
            number_format=_number_format_for_dim(dim.length),
        )
        _set_number_cell(
            sheet,
            row,
            5,
            dim.width,
            number_format=_number_format_for_dim(dim.width),
        )
        _set_number_cell(
            sheet,
            row,
            6,
            dim.height,
            number_format=_number_format_for_dim(dim.height),
        )

    sheet.column_dimensions["A"].width = 20.55
    sheet.column_dimensions["B"].width = 8.33
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions["D"].width = 10.0
    sheet.column_dimensions["E"].width = 10.0
    sheet.column_dimensions["F"].width = 10.0


def build_corrected_workbook(
    rows: Sequence[ContentRow],
    dimensions: Sequence[CartonDimension],
    *,
    shipment_id: str,
    last_box: int,
    added: Sequence,
    box_numbers: Sequence[int] = (),
) -> bytes:
    workbook = Workbook()
    contents = workbook.active
    contents.title = CONTENTS_SHEET
    dimensions_sheet = workbook.create_sheet(DIMENSIONS_SHEET)
    _write_contents(
        contents,
        rows,
        shipment_id=shipment_id,
        last_box=last_box,
        added=added,
        box_numbers=box_numbers,
    )
    _write_dimensions(dimensions_sheet, shipment_id, dimensions)
    dimensions_sheet.sheet_view.tabSelected = False
    contents.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_dnk_box_contents_compare(
    box_contents: bytes,
    manifest: bytes,
    catalog_rows: Sequence[Mapping[str, object]],
    *,
    box_contents_filename: str | None = None,
    manifest_filename: str | None = None,
) -> FbaBoxContentsDnkCompareResult:
    """Compare DNK Box Contents to the manifest and return a corrected workbook."""
    _ = manifest_filename
    parsed_box = parse_dnk_box_contents(box_contents, filename=box_contents_filename)
    parsed_manifest = parse_manifest(manifest)

    shipment_id = (
        parsed_manifest.shipment_id
        or parsed_box.shipment_id
        or _shipment_from_filename(box_contents_filename)
    )

    amz_ids = {row.sku_id for row in parsed_manifest.skus}
    upc_to_old = build_upc_to_old_sku_map(catalog_rows, amz_ids=amz_ids)
    remapped_rows = apply_old_sku_remaps(parsed_box, parsed_manifest, upc_to_old)
    reconciled = reconcile_to_amz(remapped_rows, parsed_manifest, upc_to_old)

    file_bytes = build_corrected_workbook(
        reconciled.rows,
        parsed_box.dimensions,
        shipment_id=shipment_id,
        last_box=reconciled.last_box,
        added=reconciled.added,
        box_numbers=reconciled.box_numbers,
    )

    total_qty: int | float = 0
    for row in reconciled.rows:
        total_qty += row.qty
    added_qty: int | float = 0
    for adjustment in reconciled.added:
        added_qty += adjustment.qty
    removed_qty: int | float = 0
    for adjustment in reconciled.removed:
        removed_qty += adjustment.qty
    remapped_count = sum(1 for row in reconciled.rows if row.remapped)
    upc_count = len({row.identifier for row in reconciled.rows})

    return FbaBoxContentsDnkCompareResult(
        file_bytes=file_bytes,
        filename=sanitize_download_filename(shipment_id),
        row_count=len(reconciled.rows),
        box_count=len(reconciled.box_numbers),
        upc_count=upc_count,
        total_qty=total_qty,
        remapped_count=remapped_count,
        added_count=len(reconciled.added),
        added_qty=added_qty,
        removed_count=len(reconciled.removed),
        removed_qty=removed_qty,
        last_box=reconciled.last_box,
        shipment_id=shipment_id,
    )
