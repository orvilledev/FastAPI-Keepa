"""Compare FBA Box Contents Output vs Amazon upload template → Result workbook.

Takes:
  - Output (``Box Contents`` + ``Dimensions`` from FBA Box Contents tools)
  - Amazon ``Upload file - AMZ`` (``Box packing information`` sheet)

The AMZ upload file is the source of truth: the Result Box Contents is
reconciled so the Sum-of-QTY pivot matches the AMZ expected quantities exactly.

Produces a Result workbook in the same Output format:
  - Remap Output UPCs to Old SKUs (yellow highlight) when the AMZ file uses
    that Old SKU (via the Old SKUs catalog UPC ↔ Old SKU mapping)
  - Drop UPCs/units the AMZ file does not expect, and top up shortfalls in the
    last box (green highlight) so every pivot total equals the AMZ quantity
  - Rebuild the Sum-of-QTY pivot with remapped identifiers, keeping every
    original box column so the layout still lines up with Dimensions
  - Copy Dimensions unchanged
  - Add ``added units`` / ``removed units`` sheets for the reconciliation, plus a
    ``missing items`` sheet when SKU/UPC or quantity discrepancies exist
"""
from __future__ import annotations

import io
import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_CALIBRI = Font(name="Calibri", size=11)
_CALIBRI_BOLD = Font(name="Calibri", size=11, bold=True)
_LEFT = Alignment(horizontal="left", vertical="center")
_YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

BOX_CONTENTS_SHEET = "Box Contents"
DIMENSIONS_SHEET = "Dimensions"
ADDED_UNITS_SHEET = "added units"
REMOVED_UNITS_SHEET = "removed units"
MISSING_ITEMS_SHEET = "missing items"
DEFAULT_RESULT_FILENAME = "FBA Result.xlsx"
PIVOT_START_COL = 7  # column G
_SKU_SUFFIX_RE = re.compile(r"-FNSKU$", re.IGNORECASE)
_QTY_EPSILON = 0.0001

NOT_IN_AMZ_REASON = "Not expected in AMZ upload"
SURPLUS_REASON = "Surplus vs AMZ expected quantity"
MISSING_UPC_REASON = "UPC absent from Output"
SHORTFALL_REASON = "Short vs AMZ expected quantity"


class FbaUploadCompareError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class ContentRow:
    identifier: str
    box_number: int
    qty: int | float
    remapped: bool = False
    added: bool = False


@dataclass
class _WorkRow:
    """Mutable Box Contents row used while reconciling against the AMZ file."""

    identifier: str
    box_number: int
    qty: int | float
    remapped: bool = False
    added: bool = False


@dataclass(frozen=True)
class DimensionRow:
    weight: str
    length: int | float | None
    width: int | float | None
    height: int | float | None


@dataclass(frozen=True)
class AmzSkuRow:
    sku_id: str
    expected_qty: int | float
    product_title: str = ""


@dataclass(frozen=True)
class AdjustmentRow:
    """One unit movement applied while reconciling the Output to the AMZ file."""

    identifier: str
    box_number: int
    qty: int | float
    reason: str = ""


@dataclass
class Reconciliation:
    rows: list[ContentRow] = field(default_factory=list)
    box_numbers: list[int] = field(default_factory=list)
    last_box: int = 0
    added: list[AdjustmentRow] = field(default_factory=list)
    removed: list[AdjustmentRow] = field(default_factory=list)


@dataclass
class MissingItem:
    issue: str
    output_id: str = ""
    amz_sku: str = ""
    output_qty: int | float | None = None
    amz_qty: int | float | None = None
    notes: str = ""


@dataclass
class ParsedOutput:
    rows: list[ContentRow] = field(default_factory=list)
    dimensions: list[DimensionRow] = field(default_factory=list)
    shipment_hint: str = ""


@dataclass
class ParsedAmz:
    skus: list[AmzSkuRow] = field(default_factory=list)
    shipment_id: str = ""
    total_box_count: int | None = None


@dataclass(frozen=True)
class FbaUploadCompareResult:
    file_bytes: bytes
    filename: str
    row_count: int
    box_count: int
    upc_count: int
    total_qty: int | float
    remapped_count: int
    missing_count: int
    shipment_id: str
    added_count: int = 0
    added_qty: int | float = 0
    removed_count: int = 0
    removed_qty: int | float = 0
    last_box: int = 0


def sanitize_result_filename(
    output_name: str | None,
    amz_name: str | None = None,
    fallback: str = DEFAULT_RESULT_FILENAME,
) -> str:
    """Derive ``{shipment} Result.xlsx`` from the Output or AMZ filename."""
    for candidate in (output_name, amz_name):
        cleaned = (candidate or "").replace('"', "").replace("\r", "").replace("\n", "").strip()
        base = Path(cleaned).name if cleaned else ""
        stem = Path(base).stem if base else ""
        for suffix in (" Output", " Upload file - AMZ", " Result"):
            if stem.lower().endswith(suffix.lower()):
                stem = stem[: -len(suffix)].rstrip()
                break
        stem = re.sub(r'[<>:"/\\|?*]', "", stem).strip(" .")
        if stem:
            return f"{stem} Result.xlsx"
    return fallback


def _cell_text(value: object) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return format(value, ".15g").strip()
    return str(value).strip()


def _weight_as_text(value: object) -> str:
    """Preserve Output weight text (including leading spaces) when present."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value))
        return format(value, ".15g")
    return str(value)


def _as_number(value: object) -> int | float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return int(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer() and abs(value) < 2**53:
            return int(value)
        return float(value)
    text = _cell_text(value).replace(",", "")
    if not text:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    if not math.isfinite(num):
        return None
    if num.is_integer() and abs(num) < 2**53:
        return int(num)
    return num


def _as_qty(value: object) -> int | float:
    num = _as_number(value)
    if num is None:
        return 0
    return num


def _normalize_sku_id(value: object) -> str:
    text = _cell_text(value)
    if not text:
        return ""
    return _SKU_SUFFIX_RE.sub("", text).strip()


def _set_text_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: str | None,
    *,
    bold: bool = False,
    align_left: bool = False,
    fill: PatternFill | None = None,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if align_left:
        cell.alignment = _LEFT
    if fill is not None:
        cell.fill = fill
    if value is None or value == "":
        cell.value = None
    else:
        cell.value = str(value)
        cell.data_type = "s"


def _set_number_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: int | float | None,
    *,
    bold: bool = False,
) -> None:
    cell = sheet.cell(row=row, column=column)
    cell.number_format = "General"
    cell.font = _CALIBRI_BOLD if bold else _CALIBRI
    if value is None or isinstance(value, bool):
        cell.value = None
        return
    if isinstance(value, int):
        cell.value = int(value)
    elif isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            cell.value = int(value)
        else:
            cell.value = float(value)
    else:
        raise TypeError(f"Expected a number, got {type(value).__name__}")


def parse_output_workbook(content: bytes) -> ParsedOutput:
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FbaUploadCompareError("Output file is too large (max 15 MB).")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise FbaUploadCompareError("Could not read the Output Excel file.") from exc

    if BOX_CONTENTS_SHEET not in workbook.sheetnames:
        raise FbaUploadCompareError(
            f'Output file must include a "{BOX_CONTENTS_SHEET}" sheet.'
        )

    sheet = workbook[BOX_CONTENTS_SHEET]
    header_upc = _cell_text(sheet.cell(1, 1).value).lower()
    header_box = _cell_text(sheet.cell(1, 2).value).lower()
    header_qty = _cell_text(sheet.cell(1, 3).value).lower()
    if header_upc != "upc" or "box" not in header_box or header_qty not in ("qty", "quantity"):
        raise FbaUploadCompareError(
            'Output "Box Contents" must have columns UPC, Box Number, QTY.'
        )

    rows: list[ContentRow] = []
    for r in range(2, sheet.max_row + 1):
        upc = _cell_text(sheet.cell(r, 1).value)
        box = _as_number(sheet.cell(r, 2).value)
        qty_raw = sheet.cell(r, 3).value
        if not upc and box is None and qty_raw is None:
            continue
        if not upc:
            continue
        if box is None:
            raise FbaUploadCompareError(f"Output row {r} is missing a Box Number.")
        box_number = int(box)
        rows.append(ContentRow(identifier=upc, box_number=box_number, qty=_as_qty(qty_raw)))

    if not rows:
        raise FbaUploadCompareError("Output file has no Box Contents rows.")

    dimensions: list[DimensionRow] = []
    if DIMENSIONS_SHEET in workbook.sheetnames:
        dim = workbook[DIMENSIONS_SHEET]
        for r in range(2, dim.max_row + 1):
            weight = _weight_as_text(dim.cell(r, 1).value)
            length = _as_number(dim.cell(r, 2).value)
            width = _as_number(dim.cell(r, 3).value)
            height = _as_number(dim.cell(r, 4).value)
            if not weight.strip() and length is None and width is None and height is None:
                continue
            dimensions.append(
                DimensionRow(weight=weight, length=length, width=width, height=height)
            )

    return ParsedOutput(rows=rows, dimensions=dimensions)


def parse_amz_upload(content: bytes) -> ParsedAmz:
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FbaUploadCompareError("AMZ upload file is too large (max 15 MB).")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise FbaUploadCompareError("Could not read the AMZ upload Excel file.") from exc

    sheet_name = None
    for name in workbook.sheetnames:
        if name.strip().lower() == "box packing information":
            sheet_name = name
            break
    if sheet_name is None:
        raise FbaUploadCompareError(
            'AMZ upload file must include a "Box packing information" sheet.'
        )

    sheet = workbook[sheet_name]
    shipment_id = ""
    total_box_count: int | None = None
    for r in range(1, min(6, sheet.max_row + 1)):
        a = _cell_text(sheet.cell(r, 1).value)
        if a.lower().startswith("shipment:"):
            shipment_id = a.split(":", 1)[1].strip() or _cell_text(sheet.cell(r, 2).value)
        # Total box count sits near column M (13) on Amazon templates
        for c in range(1, min(20, sheet.max_column + 1)):
            label = _cell_text(sheet.cell(r, c).value).lower()
            if "total box count" in label:
                for look in range(c + 1, min(c + 6, sheet.max_column + 1)):
                    num = _as_number(sheet.cell(r, look).value)
                    if num is not None:
                        total_box_count = int(num)
                        break

    # Find header row with SKU / Expected quantity
    header_row = None
    sku_col = expected_col = title_col = None
    for r in range(1, min(12, sheet.max_row + 1)):
        labels = {
            c: _cell_text(sheet.cell(r, c).value).lower()
            for c in range(1, min(20, sheet.max_column + 1))
        }
        sku_candidates = [c for c, v in labels.items() if v == "sku"]
        exp_candidates = [
            c for c, v in labels.items() if v in ("expected quantity", "expected qty")
        ]
        title_candidates = [
            c for c, v in labels.items() if "product title" in v or v == "title"
        ]
        if sku_candidates and exp_candidates:
            header_row = r
            sku_col = sku_candidates[0]
            expected_col = exp_candidates[0]
            title_col = title_candidates[0] if title_candidates else None
            break

    if header_row is None or sku_col is None or expected_col is None:
        raise FbaUploadCompareError(
            'AMZ "Box packing information" must include SKU and Expected quantity columns.'
        )

    skus: list[AmzSkuRow] = []
    for r in range(header_row + 1, sheet.max_row + 1):
        raw_sku = _cell_text(sheet.cell(r, sku_col).value)
        if not raw_sku:
            continue
        lower = raw_sku.lower()
        if lower.startswith("box ") or lower.startswith("name of box") or lower.startswith("box weight"):
            break
        sku_id = _normalize_sku_id(raw_sku)
        if not sku_id:
            continue
        title = _cell_text(sheet.cell(r, title_col).value) if title_col else ""
        skus.append(
            AmzSkuRow(
                sku_id=sku_id,
                expected_qty=_as_qty(sheet.cell(r, expected_col).value),
                product_title=title,
            )
        )

    if not skus:
        raise FbaUploadCompareError("AMZ upload file has no SKU rows.")

    return ParsedAmz(
        skus=skus,
        shipment_id=shipment_id,
        total_box_count=total_box_count,
    )


def build_upc_to_old_sku_map(
    catalog_rows: Sequence[Mapping[str, object]],
    *,
    amz_ids: set[str] | None = None,
) -> dict[str, str]:
    """Map UPC Code → Old SKU, preferring real Old SKUs used in the AMZ file.

    The catalog often stores both an identity row (``old_sku == upc_code``) and a
    true Old SKU row (e.g. ``9990262``) for the same UPC. Identity rows must not
    win the mapping or remaps and yellow highlights break.
    """
    amz_ids = amz_ids or set()
    candidates: dict[str, list[str]] = defaultdict(list)
    for row in catalog_rows:
        upc = _cell_text(row.get("upc_code"))
        old = _cell_text(row.get("old_sku"))
        if not upc or not old:
            continue
        if old not in candidates[upc]:
            candidates[upc].append(old)

    mapping: dict[str, str] = {}
    for upc, olds in candidates.items():
        in_amz = [old for old in olds if old != upc and old in amz_ids]
        if in_amz:
            mapping[upc] = in_amz[0]
            continue
        non_identity = [old for old in olds if old != upc]
        if non_identity:
            mapping[upc] = non_identity[0]
    return mapping


def apply_old_sku_remaps(
    output: ParsedOutput,
    amz: ParsedAmz,
    upc_to_old_sku: Mapping[str, str],
) -> list[ContentRow]:
    """Replace Output UPCs with Old SKUs when those Old SKUs appear in AMZ."""
    amz_ids = {row.sku_id for row in amz.skus}
    remapped_rows: list[ContentRow] = []
    for row in output.rows:
        old_sku = upc_to_old_sku.get(row.identifier)
        if old_sku and old_sku in amz_ids and old_sku != row.identifier:
            remapped_rows.append(
                ContentRow(
                    identifier=old_sku,
                    box_number=row.box_number,
                    qty=row.qty,
                    remapped=True,
                )
            )
        else:
            remapped_rows.append(row)
    return remapped_rows


def find_missing_items(
    output_rows: Sequence[ContentRow],
    amz: ParsedAmz,
    upc_to_old_sku: Mapping[str, str],
) -> list[MissingItem]:
    """Compare identifiers and quantities between Output and AMZ expected qty.

    This reports the Output as submitted, before ``reconcile_to_amz`` rewrites it,
    so the Result keeps a record of what did not line up.
    """
    old_to_upc = {old: upc for upc, old in upc_to_old_sku.items()}

    output_totals: dict[str, float] = defaultdict(float)
    for row in output_rows:
        output_totals[row.identifier] += float(row.qty)

    amz_totals: dict[str, float] = defaultdict(float)
    amz_titles: dict[str, str] = {}
    for row in amz.skus:
        amz_totals[row.sku_id] += float(row.expected_qty)
        if row.product_title and row.sku_id not in amz_titles:
            amz_titles[row.sku_id] = row.product_title

    def resolve_amz_key(output_id: str) -> str | None:
        if output_id in amz_totals:
            return output_id
        old = upc_to_old_sku.get(output_id)
        if old and old in amz_totals:
            return old
        upc = old_to_upc.get(output_id)
        if upc and upc in amz_totals:
            return upc
        return None

    def resolve_output_key(amz_id: str) -> str | None:
        if amz_id in output_totals:
            return amz_id
        upc = old_to_upc.get(amz_id)
        if upc and upc in output_totals:
            return upc
        old = upc_to_old_sku.get(amz_id)
        if old and old in output_totals:
            return old
        return None

    missing: list[MissingItem] = []
    matched_amz: set[str] = set()

    for out_id, out_qty in sorted(output_totals.items()):
        amz_key = resolve_amz_key(out_id)
        if amz_key is None:
            missing.append(
                MissingItem(
                    issue="Missing in AMZ upload",
                    output_id=out_id,
                    output_qty=out_qty,
                    notes="Present in Output but not found in AMZ upload SKUs",
                )
            )
            continue
        matched_amz.add(amz_key)
        amz_qty = amz_totals[amz_key]
        if abs(amz_qty - out_qty) > 0.001:
            missing.append(
                MissingItem(
                    issue="Quantity mismatch",
                    output_id=out_id,
                    amz_sku=amz_key,
                    output_qty=out_qty,
                    amz_qty=amz_qty,
                    notes=amz_titles.get(amz_key, ""),
                )
            )

    for amz_id, amz_qty in sorted(amz_totals.items()):
        if amz_id in matched_amz:
            continue
        out_key = resolve_output_key(amz_id)
        if out_key is not None:
            # Already handled via output side (including qty mismatch)
            continue
        missing.append(
            MissingItem(
                issue="Missing in Output",
                amz_sku=amz_id,
                amz_qty=amz_qty,
                notes=amz_titles.get(amz_id, "")
                or "Present in AMZ upload but not found in Output UPCs",
            )
        )

    return missing


def reconcile_to_amz(
    rows: Sequence[ContentRow],
    amz: ParsedAmz,
    upc_to_old_sku: Mapping[str, str],
) -> Reconciliation:
    """Force Box Contents totals to equal the AMZ expected quantities.

    The AMZ upload is the source of truth, so anything it does not expect is
    removed (surplus units come off the highest box numbers first) and every
    shortfall is topped up in the last box.
    """
    targets: dict[str, int | float] = {}
    for sku in amz.skus:
        targets[sku.sku_id] = targets.get(sku.sku_id, 0) + sku.expected_qty

    old_to_upc = {old: upc for upc, old in upc_to_old_sku.items()}

    def resolve(identifier: str) -> str | None:
        if identifier in targets:
            return identifier
        old = upc_to_old_sku.get(identifier)
        if old and old in targets:
            return old
        upc = old_to_upc.get(identifier)
        if upc and upc in targets:
            return upc
        return None

    # Keep every original box column: reconciling can empty a box completely,
    # and the pivot still has to line up with the AMZ boxes and Dimensions.
    box_numbers = sorted({row.box_number for row in rows})
    if not box_numbers:
        raise FbaUploadCompareError("Output file has no Box Contents rows.")
    last_box = box_numbers[-1]

    work: list[_WorkRow] = []
    grouped: dict[str, list[int]] = defaultdict(list)
    removed: list[AdjustmentRow] = []

    for row in rows:
        key = resolve(row.identifier)
        if key is None:
            removed.append(
                AdjustmentRow(
                    identifier=row.identifier,
                    box_number=row.box_number,
                    qty=row.qty,
                    reason=NOT_IN_AMZ_REASON,
                )
            )
            continue
        grouped[key].append(len(work))
        work.append(
            _WorkRow(
                identifier=key,
                box_number=row.box_number,
                qty=row.qty,
                remapped=row.remapped or key != row.identifier,
            )
        )

    added: list[AdjustmentRow] = []
    for key in sorted(targets):
        target = targets[key]
        indexes = grouped.get(key, [])
        current: int | float = 0
        for index in indexes:
            current += work[index].qty
        delta = target - current
        if abs(delta) <= _QTY_EPSILON:
            continue
        if delta < 0:
            surplus = -delta
            # Trim the tail boxes first so the earlier boxes stay as packed.
            for index in sorted(indexes, key=lambda i: (work[i].box_number, i), reverse=True):
                if surplus <= _QTY_EPSILON:
                    break
                take = min(surplus, work[index].qty)
                if take <= 0:
                    continue
                work[index].qty -= take
                surplus -= take
                removed.append(
                    AdjustmentRow(
                        identifier=key,
                        box_number=work[index].box_number,
                        qty=take,
                        reason=SURPLUS_REASON,
                    )
                )
            continue
        added.append(
            AdjustmentRow(
                identifier=key,
                box_number=last_box,
                qty=delta,
                reason=SHORTFALL_REASON if indexes else MISSING_UPC_REASON,
            )
        )

    kept = [row for row in work if row.qty]
    last_box_rows = {row.identifier: row for row in kept if row.box_number == last_box}
    appended: list[_WorkRow] = []
    for adjustment in added:
        existing = last_box_rows.get(adjustment.identifier)
        if existing is not None:
            existing.qty += adjustment.qty
            existing.added = True
            continue
        new_row = _WorkRow(
            identifier=adjustment.identifier,
            box_number=last_box,
            qty=adjustment.qty,
            added=True,
        )
        last_box_rows[adjustment.identifier] = new_row
        appended.append(new_row)
    kept.extend(appended)

    return Reconciliation(
        rows=[
            ContentRow(
                identifier=row.identifier,
                box_number=row.box_number,
                qty=row.qty,
                remapped=row.remapped,
                added=row.added,
            )
            for row in kept
        ],
        box_numbers=box_numbers,
        last_box=last_box,
        added=added,
        removed=removed,
    )


def _write_box_contents(
    sheet: Worksheet,
    rows: Sequence[ContentRow],
    box_numbers: Sequence[int] = (),
) -> int | float:
    sheet.sheet_view.showGridLines = True
    _set_text_cell(sheet, 1, 1, "UPC", bold=True)
    _set_text_cell(sheet, 1, 2, "Box Number", bold=True)
    _set_text_cell(sheet, 1, 3, "QTY", bold=True)

    total_qty: int | float = 0
    for index, row in enumerate(rows, start=2):
        if row.remapped:
            fill = _YELLOW
        elif row.added:
            fill = _GREEN
        else:
            fill = None
        _set_text_cell(sheet, index, 1, row.identifier, fill=fill)
        _set_number_cell(sheet, index, 2, row.box_number)
        _set_number_cell(sheet, index, 3, row.qty)
        total_qty += row.qty

    boxes = sorted({row.box_number for row in rows} | set(box_numbers))
    counts: dict[tuple[str, int], int | float] = {}
    identifiers: set[str] = set()
    for row in rows:
        identifiers.add(row.identifier)
        key = (row.identifier, row.box_number)
        counts[key] = counts.get(key, 0) + row.qty

    sorted_ids = sorted(identifiers)
    grand_total_col = PIVOT_START_COL + 1 + len(boxes)

    _set_text_cell(sheet, 1, PIVOT_START_COL, "Sum of QTY")
    _set_text_cell(sheet, 1, PIVOT_START_COL + 1, "Column Labels")
    _set_text_cell(sheet, 2, PIVOT_START_COL, "Row Labels")
    for offset, box in enumerate(boxes):
        _set_number_cell(sheet, 2, PIVOT_START_COL + 1 + offset, box)
    _set_text_cell(sheet, 2, grand_total_col, "Grand Total")

    column_totals: list[int | float] = [0] * len(boxes)
    current_row = 3
    for ident in sorted_ids:
        _set_text_cell(sheet, current_row, PIVOT_START_COL, ident, align_left=True)
        row_total: int | float = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((ident, box), 0)
            if qty:
                _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
                row_total += qty
                column_totals[offset] += qty
        _set_number_cell(sheet, current_row, grand_total_col, row_total)
        current_row += 1

    _set_text_cell(sheet, current_row, PIVOT_START_COL, "Grand Total", align_left=True)
    for offset, qty in enumerate(column_totals):
        _set_number_cell(sheet, current_row, PIVOT_START_COL + 1 + offset, qty)
    _set_number_cell(sheet, current_row, grand_total_col, total_qty)

    sheet.column_dimensions["A"].width = 13.11
    sheet.column_dimensions["B"].width = 13.33
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL)].width = 13.11
    sheet.column_dimensions[get_column_letter(PIVOT_START_COL + 1)].width = 15.55
    for column in range(PIVOT_START_COL + 2, grand_total_col):
        sheet.column_dimensions[get_column_letter(column)].width = 3.0
    sheet.column_dimensions[get_column_letter(grand_total_col)].width = 10.78
    return total_qty


def _write_dimensions(sheet: Worksheet, dimensions: Sequence[DimensionRow]) -> None:
    sheet.sheet_view.showGridLines = True
    for column, header in enumerate(("Weight", "Length", "Width", "Height"), start=1):
        _set_text_cell(sheet, 1, column, header)

    for index, dim in enumerate(dimensions, start=2):
        _set_text_cell(sheet, index, 1, dim.weight or None)
        _set_number_cell(sheet, index, 2, dim.length)
        _set_number_cell(sheet, index, 3, dim.width)
        _set_number_cell(sheet, index, 4, dim.height)

    sheet.column_dimensions["A"].width = 10.0
    sheet.column_dimensions["B"].width = 10.0
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions["D"].width = 10.0


def _write_added_units(sheet: Worksheet, added: Sequence[AdjustmentRow]) -> None:
    """Units topped up in the last box to reach the AMZ expected quantities."""
    sheet.sheet_view.showGridLines = True
    for col, header in enumerate(("UPC", "Box Number", "QTY"), start=1):
        _set_text_cell(sheet, 1, col, header, bold=True)

    for index, item in enumerate(added, start=2):
        _set_text_cell(sheet, index, 1, item.identifier)
        _set_number_cell(sheet, index, 2, item.box_number)
        _set_number_cell(sheet, index, 3, item.qty)

    sheet.column_dimensions["A"].width = 15.0
    sheet.column_dimensions["B"].width = 13.33
    sheet.column_dimensions["C"].width = 10.0


def _write_removed_units(sheet: Worksheet, removed: Sequence[AdjustmentRow]) -> None:
    """Units dropped from the Output because the AMZ file does not expect them."""
    sheet.sheet_view.showGridLines = True
    for col, header in enumerate(("UPC", "Box Number", "QTY", "Reason"), start=1):
        _set_text_cell(sheet, 1, col, header, bold=True)

    for index, item in enumerate(removed, start=2):
        _set_text_cell(sheet, index, 1, item.identifier)
        _set_number_cell(sheet, index, 2, item.box_number)
        _set_number_cell(sheet, index, 3, item.qty)
        _set_text_cell(sheet, index, 4, item.reason or None)

    sheet.column_dimensions["A"].width = 15.0
    sheet.column_dimensions["B"].width = 13.33
    sheet.column_dimensions["C"].width = 10.0
    sheet.column_dimensions["D"].width = 32.0


def _write_missing_items(sheet: Worksheet, missing: Sequence[MissingItem]) -> None:
    sheet.sheet_view.showGridLines = True
    headers = (
        "Issue",
        "Output UPC/SKU",
        "AMZ SKU",
        "Output Qty",
        "AMZ Expected Qty",
        "Notes",
    )
    for col, header in enumerate(headers, start=1):
        _set_text_cell(sheet, 1, col, header, bold=True)

    for index, item in enumerate(missing, start=2):
        _set_text_cell(sheet, index, 1, item.issue)
        _set_text_cell(sheet, index, 2, item.output_id or None)
        _set_text_cell(sheet, index, 3, item.amz_sku or None)
        _set_number_cell(sheet, index, 4, item.output_qty)
        _set_number_cell(sheet, index, 5, item.amz_qty)
        _set_text_cell(sheet, index, 6, item.notes or None)

    sheet.column_dimensions["A"].width = 22.0
    sheet.column_dimensions["B"].width = 18.0
    sheet.column_dimensions["C"].width = 18.0
    sheet.column_dimensions["D"].width = 12.0
    sheet.column_dimensions["E"].width = 18.0
    sheet.column_dimensions["F"].width = 48.0


def build_result_workbook(
    remapped_rows: Sequence[ContentRow],
    dimensions: Sequence[DimensionRow],
    missing: Sequence[MissingItem],
    *,
    box_numbers: Sequence[int] = (),
    added: Sequence[AdjustmentRow] = (),
    removed: Sequence[AdjustmentRow] = (),
) -> bytes:
    workbook = Workbook()
    contents = workbook.active
    contents.title = BOX_CONTENTS_SHEET
    dimensions_sheet = workbook.create_sheet(DIMENSIONS_SHEET)
    _write_box_contents(contents, remapped_rows, box_numbers)
    _write_dimensions(dimensions_sheet, dimensions)
    extra_sheets = [dimensions_sheet]
    if added:
        added_sheet = workbook.create_sheet(ADDED_UNITS_SHEET)
        _write_added_units(added_sheet, added)
        extra_sheets.append(added_sheet)
    if removed:
        removed_sheet = workbook.create_sheet(REMOVED_UNITS_SHEET)
        _write_removed_units(removed_sheet, removed)
        extra_sheets.append(removed_sheet)
    if missing:
        missing_sheet = workbook.create_sheet(MISSING_ITEMS_SHEET)
        _write_missing_items(missing_sheet, missing)
        extra_sheets.append(missing_sheet)
    for sheet in extra_sheets:
        sheet.sheet_view.tabSelected = False
    contents.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_fba_upload_compare(
    output_content: bytes,
    amz_content: bytes,
    catalog_rows: Sequence[Mapping[str, object]],
    *,
    output_filename: str | None = None,
    amz_filename: str | None = None,
) -> FbaUploadCompareResult:
    parsed_output = parse_output_workbook(output_content)
    parsed_amz = parse_amz_upload(amz_content)
    amz_ids = {row.sku_id for row in parsed_amz.skus}
    upc_to_old = build_upc_to_old_sku_map(catalog_rows, amz_ids=amz_ids)

    remapped_rows = apply_old_sku_remaps(parsed_output, parsed_amz, upc_to_old)
    # Missing-items compare uses original Output identifiers + catalog links
    missing = find_missing_items(parsed_output.rows, parsed_amz, upc_to_old)
    # The AMZ upload is the source of truth, so the pivot is forced to match it
    reconciled = reconcile_to_amz(remapped_rows, parsed_amz, upc_to_old)

    file_bytes = build_result_workbook(
        reconciled.rows,
        parsed_output.dimensions,
        missing,
        box_numbers=reconciled.box_numbers,
        added=reconciled.added,
        removed=reconciled.removed,
    )

    total_qty: int | float = 0
    for row in reconciled.rows:
        total_qty += row.qty
    added_qty: int | float = 0
    for adjustment in reconciled.added:
        added_qty += adjustment.qty
    removed_qty: int | float = 0
    for adjustment in reconciled.removed:
        removed_qty += adjustment.qty
    remapped_count = len({row.identifier for row in reconciled.rows if row.remapped})
    upc_count = len({row.identifier for row in reconciled.rows})
    result_name = sanitize_result_filename(output_filename, amz_filename)

    return FbaUploadCompareResult(
        file_bytes=file_bytes,
        filename=result_name,
        row_count=len(reconciled.rows),
        box_count=len(reconciled.box_numbers),
        upc_count=upc_count,
        total_qty=total_qty,
        remapped_count=remapped_count,
        missing_count=len(missing),
        shipment_id=parsed_amz.shipment_id,
        added_count=len(reconciled.added),
        added_qty=added_qty,
        removed_count=len(reconciled.removed),
        removed_qty=removed_qty,
        last_box=reconciled.last_box,
    )
