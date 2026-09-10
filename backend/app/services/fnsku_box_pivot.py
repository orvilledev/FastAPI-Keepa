"""Convert an FNSKU + BOX# scan sheet into a box-contents workbook.

Input (header row, case-insensitive): FNSKU | BOX#

Output matches the warehouse Excel workflow:
  - ``scanned data``: msku, FNSKU, BOX#, QTY (QTY is 1 per input row)
  - ``Sheet7``: Excel-style pivot of Sum of QTY by msku (rows) and box (columns)
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.repositories.warehouse_product_repository import merchant_sku_from_catalog_row

_MAX_UPLOAD_BYTES = 15 * 1024 * 1024
_CALIBRI = Font(name="Calibri", size=11)
_TEXT_FORMAT = "@"
_INT_FORMAT = "0"
_FLOAT_FORMAT = "0.##############"
# Excel column width for 35px at the default 7px digit width: (pixels - 5) / 7.
_PIVOT_DATA_COL_WIDTH = (35 - 5) / 7
_PIVOT_LABEL_COL_WIDTH = 19.89

_FNSKU_HEADERS = frozenset({"fnsku"})
_BOX_HEADERS = frozenset({"box#", "box #", "box", "box number", "box no", "boxno"})

PIVOT_SHEET_NAME = "Sheet7"
SCANNED_SHEET_NAME = "scanned data"
OUTPUT_FILENAME = "Output.xlsx"


class FnskuBoxPivotError(ValueError):
    """Raised for user-correctable input problems."""


@dataclass(frozen=True)
class ScanRow:
    fnsku: str
    box: Any


@dataclass(frozen=True)
class ScannedDataRow:
    msku: str
    fnsku: str
    box: Any
    qty: int = 1


@dataclass(frozen=True)
class FnskuBoxPivotResult:
    file_bytes: bytes
    filename: str
    row_count: int
    sku_count: int
    unmatched_count: int
    unmatched_fnskus: tuple[str, ...]


def _normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g").strip()
    return str(value).strip()


def _parse_box(value: object) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        return value
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return int(number)
    return number


def _box_sort_key(box: Any) -> tuple:
    if isinstance(box, int):
        return (0, box, "")
    if isinstance(box, float):
        return (0, box, "")
    return (1, 0, str(box))


def parse_fnsku_box_rows(content: bytes) -> list[ScanRow]:
    if not content:
        raise FnskuBoxPivotError("Uploaded file is empty.")
    if len(content) > _MAX_UPLOAD_BYTES:
        raise FnskuBoxPivotError("File is too large (max 15 MB).")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise FnskuBoxPivotError("Could not read the Excel file. Upload a valid .xlsx workbook.") from exc

    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index = None
    fnsku_at = None
    box_at = None
    for index, row in enumerate(rows):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            header = _normalize_header(cell)
            if header in _FNSKU_HEADERS:
                mapping["fnsku"] = col
            elif header in _BOX_HEADERS:
                mapping["box"] = col
        if "fnsku" in mapping and "box" in mapping:
            header_index = index
            fnsku_at = mapping["fnsku"]
            box_at = mapping["box"]
            break

    if header_index is None or fnsku_at is None or box_at is None:
        raise FnskuBoxPivotError(
            'The file must include "FNSKU" and "BOX#" column headers.'
        )

    parsed: list[ScanRow] = []
    for row in rows[header_index + 1 :]:
        fnsku = _cell_str(row[fnsku_at] if fnsku_at < len(row) else None)
        if not fnsku:
            continue
        box = _parse_box(row[box_at] if box_at < len(row) else None)
        parsed.append(ScanRow(fnsku=fnsku, box=box))

    if not parsed:
        raise FnskuBoxPivotError("No FNSKU rows found under the header.")
    return parsed


def apply_msku_lookup(
    rows: Sequence[ScanRow],
    catalog_by_fnsku: Mapping[str, Mapping[str, Any]],
) -> tuple[list[ScannedDataRow], tuple[str, ...]]:
    scanned: list[ScannedDataRow] = []
    unmatched: list[str] = []
    seen_unmatched: set[str] = set()
    for row in rows:
        catalog_row = catalog_by_fnsku.get(row.fnsku)
        msku = merchant_sku_from_catalog_row(catalog_row) if catalog_row is not None else ""
        if not msku and row.fnsku not in seen_unmatched:
            seen_unmatched.add(row.fnsku)
            unmatched.append(row.fnsku)
        scanned.append(
            ScannedDataRow(msku=msku, fnsku=row.fnsku, box=row.box, qty=1)
        )
    return scanned, tuple(unmatched)


def _set_text_cell(sheet: Worksheet, row: int, column: int, value: Any) -> None:
    """Write a lookup key as Excel text so VLOOKUP/XLOOKUP will not coerce it to a number."""
    cell = sheet.cell(row=row, column=column)
    cell.font = _CALIBRI
    cell.number_format = _TEXT_FORMAT
    if value is None:
        cell.value = None
        return
    text = str(value)
    if not text:
        cell.value = None
        return
    cell.value = text
    cell.data_type = "s"


def _set_number_cell(sheet: Worksheet, row: int, column: int, value: Any) -> None:
    """Write a true Excel number. Non-numeric values fall back to text."""
    cell = sheet.cell(row=row, column=column)
    cell.font = _CALIBRI
    if value is None or isinstance(value, bool):
        cell.value = None
        cell.number_format = _INT_FORMAT
        return
    if isinstance(value, int):
        cell.value = value
        cell.number_format = _INT_FORMAT
        return
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            cell.value = int(value)
            cell.number_format = _INT_FORMAT
        else:
            cell.value = value
            cell.number_format = _FLOAT_FORMAT
        return
    parsed = _parse_box(value)
    if isinstance(parsed, (int, float)):
        _set_number_cell(sheet, row, column, parsed)
        return
    _set_text_cell(sheet, row, column, value)


def _write_scanned_data(sheet: Worksheet, rows: Sequence[ScannedDataRow]) -> None:
    sheet.freeze_panes = "A2"
    for column, header in enumerate(("msku", "FNSKU", "BOX#", "QTY"), start=1):
        _set_text_cell(sheet, 1, column, header)
    for index, row in enumerate(rows, start=2):
        _set_text_cell(sheet, index, 1, row.msku or None)
        _set_text_cell(sheet, index, 2, row.fnsku)
        _set_number_cell(sheet, index, 3, row.box)
        _set_number_cell(sheet, index, 4, row.qty)
    sheet.column_dimensions["A"].width = 19.89
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 13
    sheet.column_dimensions["D"].width = 13
    sheet.column_dimensions["A"].number_format = _TEXT_FORMAT
    sheet.column_dimensions["B"].number_format = _TEXT_FORMAT
    sheet.column_dimensions["C"].number_format = _INT_FORMAT
    sheet.column_dimensions["D"].number_format = _INT_FORMAT


def _write_pivot(sheet: Worksheet, rows: Sequence[ScannedDataRow]) -> None:
    boxes: list[Any] = []
    seen_boxes: set[Any] = set()
    for row in rows:
        if row.box is None:
            continue
        if row.box not in seen_boxes:
            seen_boxes.add(row.box)
            boxes.append(row.box)
    boxes.sort(key=_box_sort_key)

    counts: dict[tuple[str, Any], int] = {}
    mskus: set[str] = set()
    for row in rows:
        msku = row.msku or ""
        if not msku or row.box is None:
            continue
        mskus.add(msku)
        key = (msku, row.box)
        counts[key] = counts.get(key, 0) + row.qty

    sorted_mskus = sorted(mskus)
    grand_total_col = 2 + len(boxes)

    _set_text_cell(sheet, 3, 1, "Sum of QTY")
    _set_text_cell(sheet, 3, 2, "Column Labels")
    _set_text_cell(sheet, 4, 1, "Row Labels")
    for offset, box in enumerate(boxes, start=2):
        _set_number_cell(sheet, 4, offset, box)
    _set_text_cell(sheet, 4, grand_total_col, "Grand Total")

    column_totals = [0] * len(boxes)
    current_row = 5
    for msku in sorted_mskus:
        _set_text_cell(sheet, current_row, 1, msku)
        row_total = 0
        for offset, box in enumerate(boxes):
            qty = counts.get((msku, box), 0)
            if qty:
                _set_number_cell(sheet, current_row, offset + 2, qty)
                row_total += qty
                column_totals[offset] += qty
        if row_total:
            _set_number_cell(sheet, current_row, grand_total_col, row_total)
        current_row += 1

    total_row = current_row
    _set_text_cell(sheet, total_row, 1, "Grand Total")
    grand = 0
    for offset, qty in enumerate(column_totals):
        if qty:
            _set_number_cell(sheet, total_row, offset + 2, qty)
            grand += qty
    if grand:
        _set_number_cell(sheet, total_row, grand_total_col, grand)

    sheet.column_dimensions["A"].width = _PIVOT_LABEL_COL_WIDTH
    for column in range(2, grand_total_col + 1):
        sheet.column_dimensions[get_column_letter(column)].width = _PIVOT_DATA_COL_WIDTH


def build_output_workbook(rows: Sequence[ScannedDataRow]) -> bytes:
    workbook = Workbook()
    pivot = workbook.active
    pivot.title = PIVOT_SHEET_NAME
    scanned = workbook.create_sheet(SCANNED_SHEET_NAME)
    _write_pivot(pivot, rows)
    _write_scanned_data(scanned, rows)
    scanned.sheet_view.tabSelected = False
    pivot.sheet_view.tabSelected = True
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def generate_fnsku_box_pivot(
    content: bytes,
    catalog_by_fnsku: Mapping[str, Mapping[str, Any]],
    filename: str = OUTPUT_FILENAME,
    scan_rows: Sequence[ScanRow] | None = None,
) -> FnskuBoxPivotResult:
    if scan_rows is None:
        scan_rows = parse_fnsku_box_rows(content)
    scanned, unmatched = apply_msku_lookup(scan_rows, catalog_by_fnsku)
    file_bytes = build_output_workbook(scanned)
    sku_count = len({row.msku for row in scanned if row.msku})
    return FnskuBoxPivotResult(
        file_bytes=file_bytes,
        filename=filename or OUTPUT_FILENAME,
        row_count=len(scanned),
        sku_count=sku_count,
        unmatched_count=len(unmatched),
        unmatched_fnskus=unmatched,
    )


def build_template_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    _set_text_cell(sheet, 1, 1, "FNSKU")
    _set_text_cell(sheet, 1, 2, "BOX#")
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["A"].number_format = _TEXT_FORMAT
    sheet.column_dimensions["B"].number_format = _INT_FORMAT
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
