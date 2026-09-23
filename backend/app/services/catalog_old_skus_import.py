"""Parse the Old SKUs workbook (OLD SKU / Vendor Name / UPC Code)."""
from __future__ import annotations

import io
import logging
from datetime import date, datetime
from typing import Any, Dict, List, Sequence, Tuple

from openpyxl import load_workbook

from app.services.catalog_old_skus_headers import HEADERS, REQUIRED_HEADER, SHEET_NAME

logger = logging.getLogger(__name__)


def _cell_str(raw: Any) -> str:
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return "TRUE" if raw else "FALSE"
    if isinstance(raw, datetime):
        if raw.hour == 0 and raw.minute == 0 and raw.second == 0 and raw.microsecond == 0:
            return raw.date().isoformat()
        return raw.isoformat(sep=" ", timespec="seconds")
    if isinstance(raw, date):
        return raw.isoformat()
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if raw.is_integer():
            return str(int(raw))
        return format(raw, ".15g").strip()
    text = str(raw).strip()
    if text.startswith("="):
        return ""
    return text


def _normalize_header(cell: Any) -> str:
    return _cell_str(cell).strip()


def _find_sheet(wb, preferred: str):
    wanted = preferred.strip().upper()
    for name in wb.sheetnames:
        if name.strip().upper() == wanted:
            return wb[name]
    return wb[wb.sheetnames[0]] if wb.sheetnames else None


def _header_index_map(
    header_row: Sequence[Any],
    expected: Sequence[str],
    required: str,
) -> Dict[str, int]:
    found: Dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = _normalize_header(cell)
        if not name or name in found:
            continue
        found[name] = idx

    if required not in found:
        raise ValueError(
            f'Spreadsheet must include a "{required}" column. '
            "Download the template and match the exact headers."
        )

    missing = [h for h in expected if h not in found]
    if missing:
        preview = ", ".join(missing)
        raise ValueError(
            f"Missing required column(s): {preview}. "
            "Use the downloadable template (exact headers from the Old SKUs list)."
        )
    return {h: found[h] for h in expected}


def parse_old_skus_spreadsheet(filename: str, content: bytes) -> Tuple[List[Dict[str, str]], int]:
    lower = (filename or "").lower()
    if not lower.endswith((".xlsx", ".xlsm", ".xls")):
        raise ValueError("Upload an .xlsx file matching the Old SKUs template.")

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        sheet = _find_sheet(wb, SHEET_NAME)
        if sheet is None:
            raise ValueError("Workbook has no sheets.")
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], 0

        mapping = _header_index_map(header_row, HEADERS, REQUIRED_HEADER)
        valid: List[Dict[str, str]] = []
        invalid = 0
        for row in rows_iter:
            row_data: Dict[str, str] = {}
            any_value = False
            for header in HEADERS:
                idx = mapping[header]
                value = _cell_str(row[idx]) if idx < len(row) else ""
                if value:
                    any_value = True
                row_data[header] = value
            if not any_value:
                continue
            key = (row_data.get(REQUIRED_HEADER) or "").strip()
            if not key:
                invalid += 1
                continue
            valid.append(row_data)
        return valid, invalid
    finally:
        wb.close()


def old_sku_row_to_record(row_data: Dict[str, str]) -> Dict[str, Any]:
    return {
        "old_sku": (row_data.get("OLD SKU") or "").strip(),
        "vendor_name": (row_data.get("Vendor Name") or "").strip(),
        "upc_code": (row_data.get("UPC Code") or "").strip(),
        "row_data": row_data,
    }


def dedupe_by_old_sku(rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    """Last row wins for duplicate OLD SKU values within one import file."""
    seen: Dict[str, Dict[str, str]] = {}
    for row in rows:
        key = (row.get("OLD SKU") or "").strip()
        if key:
            seen[key] = row
    return list(seen.values())
