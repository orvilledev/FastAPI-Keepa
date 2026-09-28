"""Tests for the SMW Shipment Analyzer (box contents request vs Amazon pack list)."""
from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest
from openpyxl import Workbook

from app.services.smw_shipment_analyzer import (
    COMPARISON_SHEET,
    CROSS_SHEET,
    DISCREPANCIES_SHEET,
    FILES_SHEET,
    KIND_CARTON,
    KIND_PACK_LIST,
    STATUS_MATCH,
    STATUS_ONLY_CARTON,
    STATUS_ONLY_PACK_LIST,
    STATUS_QTY,
    SUMMARY_SHEET,
    UNRESOLVED_SHEET,
    SmwShipmentAnalyzerError,
    analyze_advanced,
    analyze_basic,
    parse_upload,
)

# One item line in the request: (sku, upc, description, qty, weight).
Item = tuple[str, str, str, float, float]


def bc_request(
    cartons: list[list[Item]],
    *,
    shipment_id: str = "FBA19TEST001",
) -> bytes:
    """A box contents request: spaced header, ``Carton#:`` blocks, totals footer."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["PO#:", shipment_id])
    sheet.append([])
    sheet.append(
        [
            "Sku", None, "UPC", None, "Description", None, None, "Qty", None, "Weight",
            None, "Carton Length", None, "Carton Width", None, "Carton Height",
        ]
    )
    sheet.append([])
    total_units = 0.0
    total_weight = 0.0
    for index, items in enumerate(cartons, start=1):
        sheet.append(
            ["Carton#:", f"0000827399207414{index:04d}"]
            + [None] * 9
            + [24, None, 16, None, 11]
        )
        sheet.append([])
        carton_units = 0.0
        carton_weight = 0.0
        for sku, upc, description, qty, weight in items:
            sheet.append([sku, None, upc, None, description, None, None, qty, None, weight])
            sheet.append([])
            carton_units += qty
            carton_weight += weight
        sheet.append(
            [None] * 5 + ["Total", None, carton_units, None, f" {carton_weight:.2f}"]
        )
        sheet.append([])
        total_units += carton_units
        total_weight += carton_weight
    sheet.append(
        ["Total Ctns:", len(cartons), "Total Units:", total_units, "Total Weight:", total_weight]
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def pack_list(
    rows: list[tuple[str, str, list[float]]],
    *,
    shipment_id: str = "FBA19TEST001",
    boxes: int = 2,
    declared_units: float | None = None,
    declared_skus: int | None = None,
) -> bytes:
    """An Amazon pack list export: metadata preamble then a per-box unit matrix."""
    body_units = sum(sum(per_box) for _, _, per_box in rows)
    lines = [
        ["Workflow name", "wf-test"],
        ["Shipment ID", shipment_id],
        ["Shipment name", " WR TEST 1 of 2 "],
        ["Ship to", "RANCHO CORDOVA, CA"],
        ["Boxes", str(boxes)],
        ["SKUs", str(declared_skus if declared_skus is not None else len(rows))],
        ["Units", str(int(declared_units if declared_units is not None else body_units))],
        [],
        [f"Individual units ({boxes} boxes)"],
        ["SKU", "Title", "ASIN", "FNSKU", "Condition", "Prep type", "Total units"]
        + [f"Box {index} units" for index in range(1, boxes + 1)],
    ]
    for upc, title, per_box in rows:
        padded = list(per_box) + [0.0] * (boxes - len(per_box))
        lines.append(
            [f"{upc}-FNSKU", title, "B0TEST", "X00TEST", "New", "No prep needed", sum(padded)]
            + ["" if not qty else qty for qty in padded]
        )
    return "\n".join(",".join(f'"{cell}"' for cell in line) for line in lines).encode("utf-8")


def shipment_plan(
    rows: list[tuple[str, str, float]],
    *,
    shipment_number: int = 1,
    declared_units: float | None = None,
) -> bytes:
    """The other pack-list export: a Quantity column, a UPC column, no boxes.

    This one names the shipment nowhere in the file — only in its filename.
    """
    body_units = sum(qty for _, _, qty in rows)
    lines = [
        ["Shipment number", str(shipment_number)],
        ["Workflow name", "wf-test"],
        ["SKUs", str(len(rows))],
        ["Units", str(int(declared_units if declared_units is not None else body_units))],
        [],
        [
            "SKU", "Title", "ASIN", "FNSKU", "UPC/EAN/ISBN/JAN/CODABAR",
            "Condition", "Prep type", "Quantity",
        ],
    ]
    for upc, title, qty in rows:
        lines.append(
            [
                f"{upc}-FNSKU", title, "B0TEST", "X00TEST", f"UPC:{upc}",
                "New", "FC_PROVIDED", qty,
            ]
        )
    return "\n".join(",".join(f'"{cell}"' for cell in line) for line in lines).encode("utf-8")


def _sheet_rows(file_bytes: bytes, sheet_name: str) -> list[tuple]:
    workbook = openpyxl.load_workbook(BytesIO(file_bytes))
    sheet = workbook[sheet_name]
    return list(sheet.iter_rows(values_only=True))


def _tabs(file_bytes: bytes) -> list[str]:
    return openpyxl.load_workbook(BytesIO(file_bytes)).sheetnames


# --- Parsing --------------------------------------------------------------


def test_parse_detects_each_shape_without_being_told():
    request = parse_upload(
        "FBA19TEST001 - bc request.xls",
        bc_request([[("SKU-A", "193391677910", "K CL THRM", 2, 1.76)]]),
    )
    listing = parse_upload(
        "FBA19TEST001 - pack list.csv",
        pack_list([("193391677910", "Smartwool Crew", [2])], boxes=1),
    )
    assert request.kind == KIND_CARTON
    assert listing.kind == KIND_PACK_LIST
    assert request.shipment_id == listing.shipment_id == "FBA19TEST001"
    assert request.units_by_upc == listing.units_by_upc == {"193391677910": 2}


def test_parse_request_skips_carton_total_rows_and_reads_its_footer():
    parsed = parse_upload(
        "request.xls",
        bc_request(
            [
                [("SKU-A", "193391677910", "K CL THRM", 2, 1.76)],
                [("SKU-B", "605284169933", "K CL THRM", 3, 2.64)],
            ]
        ),
    )
    # 5 units of stock, not 10 — the per-carton Total rows must not be counted.
    assert parsed.total_units == 5
    assert parsed.line_count == 2
    assert parsed.box_count == 2
    assert parsed.declared_units == 5
    assert parsed.declared_box_count == 2


def test_parse_reads_the_shipment_plan_export():
    """A Quantity column, a UPC column, and the shipment id only in the filename."""
    parsed = parse_upload(
        "FBA19Q0KLWRS.csv",
        shipment_plan(
            [
                ("198268465935", "Smartwool Womens Crew, Large", 2),
                ("193392648735", "Smartwool Mens Crew, Medium", 21),
            ]
        ),
    )
    assert parsed.kind == KIND_PACK_LIST
    assert parsed.shipment_id == "FBA19Q0KLWRS"
    assert parsed.units_by_upc == {"198268465935": 2, "193392648735": 21}
    assert parsed.total_units == 23
    assert parsed.declared_units == 23
    # This export carries no box breakdown at all.
    assert parsed.box_count == 0
    assert any("taken from the filename" in note for note in parsed.notes)


def test_shipment_plan_export_compares_against_a_request():
    result = analyze_basic(
        [
            (
                "FBA19Q0KLWRS - bc request.xls",
                bc_request(
                    [[("SKU-A", "198268465935", "Crew", 2, 1.0)]],
                    shipment_id="FBA19Q0KLWRS",
                ),
            ),
            ("FBA19Q0KLWRS.csv", shipment_plan([("198268465935", "Crew Large", 2)])),
        ]
    )
    assert result.shipment_id == "FBA19Q0KLWRS"
    assert result.total_units == 2
    # Boxes cannot be compared, but that is not a discrepancy on its own.
    assert result.discrepancy_count == 0
    checks = {
        row[0]: row[3] for row in _sheet_rows(result.file_bytes, SUMMARY_SHEET) if row and row[0]
    }
    assert checks["Boxes"] == "Not compared"
    assert checks["Shipment ID"] == STATUS_MATCH


def test_advanced_says_so_when_no_request_was_uploaded():
    """Ten pack lists with no counterparts must not read as ten real discrepancies."""
    result = analyze_advanced(
        [
            ("FBA19Q0KLWRS.csv", shipment_plan([("198268465935", "Crew", 2)])),
            ("FBA19Q0M37ZZ.csv", shipment_plan([("198268465935", "Crew", 3)])),
        ]
    )
    assert result.shipment_count == 2
    assert result.upc_count == 0
    assert result.resolved_count == 0
    assert result.unresolved_count == 0
    verdict = _sheet_rows(result.file_bytes, SUMMARY_SHEET)[1][0]
    assert "Nothing could be compared" in verdict
    assert "box contents request" in verdict


def test_parse_rejects_an_unrecognized_file():
    workbook = Workbook()
    workbook.active.append(["something", "else"])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(SmwShipmentAnalyzerError, match="Could not tell what"):
        parse_upload("mystery.xlsx", output.getvalue())


# --- Basic analysis -------------------------------------------------------


def test_basic_analysis_reports_a_clean_match():
    result = analyze_basic(
        [
            (
                "bc request.xls",
                bc_request(
                    [
                        [("SKU-A", "193391677910", "Crew", 2, 1.76)],
                        [("SKU-B", "605284169933", "Bottom", 3, 2.64)],
                    ]
                ),
            ),
            (
                "pack list.csv",
                pack_list(
                    [("193391677910", "Crew", [2, 0]), ("605284169933", "Bottom", [0, 3])]
                ),
            ),
        ]
    )
    assert result.discrepancy_count == 0
    assert result.clean
    assert result.upc_count == 2
    assert result.total_units == 5
    assert result.shipment_id == "FBA19TEST001"
    assert result.filename == "FBA19TEST001 Shipment Analysis.xlsx"
    # No discrepancies means no discrepancy tab.
    assert _tabs(result.file_bytes) == [SUMMARY_SHEET, COMPARISON_SHEET]
    statuses = {row[5] for row in _sheet_rows(result.file_bytes, COMPARISON_SHEET)[1:]}
    assert statuses == {STATUS_MATCH}


def test_basic_analysis_flags_unit_and_upc_differences():
    result = analyze_basic(
        [
            (
                "bc request.xls",
                bc_request(
                    [
                        [
                            ("SKU-A", "193391677910", "Crew", 5, 4.0),
                            ("SKU-C", "198268468868", "Hood", 4, 2.0),
                        ]
                    ]
                ),
            ),
            (
                "pack list.csv",
                pack_list(
                    [
                        ("193391677910", "Crew", [2]),
                        ("605284169933", "Bottom", [7]),
                    ],
                    boxes=1,
                ),
            ),
        ]
    )
    assert result.discrepancy_count == 3
    assert DISCREPANCIES_SHEET in _tabs(result.file_bytes)

    rows = {row[0]: row for row in _sheet_rows(result.file_bytes, COMPARISON_SHEET)[1:]}
    assert rows["193391677910"][2:6] == (5, 2, 3, STATUS_QTY)
    assert rows["198268468868"][2:6] == (4, 0, 4, STATUS_ONLY_CARTON)
    assert rows["605284169933"][2:6] == (0, 7, -7, STATUS_ONLY_PACK_LIST)

    issues = [row[1] for row in _sheet_rows(result.file_bytes, DISCREPANCIES_SHEET)[1:]]
    assert issues == [
        STATUS_QTY,
        "Missing from pack list",
        "Missing from box contents request",
    ]


def test_basic_analysis_flags_a_shipment_id_mismatch():
    result = analyze_basic(
        [
            ("bc request.xls", bc_request([[("SKU-A", "193391677910", "Crew", 2, 1.0)]])),
            (
                "pack list.csv",
                pack_list([("193391677910", "Crew", [2])], shipment_id="FBA19OTHER9", boxes=1),
            ),
        ]
    )
    issues = [row[1] for row in _sheet_rows(result.file_bytes, DISCREPANCIES_SHEET)[1:]]
    assert "Shipment ID mismatch" in issues


def test_basic_analysis_flags_a_stated_total_that_disagrees_with_the_rows():
    result = analyze_basic(
        [
            ("bc request.xls", bc_request([[("SKU-A", "193391677910", "Crew", 2, 1.0)]])),
            (
                "pack list.csv",
                pack_list(
                    [("193391677910", "Crew", [2])],
                    boxes=1,
                    declared_units=3,
                    declared_skus=2,
                ),
            ),
        ]
    )
    issues = [row[1] for row in _sheet_rows(result.file_bytes, DISCREPANCIES_SHEET)[1:]]
    # The pack list preamble's SKU count is ignored. Only the unit total is checked.
    assert issues == ["Stated total disagrees with rows"]


def test_basic_analysis_needs_one_file_of_each_shape():
    request = bc_request([[("SKU-A", "193391677910", "Crew", 2, 1.0)]])
    with pytest.raises(SmwShipmentAnalyzerError, match="one box contents request"):
        analyze_basic([("a.xls", request), ("b.xls", request)])
    with pytest.raises(SmwShipmentAnalyzerError, match="exactly two files"):
        analyze_basic([("a.xls", request)])


# --- Advanced analysis ----------------------------------------------------


def test_advanced_analysis_traces_units_keyed_to_the_wrong_shipment():
    """9 units of one barcode land on shipment A but belong to shipment B."""
    uploads = [
        (
            "A - bc request.xls",
            bc_request(
                [[("SKU-A", "193391677910", "Crew", 20, 10.0)]],
                shipment_id="FBA19AAAA001",
            ),
        ),
        (
            "A - pack list.csv",
            pack_list(
                [("193391677910", "Crew", [11])],
                shipment_id="FBA19AAAA001",
                boxes=1,
            ),
        ),
        (
            "B - bc request.xls",
            bc_request(
                [[("SKU-A", "193391677910", "Crew", 4, 2.0)]],
                shipment_id="FBA19BBBB002",
            ),
        ),
        (
            "B - pack list.csv",
            pack_list(
                [("193391677910", "Crew", [13])],
                shipment_id="FBA19BBBB002",
                boxes=1,
            ),
        ),
    ]
    result = analyze_advanced(uploads)

    assert result.shipment_count == 2
    assert result.file_count == 4
    assert result.resolved_count == 1
    assert result.unresolved_count == 0
    assert CROSS_SHEET in _tabs(result.file_bytes)
    assert UNRESOLVED_SHEET not in _tabs(result.file_bytes)

    match = _sheet_rows(result.file_bytes, CROSS_SHEET)[1]
    upc, _description, units, counted_on, expected_on, confidence, finding = match
    assert upc == "193391677910"
    assert units == 9
    assert counted_on == "FBA19AAAA001"
    assert expected_on == "FBA19BBBB002"
    assert confidence.startswith("Exact")
    assert "FBA19BBBB002" in finding

    assert FILES_SHEET in _tabs(result.file_bytes)
    assert len(_sheet_rows(result.file_bytes, FILES_SHEET)) == 5


def test_advanced_analysis_matches_a_barcode_swap_inside_one_shipment():
    result = analyze_advanced(
        [
            (
                "A - bc request.xls",
                bc_request([[("SKU-A", "193391677910", "Crew Medium", 6, 3.0)]]),
            ),
            (
                "A - pack list.csv",
                pack_list([("605284169933", "Crew Medium", [6])], boxes=1),
            ),
        ]
    )
    assert result.resolved_count == 1
    assert result.unresolved_count == 0
    row = _sheet_rows(result.file_bytes, CROSS_SHEET)[1]
    assert row[5].endswith("same description, different UPC")
    assert "wrong barcode" in row[6]


def test_advanced_analysis_reports_what_it_could_not_explain():
    result = analyze_advanced(
        [
            ("A - bc request.xls", bc_request([[("SKU-A", "193391677910", "Crew", 9, 4.0)]])),
            ("A - pack list.csv", pack_list([("193391677910", "Crew", [4])], boxes=1)),
        ]
    )
    assert result.resolved_count == 0
    assert result.unresolved_count == 1
    assert CROSS_SHEET not in _tabs(result.file_bytes)

    row = _sheet_rows(result.file_bytes, UNRESOLVED_SHEET)[1]
    shipment_id, upc, _description, units, side, searched, finding = row
    assert shipment_id == "FBA19TEST001"
    assert upc == "193391677910"
    assert units == 5
    assert side == "Box contents request"
    assert "1 shipment(s)" in searched
    assert "physical recount" in finding


def test_advanced_analysis_flags_a_file_with_no_counterpart():
    result = analyze_advanced(
        [
            (
                "A - bc request.xls",
                bc_request([[("SKU-A", "193391677910", "Crew", 2, 1.0)]], shipment_id="FBA19AAAA001"),
            ),
            (
                "B - pack list.csv",
                pack_list([("193391677910", "Crew", [2])], shipment_id="FBA19BBBB002", boxes=1),
            ),
        ]
    )
    assert result.shipment_count == 2
    issues = [row[1] for row in _sheet_rows(result.file_bytes, DISCREPANCIES_SHEET)[1:]]
    assert issues == ["Unpaired file", "Unpaired file"]
    # Nothing was compared, so no UPC rows and nothing to reconcile.
    assert result.upc_count == 0
    assert result.resolved_count == 0


def test_advanced_analysis_keeps_shipments_apart():
    uploads = [
        (
            "A - bc request.xls",
            bc_request([[("SKU-A", "193391677910", "Crew", 5, 2.0)]], shipment_id="FBA19AAAA001"),
        ),
        (
            "A - pack list.csv",
            pack_list([("193391677910", "Crew", [5])], shipment_id="FBA19AAAA001", boxes=1),
        ),
        (
            "B - bc request.xls",
            bc_request([[("SKU-B", "605284169933", "Bottom", 7, 3.0)]], shipment_id="FBA19BBBB002"),
        ),
        (
            "B - pack list.csv",
            pack_list([("605284169933", "Bottom", [7])], shipment_id="FBA19BBBB002", boxes=1),
        ),
    ]
    result = analyze_advanced(uploads)
    assert result.discrepancy_count == 0
    assert result.upc_count == 2
    assert result.total_units == 12
    assert _tabs(result.file_bytes) == [SUMMARY_SHEET, COMPARISON_SHEET, FILES_SHEET]
    shipments = [row[0] for row in _sheet_rows(result.file_bytes, COMPARISON_SHEET)[1:]]
    assert shipments == ["FBA19AAAA001", "FBA19BBBB002"]
