"""Scan endpoint: every scan with a box number is stored, and never half-saved."""
import pytest
from fastapi import HTTPException

from app.api import fba_shipment_library as api
from app.models.fba_shipment_library import FbaShipmentScanRequest
from app.services.fba_shipment_library import build_accounting_workbook

USER = {"id": "user-1", "email": "worker@example.com"}
LINE_ID = "11111111-1111-1111-1111-111111111111"
UNKNOWN_ID = "22222222-2222-2222-2222-222222222222"


class FakeRepo:
    """In-memory stand-in for FbaShipmentLibraryRepository."""

    def __init__(self, *, fail_count=False, missing_box_table=False):
        self.fail_count = fail_count
        self.missing_box_table = missing_box_table
        self.lines = [
            {"id": LINE_ID, "line_number": 1, "merchant_sku": "A", "external_id": "111", "shipped": 2},
        ]
        self.counts: dict[str, int] = {}
        self.unknowns: dict[str, int] = {}
        self.box_rows: dict[str, dict] = {}
        self._next = 0

    def get_upload(self, upload_id):
        return {"id": upload_id, "created_by": USER["id"], "is_locked": False}

    def lines_for_barcode(self, upload_id, barcode):
        return [dict(row) for row in self.lines if row["external_id"] == barcode]

    def counts_for_lines(self, line_ids, user_id=None):
        return {key: value for key, value in self.counts.items() if key in line_ids}

    def list_unknown_scans(self, upload_id, user_id=None):
        return [
            {"id": UNKNOWN_ID, "external_id": key, "counted": value}
            for key, value in self.unknowns.items()
        ]

    def list_line_quantities(self, upload_id):
        return [dict(row) for row in self.lines]

    def record_box_scan(self, upload_id, user_id, external_id, box_number, line_id=None):
        if self.missing_box_table:
            raise ValueError("The FBA Order Checker tables are missing.")
        self._next += 1
        row_id = f"box-{self._next}"
        self.box_rows[row_id] = {"external_id": external_id, "box_number": box_number}
        return row_id

    def delete_box_scan(self, box_scan_id):
        self.box_rows.pop(box_scan_id, None)

    def increment_count(self, upload_id, line_id, user_id):
        if self.fail_count:
            raise ValueError("Could not record that scan. Scan it again.")
        self.counts[line_id] = self.counts.get(line_id, 0) + 1
        return self.counts[line_id]

    def increment_unknown_scan(self, upload_id, user_id, external_id):
        if self.fail_count:
            raise ValueError("Could not record that scan. Scan it again.")
        self.unknowns[external_id] = self.unknowns.get(external_id, 0) + 1
        return {"id": UNKNOWN_ID, "external_id": external_id, "counted": self.unknowns[external_id]}


@pytest.fixture()
def wire(monkeypatch):
    def install(repo: FakeRepo):
        monkeypatch.setattr(api, "FbaShipmentLibraryRepository", lambda db: repo)
        monkeypatch.setattr(api, "is_superadmin_user", lambda user, db: False)

    return install


def _scan(barcode, box):
    return api.scan_fba_shipment.__wrapped__("upload-1", FbaShipmentScanRequest(barcode=barcode, box_number=box), USER, None)


def test_a_scan_without_a_box_number_is_rejected(wire):
    repo = FakeRepo()
    wire(repo)
    with pytest.raises(HTTPException) as caught:
        _scan("111", None)
    assert caught.value.status_code == 400
    assert "box number" in caught.value.detail.lower()
    assert repo.counts == {}
    assert repo.box_rows == {}
    assert repo.unknowns == {}


def test_known_and_unknown_scans_store_their_box(wire):
    repo = FakeRepo()
    wire(repo)
    first = _scan("111", 4)
    _scan("111", 4)
    unknown = _scan("999", 9)
    assert first.box_number == 4
    assert first.line.counted == 1
    assert unknown.line.not_in_list is True
    assert sorted((row["external_id"], row["box_number"]) for row in repo.box_rows.values()) == [
        ("111", 4),
        ("111", 4),
        ("999", 9),
    ]
    assert repo.counts == {LINE_ID: 2}
    assert repo.unknowns == {"999": 1}


def test_a_failed_count_removes_the_box_record(wire):
    repo = FakeRepo(fail_count=True)
    wire(repo)
    with pytest.raises(HTTPException):
        _scan("111", 1)
    with pytest.raises(HTTPException):
        _scan("999", 1)
    assert repo.box_rows == {}
    assert repo.counts == {} and repo.unknowns == {}


def test_missing_box_table_stops_the_scan_before_it_is_counted(wire):
    repo = FakeRepo(missing_box_table=True)
    wire(repo)
    with pytest.raises(HTTPException) as caught:
        _scan("111", 1)
    assert "tables are missing" in caught.value.detail
    assert repo.counts == {}


def test_stored_scans_produce_a_workbook_that_matches(wire):
    repo = FakeRepo()
    wire(repo)
    for barcode, box in [("111", 1), ("111", 2), ("999", 2)]:
        _scan(barcode, box)
    lines = [{**row, "counted": repo.counts.get(row["id"], 0)} for row in repo.lines]
    unknown = [{"external_id": key, "counted": value} for key, value in repo.unknowns.items()]
    boxed = list(repo.box_rows.values())
    content = build_accounting_workbook(
        shipment_id="FBA1", counted_by="Worker", lines=lines, unknown_lines=unknown, box_scans=boxed
    )
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(content))["Boxed"]
    log = [tuple(c.value for c in row) for row in sheet.iter_rows(min_row=2, max_col=3)]
    assert [row for row in log if row[0] is not None] == [
        ("111", 1, 1),
        ("111", 1, 2),
        ("999", 1, 2),
    ]
