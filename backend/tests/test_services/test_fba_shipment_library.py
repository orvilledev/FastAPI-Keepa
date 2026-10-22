"""Parser tests for the official FBA shipment TSV yellow fields."""
from pathlib import Path

import pytest

from uuid import uuid4

from app.models.fba_shipment_library import FbaShipmentLineResponse
from app.services.fba_shipment_library import (
    FbaShipmentParseError,
    STATUS_ACCOUNTED,
    STATUS_MISSING,
    STATUS_NOT_IN_LIST,
    STATUS_OVERAGE,
    accounting_report_filename,
    assert_box_number_follows,
    box_count_and_missing,
    build_accounting_workbook,
    build_boxed_rows,
    labels_by_box,
    share_audience,
    shipment_open_to,
    upload_visibility_filter,
    choose_line_for_scan,
    count_progress,
    counts_kept_in_box,
    is_fully_accounted,
    parse_fba_shipment_tsv,
)

SAMPLE = Path(r"c:\Users\Administrator\Downloads\FBA19T3YGF07.tsv")

OFFICIAL = (
    "Shipment ID\tFBA19T3YGF07\n"
    "Name\tTO192902\n"
    "Plan ID\t\n"
    "Ship To\tMQY2\n"
    "Total SKUs\t2\n"
    "Total Units\t7\n"
    "\n"
    "Merchant SKU\tTitle\tASIN\tFNSKU\texternal-id \tCondition\tWho will prep?\tPrep Type\tWho will label?\tShipped\n"
    "845028000003\t\"Sockwell, Black\"\tB004WKAH0M\tB004WKAH0M\tUPC: 845028000003\tNew\tMerchant\tPolybagging\tNo prep required\t3\n"
    "845028000010\tSockwell Black\tB003YKTLGM\tB003YKTLGM\tUPC: 845028000010\tNew\tMerchant\tPolybagging\tNo prep required\t4\n"
)


def test_parser_keeps_only_yellow_fields():
    parsed = parse_fba_shipment_tsv(OFFICIAL.encode("utf-8"))

    assert parsed.shipment_id == "FBA19T3YGF07"
    assert parsed.shipment_name == "TO192902"
    assert parsed.total_shipped == 7
    assert [(line.merchant_sku, line.external_id, line.shipped) for line in parsed.lines] == [
        ("845028000003", "845028000003", 3),
        ("845028000010", "845028000010", 4),
    ]


def test_fully_accounted_requires_every_line_to_match_shipped():
    assert is_fully_accounted([]) is False
    assert is_fully_accounted([(3, 3), (1, 1)]) is True
    assert is_fully_accounted([(3, 2)]) is False
    assert is_fully_accounted([(3, 4)]) is False


def test_count_progress_missing_accounted_and_overage():
    assert count_progress(3, 1) == (2, STATUS_MISSING)
    assert count_progress(3, 2) == (1, STATUS_MISSING)
    assert count_progress(3, 3) == (0, STATUS_ACCOUNTED)
    assert count_progress(3, 4) == (0, STATUS_OVERAGE)


def test_unknown_barcode_is_marked_not_in_the_list():
    row = FbaShipmentLineResponse(
        id=uuid4(),
        line_number=0,
        merchant_sku="",
        external_id="UPC: 999000",
        shipped=0,
        counted=2,
        not_in_list=True,
    )
    assert row.external_id == "999000"
    assert row.status == STATUS_NOT_IN_LIST
    assert row.counted == 2


def test_scan_fills_the_line_that_still_has_units_left():
    first = {"id": "a", "line_number": 1, "shipped": 3, "counted": 3}
    second = {"id": "b", "line_number": 2, "shipped": 2, "counted": 0}
    assert choose_line_for_scan([second, first])["id"] == "b"
    assert choose_line_for_scan([{"id": "a", "line_number": 1, "shipped": 3, "counted": 3}])["id"] == "a"


def _read_workbook(content: bytes):
    from io import BytesIO

    from openpyxl import load_workbook

    return load_workbook(BytesIO(content))


def _accounting_counts(sheet) -> dict[str, int]:
    """Counted per external-id from the Accounting rows, including the Not in the list block."""
    totals: dict[str, int] = {}
    in_unknown_block = False
    started = False
    for row in sheet.iter_rows(values_only=True):
        if not started:
            if row[0] == "Merchant SKU":
                started = True
            continue
        if row[0] == "Not in the list":
            in_unknown_block = True
            continue
        if row[0] in (None, "", "external-id"):
            continue
        external_id, counted = (row[0], row[1]) if in_unknown_block else (row[1], row[3])
        totals[str(external_id)] = totals.get(str(external_id), 0) + int(counted or 0)
    return {key: value for key, value in totals.items() if value}


def _boxed_log(sheet) -> list[tuple[str, int, object]]:
    return [
        (str(row[0]), int(row[1]), row[2])
        for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True)
        if row[0] not in (None, "")
    ]


def _pivot_cells(sheet) -> dict[tuple[str, object], int]:
    """Read the pivot grid at G:... into {(external-id, box): count}."""
    header = [cell.value for cell in sheet[2]][7:]
    boxes = header[:-1]
    cells: dict[tuple[str, object], int] = {}
    row = 3
    while sheet.cell(row=row, column=7).value not in (None, "Grand Total"):
        external_id = sheet.cell(row=row, column=7).value
        for offset, box in enumerate(boxes):
            value = sheet.cell(row=row, column=8 + offset).value
            if value:
                cells[(str(external_id), None if box == "(blank)" else box)] = int(value)
        row += 1
    return cells


def test_boxed_tab_matches_the_users_example_workbook():
    lines = [{"merchant_sku": "845028000003", "external_id": "845028000003", "shipped": 3, "counted": 3}]
    scans = [
        {"external_id": "845028000003", "box_number": 1},
        {"external_id": "845028000003", "box_number": 1},
        {"external_id": "845028000003", "box_number": 2},
    ]
    workbook = _read_workbook(
        build_accounting_workbook(
            shipment_id="FBA19T3YGF07",
            shipment_name="TO192902",
            counted_by="Test",
            lines=lines,
            unknown_lines=[],
            box_scans=scans,
        )
    )
    assert workbook.sheetnames == ["Accounting", "Boxed"]
    accounting, boxed = workbook["Accounting"], workbook["Boxed"]
    assert [c.value for c in accounting[1]][:2] == ["Name", "TO192902"]
    assert [c.value for c in accounting[2]][:2] == ["Shipment ID", "FBA19T3YGF07"]
    assert [c.value for c in accounting[3]][:2] == ["Counted by", "Test"]
    assert accounting[4][0].value == "Generated"
    assert str(accounting[4][1].value).endswith(" UTC")
    assert accounting["D1"].value == "Box count"
    assert accounting["E1"].value == 2
    assert accounting["D2"].value == "Missing box"
    assert accounting["E2"].value is None
    assert accounting["E2"].fill.patternType is None
    assert [c.value for c in accounting[6]] == [
        "Merchant SKU", "external-id", "Shipped", "Counted", "Missing", "Status", "Boxes",
    ]
    assert [c.value for c in accounting[7]] == [
        "845028000003", "845028000003", 3, 3, 0, "All Accounted", "Box 1: 2, Box 2: 1",
    ]
    assert [c.value for c in boxed[1]][:3] == ["external-id", "Counted", "Box Number"]
    assert _boxed_log(boxed) == [
        ("845028000003", 1, 1),
        ("845028000003", 1, 1),
        ("845028000003", 1, 2),
    ]
    rows = [[c.value for c in boxed[r]][6:10] for r in range(1, 5)]
    assert rows == [
        ["Sum of Counted", "Column Labels", None, None],
        ["Row Labels", 1, 2, "Grand Total"],
        ["845028000003", 2, 1, 3],
        ["Grand Total", 2, 1, 3],
    ]
    assert boxed._pivots[0].location.ref == "G1:J4"


def test_boxed_totals_match_accounting_for_many_items_boxes_and_unknown_scans():
    lines = [
        {"merchant_sku": "A", "external_id": "111", "shipped": 3, "counted": 3},
        {"merchant_sku": "B", "external_id": "222", "shipped": 2, "counted": 1},
        {"merchant_sku": "C", "external_id": "333", "shipped": 1, "counted": 0},
        {"merchant_sku": "D", "external_id": "111", "shipped": 2, "counted": 2},
    ]
    unknown = [{"external_id": "999", "counted": 2}]
    scans = [
        {"external_id": "111", "box_number": 10},
        {"external_id": "222", "box_number": 2},
        {"external_id": "111", "box_number": 2},
        {"external_id": "999", "box_number": 2},
        {"external_id": "111", "box_number": 10},
        {"external_id": "111", "box_number": 2},
        {"external_id": "999", "box_number": 11},
        {"external_id": "111", "box_number": 11},
    ]
    workbook = _read_workbook(
        build_accounting_workbook(
            shipment_id="FBA1", counted_by="Test", lines=lines, unknown_lines=unknown, box_scans=scans
        )
    )
    accounting, boxed = workbook["Accounting"], workbook["Boxed"]
    log = _boxed_log(boxed)

    per_item: dict[str, int] = {}
    for external_id, counted, _ in log:
        per_item[external_id] = per_item.get(external_id, 0) + counted
    assert per_item == _accounting_counts(accounting) == {"111": 5, "222": 1, "999": 2}
    assert all(box is not None for _, _, box in log)

    pivot = _pivot_cells(boxed)
    expected: dict[tuple[str, object], int] = {}
    for external_id, _, box in log:
        expected[(external_id, box)] = expected.get((external_id, box), 0) + 1
    assert pivot == expected
    grand_row = 3 + len(per_item)
    assert boxed.cell(row=grand_row, column=7).value == "Grand Total"
    assert boxed.cell(row=grand_row, column=7 + 1 + 3 + 0).value == sum(per_item.values())
    assert [c.value for c in boxed[2]][7:] == [2, 10, 11, "Grand Total"]


def test_box_numbers_have_to_follow_without_skipping():
    assert_box_number_follows(1, {})
    with pytest.raises(ValueError, match="next box is 1"):
        assert_box_number_follows(3, {})
    assert_box_number_follows(2, {1: 4})
    with pytest.raises(ValueError, match="next box is 2"):
        assert_box_number_follows(3, {1: 2})
    assert_box_number_follows(1, {1: 1, 2: 1})
    assert_box_number_follows(3, {1: 1, 2: 1})
    assert_box_number_follows(2, {1: 1, 3: 1})
    with pytest.raises(ValueError, match="next box is 2"):
        assert_box_number_follows(4, {1: 1, 3: 1})
    assert_box_number_follows(3, {1: 1, 3: 1}, previous_box=3)
    with pytest.raises(ValueError, match="Box 2 cannot be left empty"):
        assert_box_number_follows(1, {1: 1, 2: 1, 3: 1}, previous_box=2)
    assert_box_number_follows(1, {1: 1, 2: 1, 3: 1}, previous_box=3)
    assert_box_number_follows(4, {1: 1, 2: 2, 3: 1}, previous_box=2)


def test_missing_box_count_is_the_gap_under_the_highest_box():
    used = [("111", number) for number in range(1, 41) if number != 15]
    assert box_count_and_missing(used) == (40, 1)
    assert box_count_and_missing([("111", 1), ("111", 2)]) == (2, 0)
    assert box_count_and_missing([("111", None), ("111", None)]) == (0, 0)


def test_existing_shared_shipments_stay_open_to_everyone_until_the_owner_chooses_people():
    owner = "aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa"
    worker = "bbbbbbbb-bbbb-bbbb-bbbb-bbbbbbbbbbbb"
    other = "cccccccc-cccc-cccc-cccc-cccccccccccc"
    already_shared = {"created_by": owner, "is_locked": False}
    legacy_null = {"created_by": owner, "is_locked": False, "shared_with": None}
    locked = {"created_by": owner, "is_locked": True, "shared_with": None}
    custom = {"created_by": owner, "is_locked": False, "shared_with": [worker]}

    assert share_audience(already_shared) is None
    assert share_audience(legacy_null) is None
    assert shipment_open_to(already_shared, other) is True
    assert shipment_open_to(legacy_null, other) is True
    assert shipment_open_to(locked, other) is False
    assert shipment_open_to(locked, owner) is True
    assert shipment_open_to(custom, worker) is True
    assert shipment_open_to(custom, other) is False
    assert shipment_open_to(custom, owner) is True
    assert shipment_open_to({"created_by": owner, "is_locked": False, "shared_with": []}, other) is False
    assert "shared_with.is.null" in upload_visibility_filter(worker)
    assert f"shared_with.cs.{{{worker}}}" in upload_visibility_filter(worker)


def test_accounting_header_puts_one_missing_box_in_red():
    lines = [{"merchant_sku": "A", "external_id": "111", "shipped": 3, "counted": 3}]
    scans = [
        {"external_id": "111", "box_number": 1},
        {"external_id": "111", "box_number": 2},
        {"external_id": "111", "box_number": 4},
    ]
    accounting = _read_workbook(
        build_accounting_workbook(
            shipment_id="FBA1", counted_by="Test", lines=lines, unknown_lines=[], box_scans=scans
        )
    )["Accounting"]
    assert accounting["D1"].value == "Box count"
    assert accounting["E1"].value == 4
    assert accounting["D2"].value == "Missing box"
    assert accounting["E2"].value == 1
    assert accounting["E2"].fill.patternType == "solid"
    assert str(accounting["E2"].fill.fgColor.rgb).upper().endswith("FF0000")
    assert accounting["D2"].fill.patternType is None


def test_units_for_count_keeps_one_slot_per_counted_unit():
    from app.services.fba_shipment_library import units_for_count

    units = units_for_count(
        3,
        [{"id": "scan-1", "box_number": 4}, {"id": "scan-2", "box_number": 9}],
    )
    assert units == [
        {"unit_number": 1, "scan_id": "scan-1", "box_number": 4},
        {"unit_number": 2, "scan_id": "scan-2", "box_number": 9},
        {"unit_number": 3, "scan_id": None, "box_number": None},
    ]
    assert units_for_count(0, [{"id": "scan-1", "box_number": 1}]) == []


def test_units_counted_without_a_box_still_add_up_to_accounting():
    lines = [{"merchant_sku": "A", "external_id": "111", "shipped": 4, "counted": 4}]
    scans = [{"external_id": "111", "box_number": 1}, {"external_id": "111", "box_number": 1}]
    workbook = _read_workbook(
        build_accounting_workbook(
            shipment_id="FBA1", counted_by="Test", lines=lines, unknown_lines=[], box_scans=scans
        )
    )
    boxed = workbook["Boxed"]
    assert _boxed_log(boxed) == [("111", 1, 1), ("111", 1, 1), ("111", 1, None), ("111", 1, None)]
    assert [c.value for c in boxed[2]][7:] == [1, "(blank)", "Grand Total"]
    assert _pivot_cells(boxed) == {("111", 1): 2, ("111", None): 2}


def test_a_box_count_ignores_units_placed_in_other_boxes():
    scans = [
        {"id": "1", "line_id": "line-a", "box_number": 1, "created_at": "1"},
        {"id": "2", "line_id": "line-a", "box_number": 1, "created_at": "2"},
        {"id": "3", "line_id": "line-a", "box_number": 4, "created_at": "3"},
        {"id": "4", "line_id": "line-a", "box_number": 4, "created_at": "4"},
    ]
    assert counts_kept_in_box(scans, 4, {"line-a": 3}, "line_id") == {"line-a": 1}
    assert counts_kept_in_box(scans, 1, {"line-a": 3}, "line_id") == {"line-a": 2}


def test_box_scans_beyond_the_counted_total_are_never_exported():
    lines = [{"merchant_sku": "A", "external_id": "111", "shipped": 2, "counted": 1}]
    scans = [
        {"external_id": "111", "box_number": 1},
        {"external_id": "111", "box_number": 2},
        {"external_id": "555", "box_number": 3},
    ]
    rows = build_boxed_rows(lines=lines, unknown_lines=[], box_scans=scans)
    assert rows == [("111", 1)]


def test_boxed_rows_are_listed_box_by_box():
    lines = [
        {"merchant_sku": "A", "external_id": "222", "shipped": 2, "counted": 2},
        {"merchant_sku": "B", "external_id": "111", "shipped": 3, "counted": 3},
    ]
    scans = [
        {"external_id": "111", "box_number": 5},
        {"external_id": "222", "box_number": 3},
        {"external_id": "111", "box_number": 3},
        {"external_id": "222", "box_number": 5},
    ]
    rows = build_boxed_rows(lines=lines, unknown_lines=[], box_scans=scans)
    assert rows == [("111", 3), ("222", 3), ("111", 5), ("222", 5), ("111", None)]


def test_boxed_tab_without_any_scans_has_only_headers():
    lines = [{"merchant_sku": "A", "external_id": "111", "shipped": 2, "counted": 0}]
    workbook = _read_workbook(
        build_accounting_workbook(shipment_id="FBA1", counted_by="Test", lines=lines, unknown_lines=[])
    )
    boxed = workbook["Boxed"]
    assert [[c.value for c in row] for row in boxed.iter_rows(max_col=3)] == [
        ["external-id", "Counted", "Box Number"]
    ]
    assert not boxed._pivots


def test_scan_request_validates_box_number():
    from pydantic import ValidationError

    from app.models.fba_shipment_library import FbaShipmentScanRequest

    assert FbaShipmentScanRequest(barcode="1", box_number=7).box_number == 7
    assert FbaShipmentScanRequest(barcode="1").box_number is None
    for bad in (0, -1, 100000):
        with pytest.raises(ValidationError):
            FbaShipmentScanRequest(barcode="1", box_number=bad)


def test_fnsku_labels_are_grouped_by_box_in_box_order():
    lines = [
        {"merchant_sku": "A", "external_id": "111", "shipped": 3, "counted": 3},
        {"merchant_sku": "B", "external_id": "222", "shipped": 2, "counted": 2},
    ]
    scans = [
        {"external_id": "222", "box_number": 10},
        {"external_id": "111", "box_number": 2},
        {"external_id": "111", "box_number": 2},
        {"external_id": "222", "box_number": 10},
        {"external_id": "111", "box_number": 2},
        {"external_id": "999", "box_number": 1},
    ]
    assert labels_by_box(lines=lines, box_scans=scans) == [
        (2, [("111", 3)]),
        (10, [("222", 2)]),
    ]
    assert labels_by_box(
        lines=lines,
        box_scans=scans,
        unknown_lines=[{"external_id": "999", "counted": 1}],
    ) == [
        (2, [("111", 3)]),
        (10, [("222", 2)]),
    ]


def test_fnsku_labels_require_a_box_for_every_counted_unit():
    lines = [{"merchant_sku": "A", "external_id": "111", "shipped": 2, "counted": 2}]
    scans = [{"external_id": "111", "box_number": 1}]
    with pytest.raises(ValueError, match="box number"):
        labels_by_box(lines=lines, box_scans=scans)


def test_fnsku_labels_for_one_box_ignore_units_outside_that_box():
    lines = [
        {"merchant_sku": "A", "external_id": "111", "shipped": 2, "counted": 2},
        {"merchant_sku": "B", "external_id": "222", "shipped": 1, "counted": 1},
    ]
    scans = [
        {"external_id": "111", "box_number": 1},
        {"external_id": "111", "box_number": 2},
    ]
    unknown = [{"merchant_sku": "", "external_id": "810173985413", "counted": 1}]
    scans.append({"external_id": "810173985413", "box_number": 1})
    assert labels_by_box(
        lines=lines,
        box_scans=scans,
        box_number=1,
        unknown_lines=unknown,
    ) == [
        (1, [("111", 1), ("810173985413", 1)]),
    ]
    with pytest.raises(ValueError, match="filter"):
        labels_by_box(lines=lines, box_scans=scans, box_number=9)


def test_fnsku_labels_for_a_sku_search_skip_other_items():
    lines = [
        {"merchant_sku": "845028000003", "external_id": "845028000003", "shipped": 2, "counted": 2},
        {"merchant_sku": "999", "external_id": "111", "shipped": 1, "counted": 1},
    ]
    scans = [
        {"external_id": "845028000003", "box_number": 1},
        {"external_id": "845028000003", "box_number": 4},
        {"external_id": "111", "box_number": 1},
    ]
    assert labels_by_box(lines=lines, box_scans=scans, search="845028") == [
        (1, [("845028000003", 1)]),
        (4, [("845028000003", 1)]),
    ]
    assert labels_by_box(
        lines=lines,
        box_scans=scans,
        box_number=1,
        search="845028",
    ) == [
        (1, [("845028000003", 1)]),
    ]


def test_report_filename_uses_the_upload_name_not_the_shipment_id():
    assert accounting_report_filename("TO192902", "FBA19T3YGF07") == "TO192902.xlsx"
    assert accounting_report_filename("  TO192902  ", "FBA19T3YGF07") == "TO192902.xlsx"
    assert accounting_report_filename('TO/1:2*?"<>|', "FBA1") == "TO 1 2.xlsx"
    assert accounting_report_filename("", "FBA19T3YGF07") == "FBA19T3YGF07 Accounting.xlsx"
    assert accounting_report_filename("", "") == "shipment Accounting.xlsx"


def test_parser_reads_a_missing_name_as_blank():
    text = OFFICIAL.replace("Name\tTO192902\n", "")
    parsed = parse_fba_shipment_tsv(text.encode("utf-8"))
    assert parsed.shipment_id == "FBA19T3YGF07"
    assert parsed.shipment_name == ""
    assert len(parsed.lines) == 2


def test_parser_rejects_a_file_without_the_item_header():
    with pytest.raises(FbaShipmentParseError, match="Merchant SKU"):
        parse_fba_shipment_tsv(b"Shipment ID\tFBA1\nName\tTO1\n")


@pytest.mark.skipif(not SAMPLE.is_file(), reason="sample TSV is not on this machine")
def test_official_sample_file():
    parsed = parse_fba_shipment_tsv(SAMPLE.read_bytes())

    assert parsed.shipment_id == "FBA19T3YGF07"
    assert parsed.shipment_name == "TO192902"
    assert parsed.lines[0].merchant_sku == "845028000003"
    assert parsed.lines[0].external_id == "845028000003"
    assert parsed.lines[0].shipped == 3
    assert parsed.lines[1].shipped == 3
    assert len(parsed.lines) == 287
    assert parsed.total_shipped == 1512
