"""Tests for FBA Box Contents Tool #2 (spaced PO# carton dump)."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from app.services.fba_box_contents_tool2 import (
    BOX_CONTENTS_SHEET,
    DEFAULT_OUTPUT_FILENAME,
    DIMENSIONS_SHEET,
    PIVOT_START_COL,
    FbaBoxContentsTool2Error,
    generate_fba_box_contents_tool2,
    parse_carton_detail_tool2,
    sanitize_download_filename,
)

SAMPLE_INPUT = Path(r"c:\Users\Administrator\Downloads\FBA19PLBV097.xls")
SAMPLE_OUTPUT = Path(r"c:\Users\Administrator\Downloads\FBA19PLBV097 Output.xlsx")


def _spaced_xlsx(
    rows: list[list[object]],
    *,
    shipment_id: str = "FBA19TEST001",
) -> bytes:
    """Build a spaced-column carton dump similar to the Tool #2 sample."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["PO#:", shipment_id])
    sheet.append([])
    sheet.append(
        [
            "Sku",
            None,
            "UPC",
            None,
            "Description",
            None,
            None,
            "Qty",
            None,
            "Weight",
            None,
            "Carton Length",
            None,
            "Carton Width",
            None,
            "Carton Height",
        ]
    )
    sheet.append([])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_sanitize_filename_adds_output_suffix():
    assert sanitize_download_filename("FBA19PLBV097.xls") == "FBA19PLBV097 Output.xlsx"
    assert sanitize_download_filename("FBA19PLBV097 Output.xlsx") == "FBA19PLBV097 Output.xlsx"
    assert sanitize_download_filename(None) == DEFAULT_OUTPUT_FILENAME


def test_parse_requires_carton_headers():
    raw = _spaced_xlsx(
        [["016349-L", None, "193392648759", None, "DESC", None, None, 50, None, 35.0]]
    )
    with pytest.raises(FbaBoxContentsTool2Error, match="Carton#"):
        parse_carton_detail_tool2(raw)


def test_generate_tool2_layout_and_types():
    raw = _spaced_xlsx(
        [
            ["Carton#:", "00001111", None, None, None, None, None, None, None, None, None, 19.88, None, 17.79, None, 14.59],
            [],
            ["016349-L", None, "193392648759", None, "M CL", None, None, 50.0, None, 35.0],
            [],
            [None, None, None, None, None, "Total", None, 50.0, None, " 38.00"],
            [],
            ["Carton#:", "00002222", None, None, None, None, None, None, None, None, None, 24.0, None, 16.0, None, 11.0],
            [],
            ["0SN110-XS", None, "605284169933", None, "K CL", None, None, 2.0, None, 1.76],
            ["0SN110-S", None, "605284169940", None, "K CL", None, None, 2.0, None, 1.76],
            [None, None, None, None, None, "Total", None, 4.0, None, " 14.82"],
        ]
    )
    result = generate_fba_box_contents_tool2(raw, "FBA19PLBV097.xls")
    assert result.filename == "FBA19PLBV097 Output.xlsx"
    assert result.box_count == 2
    assert result.row_count == 3
    assert result.total_qty == 54
    assert result.shipment_id == "FBA19TEST001"
    assert result.upc_count == 3

    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes))
    try:
        assert workbook.sheetnames == [BOX_CONTENTS_SHEET, DIMENSIONS_SHEET]
        contents = workbook[BOX_CONTENTS_SHEET]
        dims = workbook[DIMENSIONS_SHEET]

        assert contents["A1"].value == "UPC"
        assert contents["B1"].value == "Box Number"
        assert contents["C1"].value == "QTY"
        assert contents["A1"].font.bold is True
        assert contents["G1"].value == "Sum of QTY"
        assert contents["G1"].font.bold is False

        assert contents["A2"].value == "193392648759"
        assert contents["A2"].data_type == "s"
        assert contents["B2"].value == 1
        assert type(contents["B2"].value) is int
        assert contents["C2"].value == 50

        # No Total QTY footer on the left
        for row in range(1, contents.max_row + 1):
            assert contents.cell(row, 2).value != "Total QTY"

        assert [dims.cell(1, c).value for c in range(1, 6)] == [
            "Box Number", "Weight", "Length", "Width", "Height",
        ]
        assert dims["A1"].font.bold is False
        assert dims.max_column == 5  # no zero values -> no Zero Correction column
        assert result.zero_corrected_count == 0

        assert dims["A2"].value == 1
        assert dims["B2"].value == " 38.00"
        assert dims["B2"].data_type == "s"
        assert dims["C2"].value == 19.88
        assert dims["D2"].value == 17.79
        assert dims["E2"].value == 14.59
        assert dims["A3"].value == 2
        assert dims["B3"].value == " 14.82"
        assert dims["C3"].value == 24
    finally:
        workbook.close()


@pytest.mark.skipif(not SAMPLE_INPUT.exists() or not SAMPLE_OUTPUT.exists(), reason="sample files not present")
def test_matches_sample_output_values():
    result = generate_fba_box_contents_tool2(SAMPLE_INPUT.read_bytes(), SAMPLE_INPUT.name)
    assert result.filename == "FBA19PLBV097 Output.xlsx"
    assert result.shipment_id == "FBA19PLBV097"
    assert result.box_count == 29
    assert result.row_count == 348
    assert result.upc_count == 321

    got = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=True)
    expected = openpyxl.load_workbook(SAMPLE_OUTPUT, data_only=True)
    try:
        exp_contents = expected[BOX_CONTENTS_SHEET]
        got_contents = got[BOX_CONTENTS_SHEET]
        exp_dims = expected[DIMENSIONS_SHEET]
        got_dims = got[DIMENSIONS_SHEET]

        # Left side UPC / Box Number / QTY
        for row in range(1, exp_contents.max_row + 1):
            for col in (1, 2, 3):
                exp_val = exp_contents.cell(row, col).value
                got_val = got_contents.cell(row, col).value
                if exp_val is None and got_val is None:
                    continue
                if col == 1 and exp_val is not None:
                    assert str(got_val) == str(exp_val), f"left {row},{col}"
                else:
                    assert got_val == exp_val, f"left {row},{col}: {got_val!r} != {exp_val!r}"

        # Pivot block
        for row in range(1, exp_contents.max_row + 1):
            for col in range(PIVOT_START_COL, exp_contents.max_column + 1):
                exp_val = exp_contents.cell(row, col).value
                got_val = got_contents.cell(row, col).value
                if exp_val is None and got_val is None:
                    continue
                if isinstance(exp_val, str):
                    assert str(got_val) == exp_val, f"pivot {row},{col}"
                else:
                    assert got_val == exp_val, f"pivot {row},{col}: {got_val!r} != {exp_val!r}"

        assert result.zero_corrected_count == 0
        # Output now leads with a Box Number column (the sample has none).
        assert got_dims.max_row == exp_dims.max_row
        for row in range(1, exp_dims.max_row + 1):
            if row > 1:
                assert got_dims.cell(row, 1).value == row - 1
            for col in range(1, 5):
                exp_val = exp_dims.cell(row, col).value
                got_val = got_dims.cell(row, col + 1).value
                assert got_val == exp_val, f"dims {row},{col}: {got_val!r} != {exp_val!r}"
    finally:
        got.close()
        expected.close()


def _carton(box_id, length, width, height, upc, qty, weight_text):
    return [
        ["Carton#:", box_id, None, None, None, None, None, None, None, None, None, length, None, width, None, height],
        [],
        ["SKU-" + upc[-3:], None, upc, None, "DESC", None, None, float(qty), None, 1.0],
        [],
        [None, None, None, None, None, "Total", None, float(qty), None, weight_text],
        [],
    ]


def test_zero_weight_gets_suggestion_reason_and_highlight():
    rows = []
    rows += _carton("C1", 20.0, 16.0, 12.0, "111111111111", 60, " 14.10")
    rows += _carton("C2", 22.0, 15.0, 10.0, "222222222222", 72, " 14.70")
    rows += _carton("C3", 22.1, 15.1, 10.0, "222222222222", 72, " 0.00")
    rows += _carton("C4", 20.1, 16.0, 12.1, "111111111111", 60, " 14.20")
    raw = _spaced_xlsx(rows)

    result = generate_fba_box_contents_tool2(raw, "FBA19ZERO.xlsx")

    assert result.zero_corrected_count == 1
    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes))
    try:
        dims = workbook[DIMENSIONS_SHEET]
        assert [dims.cell(1, c).value for c in range(1, 7)] == [
            "Box Number", "Weight", "Length", "Width", "Height", "Zero Correction",
        ]
        assert [dims.cell(r, 1).value for r in range(2, 6)] == [1, 2, 3, 4]

        # Box 3 (zero) mirrors box 2 — same UPC and qty, near-identical size.
        assert dims["B4"].value == " 14.70"
        assert dims["B4"].data_type == "s"
        assert "Box 2" in dims["F4"].value
        assert "222222222222" in dims["F4"].value
        for col in range(2, 7):
            assert dims.cell(4, col).fill.fgColor.rgb == "FFFFFF00"

        # Untouched rows keep their values, no reason and no fill.
        for r in (2, 3, 5):
            assert dims.cell(r, 6).value is None
            assert dims.cell(r, 2).fill.fill_type is None
        assert dims["B2"].value == " 14.10"
        assert dims["B5"].value == " 14.20"
    finally:
        workbook.close()


def test_zero_weight_without_any_reference_is_flagged_not_invented():
    raw = _spaced_xlsx(_carton("C1", 20.0, 16.0, 12.0, "111111111111", 60, " 0.00"))
    result = generate_fba_box_contents_tool2(raw, "FBA19ONLY.xlsx")

    dims = openpyxl.load_workbook(BytesIO(result.file_bytes))[DIMENSIONS_SHEET]
    assert result.zero_corrected_count == 1
    assert dims["B2"].value == " 0.00"
    assert "enter it manually" in dims["F2"].value


def test_zero_weight_averages_similar_boxes_when_no_identical_carton():
    rows = []
    rows += _carton("C1", 23.0, 16.0, 12.0, "111111111111", 60, " 16.00")
    rows += _carton("C2", 23.1, 16.1, 12.0, "222222222222", 60, " 17.00")
    rows += _carton("C3", 23.0, 16.0, 12.1, "333333333333", 60, " 0.00")
    result = generate_fba_box_contents_tool2(_spaced_xlsx(rows), "FBA19AVG.xlsx")

    dims = openpyxl.load_workbook(BytesIO(result.file_bytes))[DIMENSIONS_SHEET]
    assert dims["B4"].value == " 16.50"
    assert "average of Boxes 1 and 2" in dims["F4"].value
