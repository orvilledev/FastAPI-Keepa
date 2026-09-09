"""Tests for the Shipment Manager FBA export -> WR SKU Update / PO Import conversions."""
import io

import pytest
from openpyxl import load_workbook

from app.services.shipment_manager import (
    OUTPUT_FILENAME,
    ShipmentManagerError,
    ShipmentSkuRow,
    build_po_import_workbook,
    build_workbook,
    build_wr_sku_update,
    compile_stored_rows,
    dedupe_by_upc,
    parse_fba_export,
    po_import_filename,
    resolve_po_supplier,
    supplier_from_filename,
    supplier_from_title,
    supplier_from_titles,
)

FBA_EXPORT = """Workflow name,wf5acb45b4-38bc-4c72-b878-c6b522ac6f99,,
Shipment ID,FBA19JHYH77Q,,
Shipment name,NFA WHRP 7.17.26 1 OF 5,,
Ship to,DEN8,,
Boxes,38,,
SKUs,2,,
Units,140,,
,,,
Individual units (38 boxes),,,
SKU,Title,ASIN,FNSKU,Condition,Prep type,Total units,Box 1 units
197642130629-FNSKU,"Borealis Backpack, Blk",B0CN9SD8JF,X0052Z3NU3,New,FC_PROVIDED,139,4
198268844372-FNSKU,"Antora Jacket, XL",B0FX3FSCDN,X0057D6IM5,New,FC_PROVIDED,1,1
,,,,,,,
,,,,,,Box ID,FBA19JHYH77QU000022
,,,,,,Box weight (lb):,11
"""


def _run(text: str = FBA_EXPORT, filename: str = "FBA19JHYH77Q.csv"):
    return build_wr_sku_update(filename, text.encode("utf-8"))


def _sheet(result):
    return load_workbook(io.BytesIO(result.workbook_bytes)).active


def test_parses_shipment_metadata():
    result = _run()
    assert result.shipment_id == "FBA19JHYH77Q"
    assert result.shipment_name == "NFA WHRP 7.17.26 1 OF 5"
    assert result.ship_to == "DEN8"
    assert result.box_count == 38
    assert result.sku_count == 2
    assert result.total_units == 140
    assert result.filename == OUTPUT_FILENAME


def test_maps_columns_from_the_sku_table():
    sheet = _sheet(_run())
    assert [sheet.cell(1, col).value for col in range(1, 5)] == [
        "SKU",
        "Description",
        "UPC",
        "FNSKU",
    ]
    assert sheet.cell(2, 1).value == "197642130629-FNSKU"
    assert sheet.cell(2, 2).value == "Borealis Backpack, Blk"
    assert sheet.cell(2, 4).value == "X0052Z3NU3"
    assert sheet.cell(3, 1).value == "198268844372-FNSKU"


def test_upc_is_the_sku_without_the_fnsku_suffix_as_a_number():
    cell = _sheet(_run())["C2"]
    assert cell.value == 197642130629
    assert isinstance(cell.value, int)
    assert cell.number_format == "0"


def test_dimension_columns_are_left_blank():
    sheet = _sheet(_run())
    assert sheet.max_row == 3
    assert all(
        sheet.cell(row, col).value is None
        for row in range(2, sheet.max_row + 1)
        for col in range(5, 14)
    )


def test_stops_at_the_per_box_footer():
    sheet = _sheet(_run())
    # Two SKU rows only; the Box ID / Box weight rows below the blank line are ignored.
    assert sheet.max_row == 3
    assert sheet.cell(4, 1).value is None


def test_template_formatting_is_preserved():
    sheet = _sheet(_run())
    assert sheet["A1"].fill.fgColor.rgb == "FF92D050"
    assert sheet["A1"].font.bold is True
    assert sheet.row_dimensions[1].height == 28.8
    assert round(sheet.column_dimensions["B"].width, 3) == 61.332


def test_duplicate_skus_are_collapsed():
    duplicated = FBA_EXPORT.replace(
        '198268844372-FNSKU,"Antora Jacket, XL",B0FX3FSCDN,X0057D6IM5,New,FC_PROVIDED,1,1',
        '197642130629-FNSKU,"Borealis Backpack, Blk",B0CN9SD8JF,X0052Z3NU3,New,FC_PROVIDED,139,4',
    )
    result = _run(duplicated)
    assert result.sku_count == 1
    assert result.duplicate_skus == 1


def test_duplicate_upcs_keep_the_first_row_only():
    """Two different SKUs that strip to the same UPC become one output row."""
    duplicated = FBA_EXPORT.replace(
        '198268844372-FNSKU,"Antora Jacket, XL",B0FX3FSCDN,X0057D6IM5,New,FC_PROVIDED,1,1',
        '197642130629_FNSKU,"Borealis Backpack, Navy",B0CN9VGKD5,X0053P8X57,New,FC_PROVIDED,5,1',
    )
    result = _run(duplicated)
    sheet = _sheet(result)
    assert result.sku_count == 1
    assert result.duplicate_skus == 1
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == "197642130629-FNSKU"
    assert sheet.cell(2, 3).value == 197642130629
    assert sheet.cell(3, 1).value is None


def test_non_numeric_upc_falls_back_to_text():
    text = FBA_EXPORT.replace("197642130629-FNSKU", "ABC-123-FNSKU")
    assert _sheet(_run(text))["C2"].value == "ABC-123"


def test_tab_delimited_export_is_supported():
    result = _run(FBA_EXPORT.replace(",", "\t"), filename="FBA19JHYH77Q.txt")
    assert result.shipment_id == "FBA19JHYH77Q"
    assert result.sku_count == 2


def _row(upc: str, sku: str = "", description: str = "") -> ShipmentSkuRow:
    return ShipmentSkuRow(
        sku=sku or f"{upc}-FNSKU",
        description=description or f"Item {upc}",
        upc=upc,
        fnsku=f"X{upc[-8:]}",
    )


def test_parse_fba_export_does_not_render_a_workbook():
    result = parse_fba_export("FBA19JHYH77Q.csv", FBA_EXPORT.encode("utf-8"))
    assert result.workbook_bytes == b""
    assert result.sku_count == 2
    assert [row.upc for row in result.rows] == ["197642130629", "198268844372"]


def test_dedupe_by_upc_keeps_the_first_occurrence():
    rows = [
        _row("111111111111", description="From first upload"),
        _row("222222222222"),
        _row("111111111111", description="From second upload"),
    ]
    unique = dedupe_by_upc(rows)
    assert [row.upc for row in unique] == ["111111111111", "222222222222"]
    assert unique[0].description == "From first upload"


def test_dedupe_by_upc_drops_blank_upcs():
    assert dedupe_by_upc([_row(""), _row("333333333333")]) == [_row("333333333333")]


def test_build_workbook_renders_collected_rows():
    sheet = load_workbook(io.BytesIO(build_workbook([_row("111111111111"), _row("222222222222")]))).active
    assert sheet.max_row == 3
    assert sheet.cell(2, 3).value == 111111111111
    assert sheet.cell(3, 3).value == 222222222222
    assert sheet["A1"].fill.fgColor.rgb == "FF92D050"


def test_compile_stored_rows_merges_uploads_and_drops_duplicate_upcs():
    stored = [
        {"sku": "111-FNSKU", "description": "First file", "upc": "111", "fnsku": "XA", "total_units": 2},
        {"sku": "222-FNSKU", "description": "First file B", "upc": "222", "fnsku": "XB", "total_units": 1},
        {"sku": "111-FNSKU", "description": "Second file duplicate", "upc": "111", "fnsku": "XA", "total_units": 9},
        {"sku": "333", "description": "Second file new", "upc": "333", "fnsku": "XC", "total_units": 4},
    ]
    unique, collected = compile_stored_rows(stored)
    assert collected == 4
    assert [row.upc for row in unique] == ["111", "222", "333"]
    assert unique[0].description == "First file"
    assert unique[2].description == "Second file new"


# ---- PO Import sheet ----------------------------------------------------


def _po_sheet(rows=None, *, purchase_order_number="FBA19JHYH77Q", supplier="North Face"):
    rows = parse_fba_export("x.csv", FBA_EXPORT.encode("utf-8")).rows if rows is None else rows
    workbook = load_workbook(
        io.BytesIO(
            build_po_import_workbook(
                rows, purchase_order_number=purchase_order_number, supplier=supplier
            )
        )
    )
    return workbook.active


@pytest.mark.parametrize(
    "filename, expected",
    [
        ("North Face WHRP 7.17.26 1 OF 5.csv", "North Face"),
        ("Dansko WHRP 8.1.26.xlsx", "Dansko"),
        ("north-face-whrp-7.17.26.txt", "north face"),
        ("Dansko 8.1.26 2 of 3.csv", "Dansko"),
        ("C:\\Users\\me\\Downloads\\North Face WHRP 7.17.26.csv", "North Face"),
        ("North Face.csv", "North Face"),
        # No supplier in the title: the FBA id leads, so there is nothing to take.
        ("FBA19JHYH77Q.csv", ""),
        ("", ""),
    ],
)
def test_supplier_is_the_leading_words_of_the_filename(filename, expected):
    assert supplier_from_filename(filename) == expected


def test_po_import_filename_names_the_supplier_and_purchase_order():
    assert po_import_filename("FBA19JHYH77Q", "North Face") == "PO IMPORT North Face FBA19JHYH77Q.xlsx"
    assert po_import_filename("FBA19JHYH77Q", "") == "PO IMPORT FBA19JHYH77Q.xlsx"
    assert po_import_filename("", "") == "PO IMPORT.xlsx"


def test_po_import_keeps_the_template_headers_and_banner():
    sheet = _po_sheet()
    assert sheet.title == "Purchase Order Import Template"
    assert [sheet.cell(5, col).value for col in range(1, 10)] == [
        "PurchaseOrderNumber",
        "SupplierCompanyName",
        "IssueDate",
        "PONotes",
        "ItemNumber",
        "ItemQuantity",
        "Facility",
        "ExpectedDate",
        "LineItemNotes",
    ]
    assert sheet["B1"].value.startswith("Purchase Order Import Template")
    assert sheet["A5"].font.name == "Lato Black"
    assert {str(rng) for rng in sheet.merged_cells.ranges} == {"A1:A4", "B1:I4"}
    assert round(sheet.column_dimensions["B"].width, 4) == 52.4258


def test_po_import_writes_one_line_per_sku_from_row_six():
    sheet = _po_sheet()
    assert [sheet.cell(row, 1).value for row in (6, 7)] == ["FBA19JHYH77Q", "FBA19JHYH77Q"]
    assert [sheet.cell(row, 2).value for row in (6, 7)] == ["North Face", "North Face"]
    assert [sheet.cell(row, 5).value for row in (6, 7)] == [
        "197642130629-FNSKU",
        "198268844372-FNSKU",
    ]
    assert [sheet.cell(row, 6).value for row in (6, 7)] == [139, 1]
    assert sheet.max_row == 7


def test_po_import_line_items_use_the_template_body_font():
    sheet = _po_sheet()
    assert sheet["A6"].font.name == "Open Sans"
    assert sheet["F6"].font.sz == 11


def test_po_import_fills_facility_with_whrep_ontario():
    sheet = _po_sheet()
    assert [sheet.cell(row, 7).value for row in (6, 7)] == ["WHREP Ontario", "WHREP Ontario"]
    assert sheet["G6"].font.name == "Open Sans"


def test_po_import_leaves_the_optional_columns_blank():
    sheet = _po_sheet()
    # IssueDate, PONotes, ExpectedDate, LineItemNotes. Facility is always filled.
    assert all(sheet.cell(row, col).value is None for row in (6, 7) for col in (3, 4, 8, 9))


@pytest.mark.parametrize(
    "title, expected",
    [
        ("The North Face Borealis Commuter Laptop Backpack, Blk/Blk-NPF, OSFA", "The North Face"),
        ("THE NORTH FACE Evolution Simple Dome Hoodie, Pale Gy Hth, S", "The North Face"),
        ("Smartwool Women's Classic Thermal Merino Base Layer Crew", "Smartwool"),
        ("Dansko Women's Professional Clog", "Dansko"),
        ("Oboz Bridger 7\" Insulated Waterproof Winter Boot", "Oboz"),
        ("Borealis Backpack, Blk", ""),
        ("", ""),
    ],
)
def test_supplier_is_the_brand_at_the_start_of_the_title(title, expected):
    assert supplier_from_title(title) == expected


def test_supplier_from_titles_keeps_the_first_known_brand():
    assert (
        supplier_from_titles(
            [
                "Borealis Backpack, Blk",
                "THE NORTH FACE Evolution Simple Dome Hoodie, Pale Gy Hth, S",
                "Dansko Women's Professional Clog",
            ]
        )
        == "The North Face"
    )


def test_resolve_po_supplier_prefers_the_title_over_the_filename():
    rows = [
        ShipmentSkuRow(
            sku="197642130629-FNSKU",
            description="The North Face Borealis Commuter Laptop Backpack, Blk/Blk-NPF, OSFA",
            upc="197642130629",
            fnsku="X0052Z3NU3",
            total_units=139,
        )
    ]
    assert (
        resolve_po_supplier(rows, filename="FBA19JHYH77Q.csv", shipment_name="NFA WHRP 7.17.26")
        == "The North Face"
    )


def test_resolve_po_supplier_falls_back_to_the_shipment_name_code():
    rows = [_row("197642130629", description="Borealis Backpack, Blk")]
    assert (
        resolve_po_supplier(rows, filename="FBA19JHYH77Q.csv", shipment_name="NFA WHRP 7.17.26 1 OF 5")
        == "The North Face"
    )


def test_po_import_keeps_the_instructions_sheet():
    rows = parse_fba_export("x.csv", FBA_EXPORT.encode("utf-8")).rows
    workbook = load_workbook(
        io.BytesIO(build_po_import_workbook(rows, purchase_order_number="X", supplier="Y"))
    )
    assert workbook.sheetnames == ["Purchase Order Import Template", "Instructions"]
    assert workbook["Instructions"]["A5"].value == "To prepare the file for import:"


def test_po_import_tolerates_a_missing_supplier_or_purchase_order():
    sheet = _po_sheet(purchase_order_number="", supplier="")
    assert sheet.cell(6, 1).value is None
    assert sheet.cell(6, 2).value is None
    assert sheet.cell(6, 5).value == "197642130629-FNSKU"


def test_po_import_of_no_rows_is_just_the_template():
    assert _po_sheet([]).max_row == 5


def test_missing_sku_table_is_rejected():
    with pytest.raises(ShipmentManagerError, match="Could not find the SKU table"):
        _run("Shipment ID,FBA19JHYH77Q\nBoxes,38\n")


def test_empty_file_is_rejected():
    with pytest.raises(ShipmentManagerError):
        _run("")
