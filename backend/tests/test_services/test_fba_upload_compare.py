"""Tests for FBA Upload Compare — Output reconciled to the AMZ upload file."""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.services.fba_upload_compare import (
    ADDED_UNITS_SHEET,
    BOX_CONTENTS_SHEET,
    DIMENSIONS_SHEET,
    MISSING_UPC_REASON,
    NOT_IN_AMZ_REASON,
    PIVOT_START_COL,
    REMOVED_UNITS_SHEET,
    SHORTFALL_REASON,
    SURPLUS_REASON,
    AmzSkuRow,
    ContentRow,
    ParsedAmz,
    generate_fba_upload_compare,
    reconcile_to_amz,
)

SAMPLE_OUTPUT = Path(r"c:\Users\Administrator\Downloads\FBA19Q0MTKCW Output.xlsx")
SAMPLE_AMZ = Path(
    r"c:\Users\Administrator\Downloads"
    r"\2026-09-25_13-14-35_01e45a16-ff24-4be6-aafa-d4e96ed69bfc.xlsx"
)


def _amz(*pairs: tuple[str, int]) -> ParsedAmz:
    return ParsedAmz(skus=[AmzSkuRow(sku_id=sku, expected_qty=qty) for sku, qty in pairs])


def _rows(*triples: tuple[str, int, int]) -> list[ContentRow]:
    return [ContentRow(identifier=u, box_number=b, qty=q) for u, b, q in triples]


def _totals(rows) -> dict[str, int | float]:
    totals: dict[str, int | float] = defaultdict(int)
    for row in rows:
        totals[row.identifier] += row.qty
    return dict(totals)


def test_reconcile_leaves_matching_rows_untouched():
    rows = _rows(("111", 1, 4), ("222", 2, 6))
    result = reconcile_to_amz(rows, _amz(("111", 4), ("222", 6)), {})

    assert result.rows == rows
    assert result.added == []
    assert result.removed == []
    assert result.box_numbers == [1, 2]
    assert result.last_box == 2


def test_reconcile_drops_upcs_the_amz_file_does_not_expect():
    result = reconcile_to_amz(
        _rows(("111", 1, 4), ("999", 1, 7), ("999", 2, 3)),
        _amz(("111", 4)),
        {},
    )

    assert _totals(result.rows) == {"111": 4}
    assert [(a.identifier, a.box_number, a.qty, a.reason) for a in result.removed] == [
        ("999", 1, 7, NOT_IN_AMZ_REASON),
        ("999", 2, 3, NOT_IN_AMZ_REASON),
    ]
    assert result.added == []


def test_reconcile_trims_surplus_from_the_highest_boxes_first():
    result = reconcile_to_amz(
        _rows(("111", 1, 5), ("111", 2, 4), ("111", 3, 3)),
        _amz(("111", 6)),
        {},
    )

    # 12 units down to 6: box 3 gives up 3 and box 2 gives up 3.
    assert [(r.identifier, r.box_number, r.qty) for r in result.rows] == [("111", 1, 5), ("111", 2, 1)]
    assert [(a.box_number, a.qty, a.reason) for a in result.removed] == [
        (3, 3, SURPLUS_REASON),
        (2, 3, SURPLUS_REASON),
    ]
    assert _totals(result.rows) == {"111": 6}


def test_reconcile_adds_missing_upc_to_the_last_box():
    result = reconcile_to_amz(
        _rows(("111", 1, 4), ("111", 3, 2)),
        _amz(("111", 6), ("222", 9)),
        {},
    )

    assert [(a.identifier, a.box_number, a.qty, a.reason) for a in result.added] == [
        ("222", 3, 9, MISSING_UPC_REASON)
    ]
    added_row = next(row for row in result.rows if row.identifier == "222")
    assert (added_row.box_number, added_row.qty, added_row.added) == (3, 9, True)
    assert _totals(result.rows) == {"111": 6, "222": 9}


def test_reconcile_merges_a_shortfall_into_an_existing_last_box_row():
    result = reconcile_to_amz(
        _rows(("111", 1, 4), ("111", 2, 1)),
        _amz(("111", 10)),
        {},
    )

    assert [(a.box_number, a.qty, a.reason) for a in result.added] == [(2, 5, SHORTFALL_REASON)]
    # The shortfall lands on the existing last-box row instead of duplicating it.
    assert [(r.identifier, r.box_number, r.qty, r.added) for r in result.rows] == [
        ("111", 1, 4, False),
        ("111", 2, 6, True),
    ]
    assert _totals(result.rows) == {"111": 10}


def test_reconcile_keeps_box_columns_that_empty_out():
    result = reconcile_to_amz(
        _rows(("111", 1, 4), ("999", 2, 7), ("111", 3, 1)),
        _amz(("111", 5)),
        {},
    )

    assert result.box_numbers == [1, 2, 3]
    assert 2 not in {row.box_number for row in result.rows}


def test_reconcile_matches_old_sku_identifiers_through_the_catalog():
    result = reconcile_to_amz(
        [ContentRow(identifier="9990262", box_number=1, qty=4)],
        _amz(("111", 6)),
        {"111": "9990262"},
    )

    # Resolved to the AMZ id, relabelled, and topped up rather than replaced.
    assert [(r.identifier, r.box_number, r.qty, r.remapped) for r in result.rows] == [
        ("111", 1, 6, True)
    ]
    assert result.removed == []


def test_reconcile_result_always_equals_the_amz_expected_quantities():
    amz = _amz(("111", 6), ("222", 1), ("333", 12), ("444", 3))
    result = reconcile_to_amz(
        _rows(
            ("111", 1, 20),
            ("222", 1, 1),
            ("333", 2, 4),
            ("999", 2, 8),
            ("333", 4, 1),
        ),
        amz,
        {},
    )

    assert _totals(result.rows) == {"111": 6, "222": 1, "333": 12, "444": 3}
    original = 20 + 1 + 4 + 8 + 1
    added = sum(a.qty for a in result.added)
    removed = sum(a.qty for a in result.removed)
    assert original + added - removed == sum(row.expected_qty for row in amz.skus)


@pytest.mark.skipif(
    not SAMPLE_OUTPUT.exists() or not SAMPLE_AMZ.exists(),
    reason="Sample FBA Upload Compare files not present",
)
def test_sample_result_pivot_matches_the_amz_upload_file():
    source = openpyxl.load_workbook(SAMPLE_AMZ, data_only=True)
    try:
        sheet = source["Box packing information"]
        expected: dict[str, int] = defaultdict(int)
        for row in range(6, sheet.max_row + 1):
            sku = sheet.cell(row, 1).value
            if not sku:
                continue
            text = str(sku).strip()
            if text.lower().startswith(("box ", "name of box", "box weight")):
                break
            expected[text.replace("-FNSKU", "")] += int(sheet.cell(row, 10).value)
    finally:
        source.close()

    result = generate_fba_upload_compare(
        SAMPLE_OUTPUT.read_bytes(),
        SAMPLE_AMZ.read_bytes(),
        [],
        output_filename=SAMPLE_OUTPUT.name,
        amz_filename=SAMPLE_AMZ.name,
    )
    assert result.total_qty == sum(expected.values())
    assert result.upc_count == len(expected)

    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=True)
    try:
        contents = workbook[BOX_CONTENTS_SHEET]
        assert workbook.sheetnames[:2] == [BOX_CONTENTS_SHEET, DIMENSIONS_SHEET]
        assert ADDED_UNITS_SHEET in workbook.sheetnames
        assert REMOVED_UNITS_SHEET in workbook.sheetnames

        raw: dict[str, int] = defaultdict(int)
        per_box: dict[tuple[str, int], int] = defaultdict(int)
        for row in range(2, contents.max_row + 1):
            upc = contents.cell(row, 1).value
            if upc in (None, ""):
                continue
            qty = contents.cell(row, 3).value
            box = contents.cell(row, 2).value
            raw[str(upc)] += qty
            per_box[(str(upc), box)] += qty
        assert raw == expected

        box_columns: dict[int, int] = {}
        for col in range(PIVOT_START_COL + 1, contents.max_column + 1):
            header = contents.cell(2, col).value
            if not isinstance(header, int):
                break
            box_columns[header] = col
        assert sorted(box_columns) == list(range(1, result.box_count + 1))
        grand_col = max(box_columns.values()) + 1
        assert contents.cell(2, grand_col).value == "Grand Total"

        pivot: dict[str, int] = {}
        row = 3
        while True:
            label = contents.cell(row, PIVOT_START_COL).value
            if label in (None, "") or label == "Grand Total":
                break
            label = str(label)
            for box, col in box_columns.items():
                assert (contents.cell(row, col).value or 0) == per_box.get((label, box), 0)
            pivot[label] = contents.cell(row, grand_col).value
            row += 1
        assert pivot == expected
        assert contents.cell(row, grand_col).value == sum(expected.values())

        added_sheet = workbook[ADDED_UNITS_SHEET]
        assert [added_sheet.cell(1, col).value for col in range(1, 4)] == ["UPC", "Box Number", "QTY"]
        added_units = 0
        seen: set[str] = set()
        for r in range(2, added_sheet.max_row + 1):
            upc = added_sheet.cell(r, 1).value
            if upc in (None, ""):
                continue
            assert added_sheet.cell(r, 2).value == result.last_box
            assert str(upc) not in seen
            seen.add(str(upc))
            added_units += added_sheet.cell(r, 3).value
        assert added_units == result.added_qty

        removed_sheet = workbook[REMOVED_UNITS_SHEET]
        assert [removed_sheet.cell(1, col).value for col in range(1, 5)] == [
            "UPC",
            "Box Number",
            "QTY",
            "Reason",
        ]
    finally:
        workbook.close()
