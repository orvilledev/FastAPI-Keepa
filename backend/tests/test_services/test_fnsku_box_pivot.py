"""Tests for FNSKU + BOX# → scanned data + box pivot conversion."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.services.fnsku_box_pivot import (
    DEFAULT_OUTPUT_FILENAME,
    OUTPUT_FILENAME,
    PIVOT_SHEET_NAME,
    SCANNED_SHEET_NAME,
    TEMPLATE_FILENAME,
    FnskuBoxPivotError,
    ScanRow,
    apply_msku_lookup,
    build_template_workbook,
    generate_fnsku_box_pivot,
    parse_fnsku_box_rows,
    sanitize_download_filename,
)

SAMPLE_INPUT = Path(r"c:\Users\Administrator\Downloads\Input.xlsx")
SAMPLE_OUTPUT = Path(r"c:\Users\Administrator\Downloads\Output.xlsx")
# Older sample workbook used the previous tab names.
_SAMPLE_PIVOT_SHEET = "Sheet7"
_SAMPLE_SCANNED_SHEET = "scanned data"


def _xlsx_with_rows(headers: list[str], rows: list[list[object]], sheet_name: str = "Sheet1") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parse_requires_fnsku_and_box_headers():
    raw = _xlsx_with_rows(["SKU", "QTY"], [["X001", 1]])
    with pytest.raises(FnskuBoxPivotError, match="FNSKU"):
        parse_fnsku_box_rows(raw)


def test_parse_skips_blank_fnsku_and_coerces_box_numbers():
    raw = _xlsx_with_rows(
        ["FNSKU", "BOX#"],
        [["X001", 1], ["", 2], ["X002", "3.0"], ["X003", None]],
    )
    rows = parse_fnsku_box_rows(raw)
    assert rows == [
        ScanRow(fnsku="X001", box=1),
        ScanRow(fnsku="X002", box=3),
        ScanRow(fnsku="X003", box=None),
    ]


def test_generate_matches_excel_pivot_layout():
    raw = _xlsx_with_rows(
        ["FNSKU", "BOX#"],
        [
            ["XA", 2],
            ["XA", 2],
            ["XB", 1],
            ["XC", 1],
            ["XD", None],
        ],
    )
    catalog = {
        "XA": {"fnsku": "XA", "sku": "9990001", "upc": "111"},
        "XB": {"fnsku": "XB", "sku": "", "upc": "190850809165"},
        "XC": {"fnsku": "XC", "sku": "190850809165-FNSKU", "upc": "190850809165"},
    }
    result = generate_fnsku_box_pivot(raw, catalog, filename="My Scan Sheet.xlsx")
    assert result.filename == "My Scan Sheet.xlsx"
    assert result.row_count == 5
    assert result.sku_count == 2
    assert result.unmatched_count == 1
    assert result.unmatched_fnskus == ("XD",)

    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=True)
    try:
        assert workbook.sheetnames == [PIVOT_SHEET_NAME, SCANNED_SHEET_NAME]
        scanned = workbook[SCANNED_SHEET_NAME]
        assert [scanned.cell(1, col).value for col in range(1, 5)] == ["msku", "FNSKU", "BOX#", "QTY"]
        assert [scanned.cell(2, col).value for col in range(1, 5)] == ["9990001-FNSKU", "XA", 2, 1]
        assert [scanned.cell(3, col).value for col in range(1, 5)] == ["9990001-FNSKU", "XA", 2, 1]
        assert [scanned.cell(4, col).value for col in range(1, 5)] == ["190850809165-FNSKU", "XB", 1, 1]
        assert [scanned.cell(5, col).value for col in range(1, 5)] == ["190850809165-FNSKU", "XC", 1, 1]
        assert [scanned.cell(6, col).value for col in range(1, 5)] == [None, "XD", None, 1]

        pivot = workbook[PIVOT_SHEET_NAME]
        assert pivot["A3"].value == "Sum of QTY"
        assert pivot["B3"].value == "Column Labels"
        assert pivot["A4"].value == "Row Labels"
        assert pivot["B4"].value == 1
        assert pivot["C4"].value == 2
        assert pivot["D4"].value == "Grand Total"
        assert pivot["E4"].value is None
        assert pivot["A5"].value == "190850809165-FNSKU"
        assert pivot["B5"].value == 2
        assert pivot["C5"].value is None
        assert pivot["D5"].value == 2
        assert pivot["A6"].value == "9990001-FNSKU"
        assert pivot["C6"].value == 2
        assert pivot["D6"].value == 2
        assert pivot["A7"].value == "Grand Total"
        assert pivot["B7"].value == 2
        assert pivot["C7"].value == 2
        assert pivot["D7"].value == 4
        assert all(
            cell.value != "(blank)"
            for row in pivot.iter_rows(min_row=3, max_row=7, max_col=4)
            for cell in row
        )
        assert abs((pivot.column_dimensions["B"].width or 0) - (35 - 5) / 7) < 0.01
        assert abs((pivot.column_dimensions["C"].width or 0) - (35 - 5) / 7) < 0.01
        assert abs((pivot.column_dimensions["D"].width or 0) - (35 - 5) / 7) < 0.01
    finally:
        workbook.close()


def test_output_keeps_ids_as_text_and_qty_box_as_numbers():
    raw = _xlsx_with_rows(
        ["FNSKU", "BOX#"],
        [["XA", "2"], ["XB", 3.0]],
    )
    catalog = {
        "XA": {"fnsku": "XA", "sku": "9990001", "upc": "111"},
        "XB": {"fnsku": "XB", "sku": "", "upc": "190850809165"},
    }
    result = generate_fnsku_box_pivot(raw, catalog)
    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes))
    try:
        scanned = workbook[SCANNED_SHEET_NAME]
        assert scanned["A2"].value == "9990001-FNSKU"
        assert scanned["A2"].data_type == "s"
        assert scanned["A2"].number_format == "@"
        assert scanned["B2"].value == "XA"
        assert scanned["B2"].data_type == "s"
        assert scanned["B2"].number_format == "@"
        assert scanned["C2"].value == 2
        assert isinstance(scanned["C2"].value, int)
        assert scanned["C2"].data_type == "n"
        assert scanned["D2"].value == 1
        assert isinstance(scanned["D2"].value, int)
        assert scanned["D2"].data_type == "n"
        assert scanned["C3"].value == 3
        assert scanned["C3"].data_type == "n"

        pivot = workbook[PIVOT_SHEET_NAME]
        assert pivot["A5"].data_type == "s"
        assert pivot["A5"].number_format == "@"
        assert pivot["B4"].value == 2
        assert pivot["B4"].data_type == "n"
        assert isinstance(pivot["B4"].value, int)
        assert pivot["C4"].value == 3
        assert pivot["C4"].data_type == "n"
        assert pivot["C5"].value == 1
        assert pivot["C5"].data_type == "n"
        assert isinstance(pivot["C5"].value, int)
        assert pivot["D5"].data_type == "n"
        assert pivot["A7"].data_type == "s"
        assert scanned["A1"].font.bold is True
        assert scanned["A1"].fill.fgColor.rgb in ("00404040", "404040")
        assert pivot["A4"].font.bold is True
        assert pivot["A4"].fill.fgColor.rgb in ("00404040", "404040")
        assert pivot["B4"].font.bold is True
    finally:
        workbook.close()


def test_sanitize_download_filename_keeps_upload_basename():
    assert sanitize_download_filename("Input.xlsx") == "Input.xlsx"
    assert sanitize_download_filename(r"C:\temp\Input.xlsx") == "Input.xlsx"
    assert sanitize_download_filename("") == DEFAULT_OUTPUT_FILENAME
    assert sanitize_download_filename(None) == DEFAULT_OUTPUT_FILENAME
    assert OUTPUT_FILENAME == DEFAULT_OUTPUT_FILENAME
    assert TEMPLATE_FILENAME == "FNSKU Pack Station Template.xlsx"


def test_template_has_expected_headers():
    workbook = openpyxl.load_workbook(BytesIO(build_template_workbook()))
    try:
        sheet = workbook.active
        assert sheet["A1"].value == "FNSKU"
        assert sheet["B1"].value == "BOX#"
        assert sheet["A1"].font.bold is True
        assert sheet["B1"].font.bold is True
    finally:
        workbook.close()


def test_apply_msku_lookup_reports_each_unmatched_fnsku_once():
    rows = [
        ScanRow(fnsku="XA", box=1),
        ScanRow(fnsku="MISSING", box=1),
        ScanRow(fnsku="MISSING", box=2),
    ]
    scanned, unmatched = apply_msku_lookup(rows, {"XA": {"sku": "9991", "upc": "1"}})
    assert unmatched == ("MISSING",)
    assert scanned[0].msku == "9991-FNSKU"
    assert scanned[1].msku == ""


def _pivot_map(sheet) -> dict[tuple[object, object], object]:
    headers = {}
    for col in range(2, sheet.max_column + 1):
        header = sheet.cell(4, col).value
        if header is not None and header != "(blank)":
            headers[col] = header
    mapped: dict[tuple[object, object], object] = {}
    for row in range(5, sheet.max_row + 1):
        label = sheet.cell(row, 1).value
        if not label or label == "(blank)":
            continue
        for col, header in headers.items():
            value = sheet.cell(row, col).value
            if value not in (None, ""):
                mapped[(label, header)] = value
    return mapped


@pytest.mark.skipif(not SAMPLE_INPUT.exists() or not SAMPLE_OUTPUT.exists(), reason="Sample Input/Output.xlsx not present")
def test_sample_input_matches_sample_output():
    sample = openpyxl.load_workbook(SAMPLE_OUTPUT, data_only=True)
    try:
        scanned = sample[_SAMPLE_SCANNED_SHEET]
        catalog = {}
        for row in range(2, scanned.max_row + 1):
            msku = scanned.cell(row, 1).value
            fnsku = scanned.cell(row, 2).value
            if not fnsku or fnsku in catalog:
                continue
            msku_text = str(msku or "")
            if msku_text.endswith("-FNSKU") and sum(ch.isdigit() for ch in msku_text) <= 7:
                catalog[str(fnsku)] = {"sku": msku_text[: -len("-FNSKU")], "upc": "", "fnsku": str(fnsku)}
            else:
                upc = msku_text[: -len("-FNSKU")] if msku_text.endswith("-FNSKU") else msku_text
                catalog[str(fnsku)] = {"sku": "", "upc": upc, "fnsku": str(fnsku)}
        expected_scan = [
            (
                scanned.cell(row, 1).value,
                scanned.cell(row, 2).value,
                scanned.cell(row, 3).value,
                scanned.cell(row, 4).value,
            )
            for row in range(1, scanned.max_row + 1)
        ]
        expected_pivot_map = _pivot_map(sample[_SAMPLE_PIVOT_SHEET])
    finally:
        sample.close()

    result = generate_fnsku_box_pivot(SAMPLE_INPUT.read_bytes(), catalog)
    got = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=True)
    try:
        got_scan = got[SCANNED_SHEET_NAME]
        got_rows = [
            (
                got_scan.cell(row, 1).value,
                got_scan.cell(row, 2).value,
                got_scan.cell(row, 3).value,
                got_scan.cell(row, 4).value,
            )
            for row in range(1, got_scan.max_row + 1)
        ]
        assert got_rows == expected_scan

        got_pivot = got[PIVOT_SHEET_NAME]
        assert _pivot_map(got_pivot) == expected_pivot_map
        assert got_pivot["A3"].value == "Sum of QTY"
        assert got_pivot["A4"].value == "Row Labels"
        last_col = get_column_letter(got_pivot.max_column)
        assert got_pivot[f"{last_col}4"].value == "Grand Total"
        assert result.row_count == 1665
        assert result.sku_count == 270
        assert result.unmatched_count == 0
    finally:
        got.close()
