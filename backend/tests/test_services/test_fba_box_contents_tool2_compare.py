"""Tests for NFA and SMW Output vs manifest compare / correct."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.services.fba_box_contents_tool2_compare import (
    BOX_CONTENTS_SHEET,
    DIMENSIONS_SHEET,
    generate_tool2_manifest_compare,
    sanitize_download_filename,
)

KIT = Path(r"C:\Users\Administrator\Downloads\Testing Kit")
SAMPLE_OUTPUT = KIT / "FBA19RL0N7D1 Output.xlsx"
SAMPLE_MANIFEST = KIT / "Manifest File.xlsx"
SAMPLE_CORRECTED = KIT / "FBA19RL0N7D1 - Corrected.xlsx"
YELLOW = "FFFFFF00"


def _output_bytes(rows, dims=((" 10.5", 20, 15, 10),)) -> bytes:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = BOX_CONTENTS_SHEET
    sheet.append(["UPC", "Box Number", "QTY"])
    for upc, box, qty in rows:
        sheet.append([upc, box, qty])
    dim = wb.create_sheet(DIMENSIONS_SHEET)
    dim.append(["Weight", "Length", "Width", "Height"])
    for row in dims:
        dim.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _manifest_bytes(skus, shipment="FBA1TEST", boxes=2) -> bytes:
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Box packing information"
    sheet["A2"] = f"Shipment: {shipment}"
    sheet["I3"] = "Total box count:"
    sheet["M3"] = boxes
    sheet.append([])
    sheet.append(["SKU", "Product title ", "Id", "ASIN", "FNSKU", "Condition", "Prep type",
                  "Who preps units?", "Who labels units?", "Expected quantity", "Boxed quantity"])
    for sku, qty in skus:
        sheet.append([sku, "", "", "", "", "", "", "", "", qty, 0])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rows(sheet):
    return [
        (sheet.cell(r, 1).value, sheet.cell(r, 2).value, sheet.cell(r, 3).value)
        for r in range(2, sheet.max_row + 1)
        if sheet.cell(r, 1).value is not None
    ]


def test_sanitize_download_filename():
    assert sanitize_download_filename("FBA19RL0N7D1") == "FBA19RL0N7D1 - Corrected.xlsx"
    assert sanitize_download_filename(None) == "Box Contents - Corrected.xlsx"


def test_shortfalls_append_to_last_box_and_keep_rows_untouched():
    output = _output_bytes([("111", 1, 2), ("222", 1, 0), ("222", 1, 1), ("333", 2, 1)])
    manifest = _manifest_bytes([("111", 2), ("222", 3), ("333", 1), ("444", 1)])

    result = generate_tool2_manifest_compare(output, manifest, [], output_filename="X Output.xlsx")

    assert result.filename == "FBA1TEST - Corrected.xlsx"
    assert result.total_qty == 7
    assert result.added_count == 2 and result.added_qty == 3
    assert result.removed_count == 0

    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[BOX_CONTENTS_SHEET]
    # Original rows (including the zero-qty line) stay as packed; additions are new rows.
    assert _rows(sheet) == [
        ("111", 1, 2), ("222", 1, 0), ("222", 1, 1), ("333", 2, 1),
        ("222", 2, 2), ("444", 2, 1),
    ]
    for r in (2, 3, 4, 5):
        assert sheet.cell(r, 1).fill.fill_type is None
    for r in (6, 7):
        for c in (1, 2, 3):
            assert sheet.cell(r, c).fill.fgColor.rgb == YELLOW
        assert sheet.cell(r, 1).data_type == "s"

    assert [sheet.cell(1, c).value for c in (5, 6, 7)] == ["FBA1TEST", "Box #", "Qty"]
    assert all(sheet.cell(1, c).font.b for c in (5, 6, 7))
    assert [(sheet.cell(r, 5).value, sheet.cell(r, 6).value, sheet.cell(r, 7).value) for r in (2, 3)] == [
        ("222", 2, 2), ("444", 2, 1),
    ]
    assert sheet["I1"].value == "Sum of QTY"
    assert sheet["J1"].value == "Column Labels"


def test_no_adjustments_keeps_pivot_at_column_g():
    output = _output_bytes([("111", 1, 2), ("333", 2, 1)])
    manifest = _manifest_bytes([("111", 2), ("333", 1)])

    result = generate_tool2_manifest_compare(output, manifest, [])

    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[BOX_CONTENTS_SHEET]
    assert result.added_count == 0
    assert sheet["E1"].value is None
    assert sheet["G1"].value == "Sum of QTY"


def test_surplus_and_unexpected_units_are_removed():
    output = _output_bytes([("111", 1, 3), ("999", 1, 1), ("333", 2, 2)])
    manifest = _manifest_bytes([("111", 2), ("333", 2)])

    result = generate_tool2_manifest_compare(output, manifest, [])

    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[BOX_CONTENTS_SHEET]
    assert _rows(sheet) == [("111", 1, 2), ("333", 2, 2)]
    assert result.removed_qty == 2
    assert result.total_qty == 4


def test_old_sku_remap_is_highlighted():
    output = _output_bytes([("673088034429", 1, 4), ("333", 2, 1)])
    manifest = _manifest_bytes([("50504", 4), ("333", 1)])
    catalog = [{"old_sku": "50504", "upc_code": "673088034429"}]

    result = generate_tool2_manifest_compare(output, manifest, catalog)

    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[BOX_CONTENTS_SHEET]
    assert sheet["A2"].value == "50504"
    assert sheet["A2"].fill.fgColor.rgb == YELLOW
    assert result.remapped_count == 1


def test_dimensions_copied_unchanged():
    output = _output_bytes([("111", 1, 1)], dims=((" 14.85", 23.37, 16.37, 12.75),))
    manifest = _manifest_bytes([("111", 1)], boxes=1)

    result = generate_tool2_manifest_compare(output, manifest, [])

    dim = openpyxl.load_workbook(BytesIO(result.file_bytes))[DIMENSIONS_SHEET]
    # Old-layout Output in (no Box Number) -> Box Number layout out.
    assert [dim.cell(1, c).value for c in range(1, 6)] == [
        "Box Number", "Weight", "Length", "Width", "Height",
    ]
    assert [dim.cell(2, c).value for c in range(1, 6)] == [1, " 14.85", 23.37, 16.37, 12.75]
    assert dim["B2"].data_type == "s"
    assert dim.max_column == 5
    assert result.zero_corrected_count == 0


def _output_with_box_numbers(rows, dims) -> bytes:
    """Output in the new layout: Box Number first, optional Zero Correction last."""
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = BOX_CONTENTS_SHEET
    sheet.append(["UPC", "Box Number", "QTY"])
    for upc, box, qty in rows:
        sheet.append([upc, box, qty])
    dim = wb.create_sheet(DIMENSIONS_SHEET)
    dim.append(["Box Number", "Weight", "Length", "Width", "Height", "Zero Correction"])
    for row in dims:
        dim.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_zero_weight_is_suggested_with_reason_and_highlight():
    output = _output_bytes(
        [("111", 1, 60), ("111", 2, 60), ("111", 3, 60)],
        dims=(
            (" 14.00", 20, 16, 12),
            (" 0.00", 20.1, 16, 12.1),
            (" 15.00", 25, 18, 14),
        ),
    )
    manifest = _manifest_bytes([("111", 180)], boxes=3)

    result = generate_tool2_manifest_compare(output, manifest, [])

    assert result.zero_corrected_count == 1
    dim = openpyxl.load_workbook(BytesIO(result.file_bytes))[DIMENSIONS_SHEET]
    assert [dim.cell(1, c).value for c in range(1, 7)] == [
        "Box Number", "Weight", "Length", "Width", "Height", "Zero Correction",
    ]
    assert dim["A3"].value == 2
    assert dim["B3"].value == " 14.00"  # closest identical carton is Box 1
    assert "Box 1" in dim["F3"].value
    assert all(dim.cell(3, c).fill.fgColor.rgb == YELLOW for c in range(2, 7))
    for r in (2, 4):
        assert dim.cell(r, 6).value is None
        assert dim.cell(r, 2).fill.fill_type is None


def test_existing_zero_correction_is_carried_over_unchanged():
    output = _output_with_box_numbers(
        [("111", 1, 60), ("111", 2, 60)],
        dims=(
            (1, " 14.00", 20, 16, 12, None),
            (2, " 15.00", 20.1, 16, 12.1, "Similar to Box 1"),
        ),
    )
    manifest = _manifest_bytes([("111", 120)], boxes=2)

    result = generate_tool2_manifest_compare(output, manifest, [])

    assert result.zero_corrected_count == 0  # nothing new to suggest
    dim = openpyxl.load_workbook(BytesIO(result.file_bytes))[DIMENSIONS_SHEET]
    assert [dim.cell(3, c).value for c in range(1, 7)] == [
        2, " 15.00", 20.1, 16, 12.1, "Similar to Box 1",
    ]
    assert dim["B3"].fill.fgColor.rgb == YELLOW


@pytest.mark.skipif(
    not (SAMPLE_OUTPUT.exists() and SAMPLE_MANIFEST.exists() and SAMPLE_CORRECTED.exists()),
    reason="Testing Kit sample files missing",
)
def test_matches_hand_corrected_sample():
    result = generate_tool2_manifest_compare(
        SAMPLE_OUTPUT.read_bytes(),
        SAMPLE_MANIFEST.read_bytes(),
        [],
        output_filename=SAMPLE_OUTPUT.name,
    )
    assert result.filename == "FBA19RL0N7D1 - Corrected.xlsx"
    assert result.total_qty == 849
    assert result.added_count == 2
    assert result.last_box == 98

    got = openpyxl.load_workbook(BytesIO(result.file_bytes))
    expected = openpyxl.load_workbook(SAMPLE_CORRECTED)
    assert got.sheetnames == expected.sheetnames

    g, e = got[BOX_CONTENTS_SHEET], expected[BOX_CONTENTS_SHEET]
    assert g.max_row == e.max_row
    for r in range(1, e.max_row + 1):
        for c in range(1, 8):
            gc, ec = g.cell(r, c), e.cell(r, c)
            assert (None if gc.value is None else str(gc.value)) == (
                None if ec.value is None else str(ec.value)
            ), f"{gc.coordinate}: {gc.value!r} != {ec.value!r}"
            exp_fill = ec.fill.fgColor.rgb if ec.fill.fill_type else None
            got_fill = gc.fill.fgColor.rgb if gc.fill.fill_type else None
            assert got_fill == exp_fill, gc.coordinate

    gd, ed = got[DIMENSIONS_SHEET], expected[DIMENSIONS_SHEET]
    # The hand-corrected sample has no Box Number column; ours leads with one.
    for r in range(1, ed.max_row + 1):
        for c in range(1, 5):
            assert gd.cell(r, c + 1).value == ed.cell(r, c).value
