"""Tests for Ship To Address catalog spreadsheet parsing and template shape."""
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.services.catalog_ship_to_headers import HEADERS, TEMPLATE_FILENAME
from app.services.catalog_ship_to_import import (
    compose_full_address,
    parse_ship_to_spreadsheet,
    ship_to_row_to_record,
)

_TEMPLATE = (
    Path(__file__).resolve().parents[2]
    / "app"
    / "static"
    / "catalog_templates"
    / TEMPLATE_FILENAME
)


def _workbook_bytes(rows: list[list[object]], sheet_name: str = "SHIP TO ") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_ship_to_spreadsheet_valid_rows():
    content = _workbook_bytes(
        [
            list(HEADERS),
            ["ABE1", "=CONCAT(C2)", "6370 Hedgewood Dr", "Allentown", "PA", "18106-9266"],
            ["ABE2", "", "705 Boulder Drive", "Breinigsville", "PA", "18031"],
        ]
    )
    rows, invalid = parse_ship_to_spreadsheet("ship-to.xlsx", content)
    assert invalid == 0
    assert len(rows) == 2
    assert rows[0]["Code"] == "ABE1"
    assert rows[0]["Full Address"] == "6370 Hedgewood Dr, Allentown, PA, 18106-9266"
    assert rows[1]["City"] == "Breinigsville"


def test_parse_ship_to_spreadsheet_skips_rows_without_code():
    content = _workbook_bytes(
        [
            list(HEADERS),
            ["ABE1", "", "6370 Hedgewood Dr", "Allentown", "PA", "18106-9266"],
            ["", "", "Missing code", "Allentown", "PA", "18106"],
        ]
    )
    rows, invalid = parse_ship_to_spreadsheet("ship-to.xlsx", content)
    assert len(rows) == 1
    assert invalid == 1


def test_parse_ship_to_spreadsheet_requires_code_header():
    content = _workbook_bytes(
        [
            ["Full Address", "Address 1", "City", "State", "Postal Code"],
            ["x", "1 Main", "Town", "PA", "18000"],
        ]
    )
    try:
        parse_ship_to_spreadsheet("ship-to.xlsx", content)
        assert False, "expected ValueError"
    except ValueError as exc:
        assert "Code" in str(exc)


def test_ship_to_row_to_record_and_compose():
    row = {
        "Code": "ABE1",
        "Full Address": compose_full_address("6370 Hedgewood Dr", "Allentown", "PA", "18106-9266"),
        "Address 1": "6370 Hedgewood Dr",
        "City": "Allentown",
        "State": "PA",
        "Postal Code": "18106-9266",
    }
    record = ship_to_row_to_record(row)
    assert record["code"] == "ABE1"
    assert record["postal_code"] == "18106-9266"
    assert record["row_data"]["Full Address"].startswith("6370 Hedgewood Dr")


def test_template_is_original_format_with_header_and_first_entry_only():
    assert _TEMPLATE.is_file()
    wb = load_workbook(_TEMPLATE)
    ws = wb.active
    assert ws.title.strip().upper() == "SHIP TO"
    assert ws.max_row == 2
    assert ws.max_column == 6
    assert [ws.cell(1, c).value for c in range(1, 7)] == list(HEADERS)
    assert ws["A2"].value == "ABE1"
    assert ws["B2"].value == '=_xlfn.CONCAT(C2&", ", D2&", ", E2&", ", F2)'
    assert ws["C2"].value == "6370 Hedgewood Dr"
    assert ws["D2"].value == "Allentown"
    assert ws["E2"].value == "PA"
    assert ws["F2"].value == "18106-9266"
    assert ws.column_dimensions["A"].width == 18.5546875
    assert ws.column_dimensions["C"].width == 78.0
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.name == "Calibri"
    assert ws.row_dimensions[1].height == 19.2
