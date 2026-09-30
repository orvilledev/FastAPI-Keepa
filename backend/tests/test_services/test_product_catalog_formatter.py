"""Product Catalog Formatter — shipment-plan rows become unique catalog UPCs."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.services.product_catalog_formatter import (
    OUTPUT_HEADERS,
    SHEET_NAME,
    ProductCatalogFormatterError,
    extract_barcode,
    format_catalogs,
    format_wr_sku_catalogs,
)

SAMPLE_INPUT = Path(
    r"c:\Users\Administrator\Downloads\Input file product catalog formatter.xlsx"
)
SAMPLE_OUTPUT = Path(r"c:\Users\Administrator\Downloads\Product Catalog Formatter.xlsx")
SAMPLE_WR_INPUT = Path(
    r"c:\Users\Administrator\Downloads\NFA_WHRP_9.14.26_CLUSTERED_WR_SKU_UPDATE.xlsx"
)
SAMPLE_WR_OUTPUT = Path(r"c:\Users\Administrator\Downloads\Product Catalog Formatter (1).xlsx")


def _plan(rows: list[tuple[str, str, str, str, str]], *, extra_footer: bool = True) -> bytes:
    """A shipment-plan workbook shaped like the Amazon export."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "Shipment number"
    sheet["B1"] = 4
    sheet["A2"] = "Workflow name"
    sheet["B2"] = "wf-test"
    sheet["A3"] = "SKUs"
    sheet["B3"] = len(rows)
    sheet["A4"] = "Units"
    sheet["B4"] = 2
    headers = [
        "SKU",
        "Title",
        "ASIN",
        "FNSKU",
        "UPC/EAN/ISBN/JAN/CODABAR",
        "Condition",
        "Prep type",
        "Quantity",
    ]
    for col, header in enumerate(headers, start=1):
        sheet.cell(6, col, header)
    for offset, (sku, title, fnsku, barcode, condition) in enumerate(rows):
        row = 7 + offset
        sheet.cell(row, 1, sku)
        sheet.cell(row, 2, title)
        sheet.cell(row, 3, "B0TEST")
        sheet.cell(row, 4, fnsku)
        sheet.cell(row, 5, barcode)
        sheet.cell(row, 6, condition)
        sheet.cell(row, 7, "FC_PROVIDED")
        sheet.cell(row, 8, 2)
    if extra_footer:
        start = 8 + len(rows) + 4
        for label in (
            "Box ID",
            "Box name",
            "Box weight (lb):",
            "Box length (inch):",
            "Box width (inch):",
            "Box height (inch):",
        ):
            sheet.cell(start, 10, label)
            start += 1
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _catalog_rows(file_bytes: bytes) -> list[tuple]:
    workbook = openpyxl.load_workbook(BytesIO(file_bytes))
    assert workbook.sheetnames == [SHEET_NAME]
    return list(workbook[SHEET_NAME].iter_rows(values_only=True))


def test_extract_barcode_prefers_upc_and_keeps_leading_zeros():
    assert extract_barcode("UPC:198268465935") == "198268465935"
    assert extract_barcode("EAN:0198268465935") == "0198268465935"
    assert extract_barcode("EAN:1111111111111 UPC:222222222222") == "222222222222"
    assert extract_barcode("ISBN:0-306-40615-2") == "0306406152"
    assert extract_barcode(198268465935) == "198268465935"
    assert extract_barcode("") == ""
    assert extract_barcode("Box ID") == ""


def test_sample_shape_keeps_sku_and_strips_the_upc_prefix():
    result = format_catalogs(
        [
            (
                "input.xlsx",
                _plan(
                    [
                        (
                            "198268465935-FNSKU",
                            "Smartwool Womens Classic Thermal Merino Base Layer Crew Dark Teal Heather Large",
                            "X0058EHT67",
                            "UPC:198268465935",
                            "New",
                        )
                    ]
                ),
            )
        ]
    )
    rows = _catalog_rows(result.file_bytes)
    assert rows[0] == OUTPUT_HEADERS
    assert rows[1] == (
        "198268465935",
        "198268465935-FNSKU",
        "X0058EHT67",
        "Smartwool Womens Classic Thermal Merino Base Layer Crew Dark Teal Heather Large",
        "New",
    )
    assert result.row_count == 1
    assert result.duplicates_removed == 0
    assert result.skipped_rows == 0


def test_multiple_files_keep_the_first_copy_of_each_upc():
    first = _plan(
        [
            ("111111111111-FNSKU", "First title", "X00AAA", "UPC:111111111111", "New"),
            ("222222222222-FNSKU", "Second title", "X00BBB", "EAN:2222222222222", "New"),
        ]
    )
    second = _plan(
        [
            ("111111111111-FNSKU", "Later title should be dropped", "X00ZZZ", "UPC:111111111111", "Used"),
            ("333333333333-FNSKU", "Third title", "X00CCC", "UPC:019826846593", "New"),
        ]
    )
    result = format_catalogs([("a.xlsx", first), ("b.xlsx", second)])
    rows = _catalog_rows(result.file_bytes)
    assert [row[0] for row in rows[1:]] == ["111111111111", "2222222222222", "019826846593"]
    assert rows[1][3] == "First title"
    assert rows[1][2] == "X00AAA"
    assert rows[1][4] == "New"
    assert rows[3][0] == "019826846593"
    assert rows[3][1] == "333333333333-FNSKU"
    assert result.file_count == 2
    assert result.source_rows == 4
    assert result.row_count == 3
    assert result.duplicates_removed == 1


def test_upc_identifiers_stay_text_in_the_workbook():
    result = format_catalogs(
        [("input.xlsx", _plan([("019826846593-FNSKU", "Crew", "X1", "UPC:019826846593", "New")]))]
    )
    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[SHEET_NAME]
    cell = sheet["A2"]
    assert cell.value == "019826846593"
    assert cell.data_type == "s"
    assert cell.number_format == "@"


def test_rows_without_a_barcode_are_skipped_and_not_invented():
    result = format_catalogs(
        [
            (
                "input.xlsx",
                _plan(
                    [
                        ("SELLER-SKU", "No barcode", "X1", "", "New"),
                        ("444444444444-FNSKU", "Has digits in the SKU", "X2", "", "New"),
                    ]
                ),
            )
        ]
    )
    rows = _catalog_rows(result.file_bytes)
    assert len(rows) == 2
    assert rows[1][:3] == ("444444444444", "444444444444-FNSKU", "X2")
    assert result.skipped_rows == 1


def test_missing_header_names_the_file():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Hello"
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(ProductCatalogFormatterError, match="plain.xlsx"):
        format_catalogs([("plain.xlsx", buffer.getvalue())])


def test_scientific_notation_is_rejected():
    with pytest.raises(ProductCatalogFormatterError, match="scientific notation"):
        extract_barcode("1.98268E+11")


@pytest.mark.skipif(
    not SAMPLE_INPUT.exists() or not SAMPLE_OUTPUT.exists(),
    reason="Sample catalog formatter workbooks are not present",
)
def test_real_sample_matches_the_catalog_workbook():
    result = format_catalogs([("input.xlsx", SAMPLE_INPUT.read_bytes())])
    expected = openpyxl.load_workbook(SAMPLE_OUTPUT, data_only=True)
    got = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=True)
    assert got.sheetnames == expected.sheetnames == [SHEET_NAME]
    got_rows = list(got[SHEET_NAME].iter_rows(min_row=1, max_col=5, values_only=True))
    expected_rows = list(expected[SHEET_NAME].iter_rows(min_row=1, max_col=5, values_only=True))
    assert got_rows == expected_rows
    assert result.row_count == 1
    assert result.duplicates_removed == 0


def _wr_file(rows: list[tuple[str, str, object, str]]) -> bytes:
    """A WR SKU Update workbook: SKU, Description, UPC, FNSKU, plus unused dimensions."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    headers = [
        "SKU",
        "Description",
        "UPC",
        "FNSKU",
        "Item Length (in.)",
        "Item Width (in.)",
        "Item Height (in.)",
        "Item Weight (lbs)",
        "Carton Length (in.)",
        "Carton Width (in.)",
        "Carton Height (in.)",
        "Carton Weight (lbs)",
        "Units per carton",
    ]
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
    for offset, (sku, description, upc, fnsku) in enumerate(rows):
        row = 2 + offset
        sheet.cell(row, 1, sku)
        sheet.cell(row, 2, description)
        sheet.cell(row, 3, upc)
        sheet.cell(row, 4, fnsku)
        sheet.cell(row, 8, 0.6)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_wr_sku_update_maps_to_the_catalog_and_sets_condition_new():
    description = "The North Face Women''s Aconcagua Jacket, Smoked Pearl, Large"
    result = format_wr_sku_catalogs(
        [
            (
                "update.xlsx",
                _wr_file(
                    [
                        ("198268878087-FNSKU", description, 198268878087, "X005957U4B"),
                    ]
                ),
            )
        ]
    )
    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[SHEET_NAME]
    assert [sheet.cell(1, col).value for col in range(1, 6)] == list(OUTPUT_HEADERS)
    assert sheet["A2"].value == 198268878087
    assert sheet["A2"].number_format == "0"
    assert sheet["B2"].value == "198268878087-FNSKU"
    assert sheet["C2"].value == "X005957U4B"
    assert sheet["D2"].value == description
    assert sheet["E2"].value == "New"
    assert sheet.max_column == 5


def test_wr_files_keep_the_first_copy_of_each_upc():
    first = _wr_file(
        [
            ("111111111111-FNSKU", "First title", 111111111111, "X00AAA"),
            ("222222222222-FNSKU", "Second title", "222222222222", "X00BBB"),
        ]
    )
    second = _wr_file(
        [
            ("111111111111-FNSKU", "Later title", 111111111111, "X00ZZZ"),
            ("019826846593-FNSKU", "Leading zero", "019826846593", "X00CCC"),
        ]
    )
    result = format_wr_sku_catalogs([("a.xlsx", first), ("b.xlsx", second)])
    sheet = openpyxl.load_workbook(BytesIO(result.file_bytes))[SHEET_NAME]
    assert [sheet.cell(row, 1).value for row in range(2, 5)] == [
        111111111111,
        222222222222,
        "019826846593",
    ]
    assert sheet["D2"].value == "First title"
    assert sheet["C2"].value == "X00AAA"
    assert sheet["E2"].value == "New"
    assert sheet["A4"].number_format != "0" or isinstance(sheet["A4"].value, str)
    assert result.duplicates_removed == 1
    assert result.row_count == 3


def test_wr_parser_does_not_accept_a_shipment_plan():
    content = _plan(
        [("198268465935-FNSKU", "Crew", "X0058EHT67", "UPC:198268465935", "New")]
    )
    with pytest.raises(ProductCatalogFormatterError, match="WR SKU Update"):
        format_wr_sku_catalogs([("plan.xlsx", content)])


@pytest.mark.skipif(
    not SAMPLE_WR_INPUT.exists() or not SAMPLE_WR_OUTPUT.exists(),
    reason="Sample WR SKU Update workbooks are not present",
)
def test_real_wr_sample_matches_the_catalog_workbook():
    result = format_wr_sku_catalogs([("input.xlsx", SAMPLE_WR_INPUT.read_bytes())])
    expected = openpyxl.load_workbook(SAMPLE_WR_OUTPUT, data_only=False)
    got = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=False)
    assert got.sheetnames == [SHEET_NAME]
    assert [got.active.cell(1, col).value for col in range(1, 6)] == [
        expected.active.cell(1, col).value for col in range(1, 6)
    ]
    assert [got.active.cell(2, col).value for col in range(1, 6)] == [
        expected.active.cell(2, col).value for col in range(1, 6)
    ]
    assert got.active["A2"].number_format == expected.active["A2"].number_format
    assert result.row_count == 1
    assert result.duplicates_removed == 0
