"""Tests for the NFA and SMW Carton Numbers generator."""
from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from app.services.fba_carton_numbers import (
    OUTPUT_FILENAME,
    FbaCartonNumbersError,
    generate_carton_numbers,
    parse_carton_numbers,
    sheet_title,
)

SAMPLE_DIR = Path(r"c:\Users\Administrator\Downloads\EXTERNALFW__EXTERNALFBA")
SAMPLE_OUTPUT = Path(r"c:\Users\Administrator\Downloads\Carton Numbers.xlsx")
SAMPLE_ORDER = [
    "FBA19RS2FS05.xls",
    "FBA19RS21F99.xls",
    "FBA19RS3KGRK.xls",
    "FBA19RS5VKHP.xls",
    "FBA19RS4L75B.xls",
]


def carton_dump(
    carton_ids: list[object],
    *,
    shipment_id: str | None = "FBA19TEST001",
    footer: int | None = None,
) -> bytes:
    """A spaced-column carton dump like the real FBA exports."""
    workbook = Workbook()
    sheet = workbook.active
    if shipment_id is not None:
        sheet.append(["PO#:", shipment_id])
    sheet.append([])
    sheet.append(["Sku", None, "UPC", None, "Description", None, None, "Qty", None, "Weight"])
    sheet.append([])
    for index, carton_id in enumerate(carton_ids):
        sheet.append(["Carton#:", carton_id, None, None, None, None, None, None, None, None, None, 21.5])
        sheet.append([])
        sheet.append([f"SKU-{index}", None, "193392648759", None, "ITEM", None, None, 2.0, None, 1.0])
        sheet.append([None, None, None, None, None, "Total", None, 2.0, None, " 1.00"])
        sheet.append([])
    if footer is not None:
        sheet.append(["Total Ctns:", footer, "Total Units:", 2 * footer])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def zip_of(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def load(result) -> openpyxl.Workbook:
    return openpyxl.load_workbook(BytesIO(result.file_bytes))


def test_one_sheet_per_shipment_with_sample_layout():
    result = generate_carton_numbers(
        [
            ("FBA19AAAA001.xlsx", carton_dump(["00001928250570412008", "00008273992079477459"], shipment_id="FBA19AAAA001")),
            ("FBA19BBBB002.xlsx", carton_dump(["00001928250570411810"], shipment_id="FBA19BBBB002")),
        ]
    )
    assert result.filename == OUTPUT_FILENAME == "Carton Numbers.xlsx"
    assert result.shipment_count == 2
    assert result.box_count == 3

    workbook = load(result)
    try:
        assert workbook.sheetnames == ["FBA19AAAA001", "FBA19BBBB002"]
        sheet = workbook["FBA19AAAA001"]
        assert sheet["A1"].value == "Original Carton"
        assert sheet["B1"].value is None
        assert sheet["C1"].value == "Box Number"
        assert sheet["A1"].font.b and sheet["C1"].font.b
        assert sheet["A2"].value == "Carton#:"
        assert sheet["B2"].value == "00001928250570412008"
        assert sheet["B2"].data_type == "s"
        assert sheet["C2"].value == 1 and type(sheet["C2"].value) is int
        assert sheet["B3"].value == "00008273992079477459"
        assert sheet["C3"].value == 2
        assert sheet.max_row == 3
        assert sheet.max_column == 3
        assert sheet.column_dimensions["B"].width == pytest.approx(21.33203125)
    finally:
        workbook.close()


def test_box_numbers_follow_carton_order_and_keep_leading_zeros():
    result = generate_carton_numbers(
        [("a.xlsx", carton_dump(["00000000000000000009", "00000000000000000002", "00000000000000000005"]))]
    )
    workbook = load(result)
    try:
        sheet = workbook["FBA19TEST001"]
        assert [sheet.cell(r, 2).value for r in range(2, 5)] == [
            "00000000000000000009",
            "00000000000000000002",
            "00000000000000000005",
        ]
        assert [sheet.cell(r, 3).value for r in range(2, 5)] == [1, 2, 3]
    finally:
        workbook.close()


def test_zip_input_matches_loose_files_and_skips_junk():
    one = carton_dump(["00000000000000000001"], shipment_id="FBA19AAAA001")
    two = carton_dump(["00000000000000000002"], shipment_id="FBA19BBBB002")
    archive = zip_of(
        {
            "folder/FBA19AAAA001.xlsx": one,
            "folder/FBA19BBBB002.xlsx": two,
            "folder/readme.txt": b"hello",
            "__MACOSX/folder/._FBA19AAAA001.xlsx": b"junk",
            "folder/~$FBA19AAAA001.xlsx": b"lock",
        }
    )
    from_zip = generate_carton_numbers([("shipments.zip", archive)])
    loose = generate_carton_numbers([("FBA19AAAA001.xlsx", one), ("FBA19BBBB002.xlsx", two)])

    assert from_zip.shipment_count == 2
    assert from_zip.skipped_count == 1
    assert [s.carton_ids for s in from_zip.shipments] == [s.carton_ids for s in loose.shipments]
    assert [s.shipment_id for s in from_zip.shipments] == ["FBA19AAAA001", "FBA19BBBB002"]


def test_loose_files_and_zip_can_be_mixed_in_upload_order():
    loose = carton_dump(["00000000000000000001"], shipment_id="FBA19AAAA001")
    zipped = carton_dump(["00000000000000000002"], shipment_id="FBA19BBBB002")
    result = generate_carton_numbers(
        [("FBA19AAAA001.xlsx", loose), ("more.zip", zip_of({"FBA19BBBB002.xlsx": zipped}))]
    )
    assert [s.shipment_id for s in result.shipments] == ["FBA19AAAA001", "FBA19BBBB002"]


def test_shipment_id_falls_back_to_file_name():
    result = generate_carton_numbers(
        [("FBA19FALLBACK.xlsx", carton_dump(["00000000000000000001"], shipment_id=None))]
    )
    assert result.shipments[0].shipment_id == "FBA19FALLBACK"


def test_duplicate_shipment_across_files_is_rejected():
    dump = carton_dump(["00000000000000000001"], shipment_id="FBA19SAME001")
    with pytest.raises(FbaCartonNumbersError, match="FBA19SAME001.*both"):
        generate_carton_numbers([("one.xlsx", dump), ("two.xlsx", dump)])


def test_footer_total_must_match_carton_rows():
    good = carton_dump(["00000000000000000001", "00000000000000000002"], footer=2)
    assert parse_carton_numbers(good, "ok.xlsx").carton_ids == (
        "00000000000000000001",
        "00000000000000000002",
    )
    bad = carton_dump(["00000000000000000001", "00000000000000000002"], footer=3)
    with pytest.raises(FbaCartonNumbersError, match="3 cartons.*2 Carton#"):
        parse_carton_numbers(bad, "bad.xlsx")


def test_numeric_carton_number_is_rejected_instead_of_losing_digits():
    with pytest.raises(FbaCartonNumbersError, match="stored as a number"):
        parse_carton_numbers(carton_dump([1928250570412008]), "numeric.xlsx")


def test_file_without_carton_rows_is_rejected_with_its_name():
    with pytest.raises(FbaCartonNumbersError, match=r'"empty\.xlsx".*Carton#'):
        parse_carton_numbers(carton_dump([]), "empty.xlsx")


def test_unsupported_and_empty_uploads_are_rejected():
    with pytest.raises(FbaCartonNumbersError, match="No files"):
        generate_carton_numbers([])
    with pytest.raises(FbaCartonNumbersError, match="not supported"):
        generate_carton_numbers([("notes.pdf", b"%PDF")])
    with pytest.raises(FbaCartonNumbersError, match="is empty"):
        generate_carton_numbers([("a.xlsx", b"")])
    with pytest.raises(FbaCartonNumbersError, match="No Excel files"):
        generate_carton_numbers([("only-text.zip", zip_of({"notes.txt": b"hi"}))])


def test_bad_and_nested_zips_are_rejected():
    with pytest.raises(FbaCartonNumbersError, match="not a valid zip"):
        generate_carton_numbers([("broken.zip", b"definitely not a zip")])
    inner = zip_of({"a.xlsx": carton_dump(["00000000000000000001"])})
    with pytest.raises(FbaCartonNumbersError, match="zip inside a zip"):
        generate_carton_numbers([("outer.zip", zip_of({"inner.zip": inner}))])


def test_a_bad_file_inside_a_zip_names_the_zip_and_file():
    archive = zip_of({"good.xlsx": carton_dump(["00000000000000000001"]), "bad.xlsx": carton_dump([])})
    with pytest.raises(FbaCartonNumbersError, match=r"shipments\.zip > bad\.xlsx"):
        generate_carton_numbers([("shipments.zip", archive)])


def test_sheet_title_is_excel_safe():
    assert sheet_title("FBA19AAAA001") == "FBA19AAAA001"
    assert sheet_title("FBA/19:A*B?[C]") == "FBA-19-A-B--C-"
    assert len(sheet_title("X" * 80)) == 31
    assert sheet_title("   ") == "Shipment"


@pytest.mark.skipif(
    not SAMPLE_OUTPUT.exists() or not all((SAMPLE_DIR / name).exists() for name in SAMPLE_ORDER),
    reason="sample files not present",
)
@pytest.mark.parametrize("as_zip", [False, True])
def test_matches_sample_carton_numbers_workbook(as_zip: bool):
    files = [(name, (SAMPLE_DIR / name).read_bytes()) for name in SAMPLE_ORDER]
    uploads = [("EXTERNALFBA.zip", zip_of(dict(files)))] if as_zip else files
    result = generate_carton_numbers(uploads)

    got = load(result)
    expected = openpyxl.load_workbook(SAMPLE_OUTPUT)
    try:
        assert got.sheetnames == expected.sheetnames
        for name in expected.sheetnames:
            exp_sheet, got_sheet = expected[name], got[name]
            for row in range(1, exp_sheet.max_row + 1):
                for col in (1, 2, 3):
                    exp_val = exp_sheet.cell(row, col).value
                    got_val = got_sheet.cell(row, col).value
                    assert got_val == exp_val, f"{name} {row},{col}: {got_val!r} != {exp_val!r}"
                    assert type(got_val) is type(exp_val), f"{name} {row},{col} type"
            assert got_sheet.max_row == len(got_sheet["B"]) and got_sheet.max_row >= 2
    finally:
        got.close()
        expected.close()
