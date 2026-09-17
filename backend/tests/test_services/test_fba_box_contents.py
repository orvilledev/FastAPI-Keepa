"""Tests for FBA Carton Detail → Box Contents + Dimensions conversion."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.services.fba_box_contents import (
    BOX_CONTENTS_SHEET,
    DEFAULT_OUTPUT_FILENAME,
    DIMENSIONS_SHEET,
    PIVOT_START_COL,
    FbaBoxContentsError,
    generate_fba_box_contents,
    parse_carton_detail,
    sanitize_download_filename,
)

SAMPLE_INPUT = Path(r"c:\Users\Administrator\Downloads\FBA19NT5WH3J.xls")
SAMPLE_OUTPUT = Path(r"c:\Users\Administrator\Downloads\FBA19NT5WH3J Output.xlsx")


def _xlsx_carton_detail(rows: list[list[object]], shipment_id: str = "FBATEST1") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["FBA Carton Detail", "Page -1 of 1", shipment_id])
    sheet.append(["Sku", "Size", "UPC", "Description", "Qty", "Weight", "Carton Length", "Carton Width", "Carton Height"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _assert_text(cell) -> None:
    assert cell.value is None or isinstance(cell.value, str), f"{cell.coordinate} should be text, got {type(cell.value).__name__}"
    if cell.value not in (None, ""):
        assert cell.data_type == "s"
        assert not isinstance(cell.value, (int, float, bool))


def _assert_number(cell) -> None:
    assert cell.value is None or isinstance(cell.value, (int, float)), (
        f"{cell.coordinate} should be a number, got {type(cell.value).__name__}: {cell.value!r}"
    )
    assert not isinstance(cell.value, bool)
    if isinstance(cell.value, str):
        raise AssertionError(f"{cell.coordinate} is a string, expected a number")


def test_sanitize_filename_adds_output_suffix():
    assert sanitize_download_filename("FBA19NT5WH3J.xls") == "FBA19NT5WH3J Output.xlsx"
    assert sanitize_download_filename("FBA19NT5WH3J Output.xlsx") == "FBA19NT5WH3J Output.xlsx"
    assert sanitize_download_filename(None) == DEFAULT_OUTPUT_FILENAME


def _assert_header_style(cell) -> None:
    assert cell.font.bold is True, f"{cell.coordinate} header should be bold"
    assert (cell.fill.fgColor.rgb or "").upper().endswith("E0B0FF"), (
        f"{cell.coordinate} header fill should be #E0B0FF, got {cell.fill.fgColor.rgb!r}"
    )


def _assert_not_header_style(cell) -> None:
    fill_rgb = (cell.fill.fgColor.rgb or "").upper() if cell.fill.fgColor else ""
    assert not fill_rgb.endswith("E0B0FF"), f"{cell.coordinate} data cell should not use header fill"


def test_parse_requires_carton_headers():
    raw = _xlsx_carton_detail([["001830-M", "196248203201", "SOCK", 48, 5.2]])
    with pytest.raises(FbaBoxContentsError, match="Carton#"):
        parse_carton_detail(raw)


def test_generate_writes_strict_types_and_tabs():
    raw = _xlsx_carton_detail(
        [
            ["Carton#:", "00001111", 18.71, 13.71, 7.29],
            ["001830-M", "196248203201", "W EVY CBLE CRW", 48.0, 5.2896],
            ["Total", 48, " 8.15"],
            ["Carton#:", "00002222", 24.0, 16.0, 11.0],
            ["001735-L", "196009151451", "EVY ANCH CRW", 10, 1.23],
            ["001735-L", "196009151451", "EVY ANCH CRW", 5, 0.61],
            ["001572-S", "195440266199", "W HIKE", 0, 0.0],
            ["Total", 15, " 3.00"],
        ]
    )
    result = generate_fba_box_contents(raw, "FBA19NT5WH3J.xls")
    assert result.filename == "FBA19NT5WH3J Output.xlsx"
    assert result.box_count == 2
    assert result.row_count == 4
    assert result.total_qty == 63
    assert result.shipment_id == "FBATEST1"

    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes))
    try:
        assert workbook.sheetnames == [BOX_CONTENTS_SHEET, DIMENSIONS_SHEET]
        contents = workbook[BOX_CONTENTS_SHEET]
        assert contents["A1"].value == "UPC"
        assert contents["B1"].value == "Box #"
        assert contents["C1"].value == "QTY"
        _assert_text(contents["A1"])
        _assert_text(contents["B1"])
        _assert_text(contents["C1"])

        _assert_text(contents["A2"])
        _assert_number(contents["B2"])
        _assert_number(contents["C2"])
        assert contents["A2"].value == "196248203201"
        assert contents["B2"].value == 1
        assert type(contents["B2"].value) is int
        assert contents["C2"].value == 48
        assert type(contents["C2"].value) is int

        assert contents["A3"].value == "196009151451"
        assert contents["B3"].value == 2
        assert contents["C3"].value == 10
        assert contents["C5"].value == 0
        assert type(contents["C5"].value) is int

        assert contents["B7"].value == "Total QTY"
        assert contents["C7"].value == 63
        assert type(contents["C7"].value) is int

        assert contents["G1"].value == "Sum of QTY"
        assert contents["H1"].value == "Column Labels"
        assert contents["G2"].value == "Row Labels"
        assert contents["H2"].value == 1
        assert type(contents["H2"].value) is int
        assert contents["I2"].value == 2
        last_col = get_column_letter(PIVOT_START_COL + 1 + 2)
        assert contents[f"{last_col}2"].value == "Grand Total"

        # Pivot UPCs are strings, sorted, and duplicate box rows are summed.
        pivot_upcs = [contents.cell(row, PIVOT_START_COL).value for row in range(3, 6)]
        assert pivot_upcs == ["195440266199", "196009151451", "196248203201"]
        for row in range(3, 6):
            _assert_text(contents.cell(row, PIVOT_START_COL))
        assert contents["I4"].value == 15
        assert type(contents["I4"].value) is int
        assert contents["H5"].value == 48
        assert contents["I3"].value is None  # qty 0 is left blank in the pivot

        dims = workbook[DIMENSIONS_SHEET]
        assert [dims.cell(1, col).value for col in range(1, 6)] == [
            "Box #",
            "Weight",
            "Length",
            "Width",
            "Height",
        ]
        for col in range(1, 6):
            _assert_text(dims.cell(1, col))
        _assert_number(dims["A2"])
        _assert_number(dims["B2"])
        _assert_number(dims["C2"])
        assert dims["A2"].value == 1
        assert type(dims["A2"].value) is int
        assert dims["B2"].value == 5.3  # ROUNDUP(5.2896, 1)
        assert dims["C2"].value == 18.71
        assert dims["D3"].value == 16
        assert type(dims["D3"].value) is int
        assert dims["E3"].value == 11
        assert type(dims["E3"].value) is int
    finally:
        workbook.close()


def test_output_headers_use_mauve_bold_fill():
    raw = _xlsx_carton_detail(
        [
            ["Carton#:", "00001111", 18.71, 13.71, 7.29],
            ["001830-M", "196248203201", "W EVY CBLE CRW", 48.0, 5.2896],
            ["Total", 48, " 8.15"],
            ["Carton#:", "00002222", 24.0, 16.0, 11.0],
            ["001735-L", "196009151451", "EVY ANCH CRW", 10, 1.23],
            ["Total", 10, " 3.00"],
        ]
    )
    result = generate_fba_box_contents(raw, "style-check.xls")
    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes))
    try:
        contents = workbook[BOX_CONTENTS_SHEET]
        dims = workbook[DIMENSIONS_SHEET]

        for coord in ("A1", "B1", "C1", "G1", "H1", "G2", "H2", "I2", "J2"):
            _assert_header_style(contents[coord])
        assert contents["J2"].value == "Grand Total"

        for col in range(1, 6):
            _assert_header_style(dims.cell(1, col))

        # Data / totals keep values and are not painted with the header fill.
        _assert_not_header_style(contents["A2"])
        _assert_not_header_style(contents["B2"])
        _assert_not_header_style(contents["C2"])
        _assert_not_header_style(contents["G3"])
        _assert_not_header_style(dims["A2"])
        _assert_not_header_style(dims["B2"])
        assert contents["A2"].value == "196248203201"
        assert contents["B2"].value == 1
        assert contents["C2"].value == 48
        assert dims["B2"].value == 5.3
    finally:
        workbook.close()


def test_numeric_upc_is_written_as_excel_text():
    raw = _xlsx_carton_detail(
        [
            ["Carton#:", "BOX1", 10, 8, 6],
            ["SKU1", 196248203201, "DESC", 2, 1.0],
            ["Total", 2, "1.0"],
        ]
    )
    result = generate_fba_box_contents(raw, "numeric-upc.xls")
    workbook = openpyxl.load_workbook(BytesIO(result.file_bytes))
    try:
        cell = workbook[BOX_CONTENTS_SHEET]["A2"]
        assert cell.value == "196248203201"
        assert isinstance(cell.value, str)
        assert cell.data_type == "s"
        assert workbook[BOX_CONTENTS_SHEET]["B2"].value == 1
        assert type(workbook[BOX_CONTENTS_SHEET]["B2"].value) is int
    finally:
        workbook.close()


def test_size_column_layout_still_reads_upc():
    raw = _xlsx_carton_detail(
        [
            ["Carton#:", "BOX1", 10, 8, 6],
            ["001830-M", "M", "196248203201", "DESC", 2, 1.04],
            ["Total", 2, " 2.00"],
        ]
    )
    parsed = parse_carton_detail(raw)
    assert parsed.rows[0].upc == "196248203201"
    assert parsed.rows[0].qty == 2
    assert parsed.cartons[0].weight == 1.1


def test_weight_uses_roundup_not_half_even():
    raw = _xlsx_carton_detail(
        [
            ["Carton#:", "BOX1", 24, 16, 11],
            ["A-L", "195440266199", "DESC", 1, 11.816],
            ["Total", 1, " 15.50"],
        ]
    )
    parsed = parse_carton_detail(raw)
    assert parsed.cartons[0].weight == 11.9


@pytest.mark.skipif(not SAMPLE_INPUT.exists() or not SAMPLE_OUTPUT.exists(), reason="Sample FBA files not present")
def test_sample_input_matches_sample_output():
    expected = openpyxl.load_workbook(SAMPLE_OUTPUT, data_only=True)
    try:
        exp_contents = expected[BOX_CONTENTS_SHEET]
        exp_dims = expected[DIMENSIONS_SHEET]
        expected_rows = []
        for row in range(2, exp_contents.max_row + 1):
            upc = exp_contents.cell(row, 1).value
            box = exp_contents.cell(row, 2).value
            qty = exp_contents.cell(row, 3).value
            if upc is None and box == "Total QTY":
                expected_total = qty
                break
            if upc is None:
                continue
            expected_rows.append((str(upc), box, qty))

        expected_pivot: dict[tuple[object, object], object] = {}
        box_headers = {}
        for col in range(PIVOT_START_COL + 1, exp_contents.max_column + 1):
            header = exp_contents.cell(2, col).value
            if header is not None and header != "Grand Total":
                box_headers[col] = header
        grand_col = max(box_headers) + 1
        for row in range(3, exp_contents.max_row + 1):
            label = exp_contents.cell(row, PIVOT_START_COL).value
            if not label or label == "Grand Total":
                continue
            for col, header in box_headers.items():
                value = exp_contents.cell(row, col).value
                if value not in (None, ""):
                    expected_pivot[(str(label), header)] = value
            expected_pivot[(str(label), "Grand Total")] = exp_contents.cell(row, grand_col).value

        expected_dims = []
        for row in range(2, exp_dims.max_row + 1):
            if exp_dims.cell(row, 1).value is None:
                continue
            expected_dims.append(
                tuple(exp_dims.cell(row, col).value for col in range(1, 6))
            )
    finally:
        expected.close()

    result = generate_fba_box_contents(SAMPLE_INPUT.read_bytes(), SAMPLE_INPUT.name)
    got = openpyxl.load_workbook(BytesIO(result.file_bytes), data_only=True)
    try:
        contents = got[BOX_CONTENTS_SHEET]
        dims = got[DIMENSIONS_SHEET]

        got_rows = []
        got_total = None
        for row in range(2, contents.max_row + 1):
            upc = contents.cell(row, 1).value
            box = contents.cell(row, 2).value
            qty = contents.cell(row, 3).value
            if upc is None and box == "Total QTY":
                got_total = qty
                break
            if upc is None:
                continue
            assert isinstance(upc, str)
            assert isinstance(box, int)
            assert isinstance(qty, (int, float)) and not isinstance(qty, bool)
            got_rows.append((upc, box, qty))

        assert got_rows == expected_rows
        assert got_total == expected_total
        assert result.total_qty == expected_total

        got_pivot: dict[tuple[object, object], object] = {}
        box_headers = {}
        for col in range(PIVOT_START_COL + 1, contents.max_column + 1):
            header = contents.cell(2, col).value
            if header is not None and header != "Grand Total":
                box_headers[col] = header
                assert isinstance(header, int)
        grand_col = max(box_headers) + 1
        for row in range(3, contents.max_row + 1):
            label = contents.cell(row, PIVOT_START_COL).value
            if not label or label == "Grand Total":
                continue
            assert isinstance(label, str)
            for col, header in box_headers.items():
                value = contents.cell(row, col).value
                if value not in (None, ""):
                    assert isinstance(value, (int, float)) and not isinstance(value, bool)
                    got_pivot[(label, header)] = value
            got_pivot[(label, "Grand Total")] = contents.cell(row, grand_col).value

        assert got_pivot == expected_pivot

        got_dims = []
        for row in range(2, dims.max_row + 1):
            if dims.cell(row, 1).value is None:
                continue
            values = tuple(dims.cell(row, col).value for col in range(1, 6))
            for value in values:
                assert isinstance(value, (int, float)) and not isinstance(value, bool)
            got_dims.append(values)

        assert len(got_dims) == len(expected_dims)
        for got_row, exp_row in zip(got_dims, expected_dims):
            assert got_row[0] == exp_row[0]
            assert got_row[1] == pytest.approx(exp_row[1], rel=0, abs=1e-9)
            assert got_row[2:] == exp_row[2:]
    finally:
        got.close()
